import base64
import io
import json
import sqlite3
import struct
import zlib

import numpy as np
import pytest
from openpyxl import load_workbook

from spc.data import Dataset
from spc.db import Database
from spc.db.database import SCHEMA
from spc.profile import (
    ExtraField, ReportTemplate, default_target_table, merge_target_table, missing_fields, resolve_analysis, snapshot,
    table_for_service, validate_analysis,
)
from spc.report import generate, report_xlsx, reproduce
from spc.report.meta import ReportMeta
from spc.service import AnalysisRequest, analyze
from tests.conftest import PASSWORD, logged_in_client, make_app
from tests.test_api import err, upload


def png_uri(size=1, extra=b"") -> str:
    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))
    raw = b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", size, size, 8, 2, 0, 0, 0)) \
        + chunk(b"IDAT", zlib.compress(b"\x00\xff\x00\x00" * size * size)) + chunk(b"IEND", b"") + extra
    return "data:image/png;base64," + base64.b64encode(raw).decode()


REPORT = {
    "organization": "Acme Quality Lab", "title": "Acme process capability report", "form_no": "QF-7.3-12", "revision": "C",
    "footer": "Acme internal. Valid without signature.", "accent": "#aa2255", "logo": png_uri(), "language": "zh-TW",
    "show_element_21": False, "show_element_22": False,
    "extra_fields": [{"key": "drawing_no", "label_en": "Drawing number", "label_zh": "圖面編號", "required": True},
                     {"key": "ppap_level", "label_en": "PPAP level", "label_zh": "PPAP 等級", "required": False}],
    "required": ["part_number", "characteristic"],
}
ANALYSIS = {"alpha": 0.01, "edition": "draft", "target_confidence": 0.999, "stability_mode": "strict",
            "targets": {"production": {"major": [2.0, 1.8]}}}


# ------------------------------------------------------------------ the profile rules

def test_report_template_round_trips_and_normalises():
    t = ReportTemplate.from_dict(REPORT)
    assert t.organization == "Acme Quality Lab" and t.extra_fields[0] == ExtraField("drawing_no", "Drawing number", "圖面編號", True)
    assert ReportTemplate.from_dict(t.to_dict()) == t
    assert ReportTemplate.from_dict(None) == ReportTemplate() and ReportTemplate().show_element_21 is True
    assert t.extra_fields[0].label("zh-TW") == "圖面編號" and t.extra_fields[1].label("en") == "PPAP level"
    assert ReportTemplate.from_dict({"organization": "  padded  "}).organization == "padded"


@pytest.mark.parametrize("bad", [
    {"unknown": 1}, {"accent": "red"}, {"accent": "#12345"}, {"language": "fr"}, {"organization": "x" * 201},
    {"show_element_21": "yes"}, {"required": ["nonsense"]}, {"required": ["process", "process"]},
    {"extra_fields": [{"key": "Bad Key", "label_en": "x", "label_zh": "", "required": False}]},
    {"extra_fields": [{"key": "a", "label_en": "", "label_zh": "", "required": False}]},
    {"extra_fields": [{"key": "a", "label_en": "x"}, {"key": "a", "label_en": "y"}]},
    {"extra_fields": [{"key": f"k{i}", "label_en": "x"} for i in range(13)]},
    {"extra_fields": [{"key": "a", "label_en": "x", "evil": 1}]},
])
def test_report_template_refuses_bad_settings(bad):
    with pytest.raises(ValueError):
        ReportTemplate.from_dict(bad)


@pytest.mark.parametrize("logo", [
    "data:image/svg+xml;base64,PHN2Zz48c2NyaXB0Pjwvc2NyaXB0Pjwvc3ZnPg==",  # an SVG may hold script
    "data:image/png;base64,!!!notbase64", "https://example.com/logo.png", "javascript:alert(1)",
    "data:image/png;base64," + base64.b64encode(b"GIF89a not a png").decode(),  # content does not match the type
    png_uri(extra=b"\x00" * (150 * 1024)),  # too large
])
def test_the_logo_must_be_a_real_small_png_or_jpeg(logo):
    with pytest.raises(ValueError):
        ReportTemplate.from_dict({"logo": logo})


def test_target_table_is_complete_and_validated():
    table = merge_target_table({"production": {"major": [2.0, 1.8]}})
    assert table["production"]["major"] == [2.0, 1.8]
    assert table["production"]["critical"] == [1.67, 1.67] and table["machine"]["critical"] == [2.33, 2.0]  # the draft's cells stay
    assert sum(len(row) for row in table.values()) == 12  # every cell is there: the snapshot is complete
    assert merge_target_table(None) == default_target_table()
    for bad in ({"production": {"major": [1.0, 2.0]}}, {"nonsense": {}}, {"production": {"huge": [1, 1]}},
                {"production": {"major": [0, 0]}}, {"production": {"major": [1]}}, {"production": {"major": [True, 1]}}):
        with pytest.raises(ValueError):
            merge_target_table(bad)
    from spc.core.capability.naming import Stage

    svc = table_for_service(table)
    assert svc[Stage.PRODUCTION]["major"] == (2.0, 1.8) and table_for_service(None) is None


def test_analysis_settings_are_validated():
    assert validate_analysis(ANALYSIS)["alpha"] == 0.01
    assert validate_analysis({}) == {} and validate_analysis(None) == {}
    for bad in ({"alpha": 1.5}, {"alpha": "x"}, {"alpha": True}, {"edition": "beta"}, {"stability_mode": "x"},
                {"rules": {"nonsense": 1}}, {"nonsense": 1}, {"targets": {"production": {"major": [1, 2]}}}):
        with pytest.raises(ValueError):
            validate_analysis(bad)


def test_resolve_takes_the_profile_for_unset_fields_and_names_deviations():
    defaults = {"alpha": 0.0027, "edition": "draft", "stability_mode": "random_range", "lsl": 1.0}
    analysis = {"alpha": 0.01, "edition": "draft", "stability_mode": "strict"}
    values, dev, table = resolve_analysis({"lsl": 1.0}, defaults, analysis)
    assert values["alpha"] == 0.01 and values["stability_mode"] == "strict" and dev == [] and table is None
    values, dev, _ = resolve_analysis({"alpha": 0.05, "edition": "draft"}, defaults, analysis)
    assert values["alpha"] == 0.05 and dev == ["alpha"]  # the caller's own value wins and is reported
    _, dev, table = resolve_analysis({}, defaults, {"targets": {"production": {"major": [2.0, 1.8]}}})
    assert table["production"]["major"] == [2.0, 1.8] and dev == []
    other = merge_target_table({"production": {"major": [3.0, 2.5]}})
    _, dev, table = resolve_analysis({"target_table": other}, defaults, {"targets": {"production": {"major": [2.0, 1.8]}}})
    assert dev == ["targets"] and table == other


def test_required_fields():
    tpl = ReportTemplate.from_dict(REPORT)
    meta = ReportMeta(part_number="P-1", characteristic="diameter", extra={"drawing_no": "D-9"})
    assert missing_fields(meta, tpl, meta.extra) == []
    assert missing_fields(ReportMeta(), tpl, {}) == ["part_number", "characteristic", "extra:drawing_no"]
    assert missing_fields(ReportMeta(part_number=" ", characteristic="x"), tpl, {"drawing_no": " "}) == ["part_number", "extra:drawing_no"]
    assert missing_fields(ReportMeta(target=0.0, part_number="p", characteristic="c"),
                          ReportTemplate.from_dict({"required": ["target"]}), {}) == []  # zero is a value


# ------------------------------------------------------------------ the analysis with the customer's targets

def sample(n=125):
    rng = np.random.default_rng(1)
    return Dataset.from_values(rng.normal(10, 0.1, n), subgroup=[f"L{i // 5 + 1}" for i in range(n)])


def test_customer_targets_replace_the_draft_values_in_the_verdict():
    base = dict(stage="production", lsl=9.5, usl=10.5, model="A1", characteristic_class="major")
    default = analyze(sample(), AnalysisRequest(**base))["targets"]
    assert (default["p"], default["pk"]) == (1.33, 1.33)
    table = merge_target_table({"production": {"major": [2.5, 2.4]}})
    custom = analyze(sample(), AnalysisRequest(**base, target_table=table))
    assert (custom["targets"]["p"], custom["targets"]["pk"]) == (2.5, 2.4) and custom["targets"]["verdict_pk"] == "fails"
    assert default["verdict_pk"] != custom["targets"]["verdict_pk"]


# ------------------------------------------------------------------ the database

def old_database(path, version):
    """A database as an older release made it: no profiles before 2, no monitors and no 'operator' role before 3, no studies before 4, no measurement systems before 5, no control plans and SPC roles before 6, no lots before 10, no improvement cycles before 11, no import templates before 12, no cycle marks on the points before 13 (the column is added only when the schema text did not make it)."""
    conn = sqlite3.connect(path)
    for statement in SCHEMA.split(";"):
        if not statement.strip() or ("studies" in statement and version < 4) or ("msa_systems" in statement and version < 5) or (("control_plans" in statement or "spc_people" in statement) and version < 6) or ("validation_" in statement and version < 7) or ("equipment_" in statement and version < 8) or (("trusted_signers" in statement or "report_signatures" in statement) and version < 9) or ("TABLE lots" in statement and version < 10) or ("TABLE improvements" in statement and version < 11) or ("TABLE import_templates" in statement and version < 12) or (version < 3 and "monitor" in statement):
            continue
        if version < 2 and "profiles" in statement:
            continue
        conn.execute(statement if version >= 3 else statement.replace("'viewer', 'operator'", "'viewer'"))
    conn.execute(f"PRAGMA user_version = {version}")
    conn.execute("INSERT INTO users (username, password_hash, role, created_at, updated_at) VALUES ('old', 'x', 'admin', 'a', 'a')")
    conn.execute("INSERT INTO sessions (token_hash, user_id, csrf, created_at, last_seen, expires_at) VALUES ('t', 1, 'c', 1, 1, 9e9)")
    conn.commit()
    conn.close()


@pytest.mark.parametrize("version", [1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12])
def test_an_old_database_is_migrated_without_losing_anything(tmp_path, version):
    path = tmp_path / "old.sqlite3"
    old_database(path, version)
    db = Database(path)
    assert db.one("PRAGMA user_version")[0] == 13
    assert db.one("SELECT username FROM users")["username"] == "old"  # nothing was lost
    assert db.one("SELECT COUNT(*) AS n FROM sessions")["n"] == 1  # the rebuild did not cascade into the sessions
    assert db.all("SELECT * FROM profiles") == [] and db.all("SELECT * FROM monitors") == [] and db.all("SELECT * FROM studies") == [] and db.all("SELECT * FROM msa_systems") == [] and db.all("SELECT * FROM control_plans") == [] and db.all("SELECT * FROM spc_people") == [] and db.all("SELECT * FROM validation_cases") == [] and db.all("SELECT * FROM validation_runs") == [] and db.all("SELECT * FROM equipment_links") == [] and db.all("SELECT * FROM trusted_signers") == [] and db.all("SELECT * FROM report_signatures") == [] and db.all("SELECT * FROM lots") == [] and db.all("SELECT * FROM improvements") == [] and db.all("SELECT * FROM import_templates") == []
    db.execute("INSERT INTO users (username, password_hash, role, created_at, updated_at) VALUES ('op', 'x', 'operator', 'a', 'a')")
    with pytest.raises(sqlite3.IntegrityError):
        db.execute("INSERT INTO users (username, password_hash, role, created_at, updated_at) VALUES ('bad', 'x', 'boss', 'a', 'a')")
    db.close()
    assert Database(path).one("PRAGMA user_version")[0] == 13  # opens again without migrating twice


def test_points_of_version_12_get_the_cycle_column_and_keep_their_data(tmp_path):
    """A real older file has no `cycle` column on the points (the schema text of today has it, so the test takes it out)."""
    path = tmp_path / "v12.sqlite3"
    old_database(path, 12)
    conn = sqlite3.connect(path)
    conn.execute("ALTER TABLE monitor_points DROP COLUMN cycle")
    assert "cycle" not in [r[1] for r in conn.execute("PRAGMA table_info(monitor_points)")]
    conn.commit()
    conn.close()
    db = Database(path)
    assert db.one("PRAGMA user_version")[0] == 13
    assert "cycle" in [r[1] for r in db.all("PRAGMA table_info(monitor_points)")]
    assert db.one("SELECT username FROM users")["username"] == "old"
    db.close()
    assert Database(path).one("PRAGMA user_version")[0] == 13


# ------------------------------------------------------------------ the API

@pytest.fixture
def env():
    app = make_app(max_upload=200_000)
    return app, logged_in_client(app, "admin"), logged_in_client(app, "eng"), logged_in_client(app, "view")


def make_profile(admin, name="Acme", analysis=None, report=None):
    r = admin.post("/api/profiles", json={"name": name, "analysis": ANALYSIS if analysis is None else analysis,
                                          "report": REPORT if report is None else report})
    assert r.status_code == 200, r.text
    return r.json()


def test_everybody_reads_profiles_and_only_an_admin_changes_them(env):
    app, admin, eng, view = env
    p = make_profile(admin)
    assert p["revision"] == 1 and p["analysis"]["alpha"] == 0.01 and p["report"]["organization"] == "Acme Quality Lab"
    for who in (eng, view):
        assert [x["name"] for x in who.get("/api/profiles").json()["profiles"]] == ["Acme"]
        assert who.post("/api/profiles", json={"name": "X"}).status_code == 403
        assert who.put(f"/api/profiles/{p['id']}", json={"name": "X"}).status_code == 403
        assert who.delete(f"/api/profiles/{p['id']}").status_code == 403
    assert err(admin.post("/api/profiles", json={"name": "acme"}))["code"] == "profile_name_taken"  # names ignore case
    assert err(admin.put("/api/profiles/999", json={"name": "Z"}))["code"] == "profile_not_found"
    assert err(admin.delete("/api/profiles/999"))["code"] == "profile_not_found"


def test_update_raises_the_revision_and_everything_is_audited(env):
    app, admin, eng, view = env
    p = make_profile(admin)
    new_report = {**REPORT, "revision": "D"}
    u = admin.put(f"/api/profiles/{p['id']}", json={"name": "Acme GmbH", "analysis": {"alpha": 0.05}, "report": new_report}).json()
    assert u["revision"] == 2 and u["name"] == "Acme GmbH" and u["analysis"] == {"alpha": 0.05} and u["report"]["revision"] == "D"
    other = make_profile(admin, "Other", {}, {})
    assert err(admin.put(f"/api/profiles/{other['id']}", json={"name": "ACME GMBH"}))["code"] == "profile_name_taken"
    assert admin.delete(f"/api/profiles/{other['id']}").status_code == 200
    actions = [(e["action"], e["target"]) for e in admin.get("/api/audit").json()["entries"] if e["action"].startswith("profile_")]
    assert ("profile_created", "Acme") in actions and ("profile_updated", "Acme GmbH") in actions and ("profile_deleted", "Other") in actions
    assert admin.get("/api/audit/verify").json()["ok"]


@pytest.mark.parametrize("body", [
    {"name": "A", "analysis": {"alpha": 2}}, {"name": "A", "analysis": {"nonsense": 1}},
    {"name": "A", "report": {"accent": "red"}}, {"name": "A", "report": {"logo": "javascript:alert(1)"}},
    {"name": "A", "analysis": {"targets": {"production": {"major": [1, 2]}}}},
])
def test_bad_profiles_are_refused_with_a_message(env, body):
    r = env[1].post("/api/profiles", json=body)
    assert r.status_code == 400 and err(r)["code"] == "invalid_input"


def test_meta_gives_the_draft_target_table_for_the_editor(env):
    table = env[2].get("/api/meta").json()["default_targets"]
    assert table == default_target_table() and table["production"]["major"] == [1.33, 1.33]


def test_analysis_with_a_profile_applies_it_and_names_deviations(env):
    app, admin, eng, view = env
    p = make_profile(admin)
    ds = upload(eng).json()
    url = f"/api/datasets/{ds['id']}/analyze"
    body = {"lsl": 9, "usl": 11, "model": "A1", "characteristic_class": "major", "profile_id": p["id"]}
    r = eng.post(url, json=body).json()
    assert r["params"]["alpha"] == 0.01 and r["params"]["stability_mode"] == "strict" and r["params"]["target_confidence"] == 0.999
    assert r["params"]["customer"] == "Acme" and r["stability"]["mode"] == "strict"
    assert (r["targets"]["p"], r["targets"]["pk"]) == (2.0, 1.8)  # the customer's targets
    assert r["profile"] == {"id": p["id"], "name": "Acme", "revision": 1, "deviations": []}
    r2 = eng.post(url, json={**body, "alpha": 0.05}).json()
    assert r2["params"]["alpha"] == 0.05 and r2["profile"]["deviations"] == ["alpha"]
    assert "profile" not in eng.post(url, json={"lsl": 9, "usl": 11}).json()
    assert err(eng.post(url, json={**body, "profile_id": 999}))["code"] == "profile_not_found"
    assert viewer_can_analyse(view, url, body)


def viewer_can_analyse(view, url, body):
    return view.post(url, json=body).status_code == 200


def report_body(p, **meta):
    return {"analysis": {"lsl": 9, "usl": 11, "model": "A1", "characteristic_class": "major", "profile_id": p["id"]},
            "meta": {"process": "turning", "part_number": "P-7", "characteristic": "diameter", "uncertainty": 0.02,
                     "extra": {"drawing_no": "D-100", "ppap_level": "3"}, **meta}}


def test_report_needs_the_fields_the_profile_requires(env):
    app, admin, eng, view = env
    p = make_profile(admin)
    ds = upload(eng).json()
    url = f"/api/datasets/{ds['id']}/reports"
    r = eng.post(url, json=report_body(p, part_number="", extra={}))
    e = err(r)
    assert r.status_code == 400 and e["code"] == "report_field_required" and e["params"]["fields"] == ["part_number", "extra:drawing_no"]
    r = eng.post(url, json=report_body(p, extra={"drawing_no": "D", "stray": "x"}))
    assert err(r)["code"] == "report_field_unknown" and err(r)["params"]["fields"] == ["stray"]
    nop = eng.post(url, json={"analysis": {"lsl": 9, "usl": 11}, "meta": {"extra": {"drawing_no": "D"}}})
    assert err(nop)["code"] == "report_field_unknown"  # extra fields belong to a profile
    assert eng.post(url, json=report_body(p)).status_code == 200
    assert view.post(url, json=report_body(p)).status_code == 403


def test_report_has_the_layout_of_the_customer(env):
    app, admin, eng, view = env
    p = make_profile(admin)
    ds = upload(eng).json()
    made = eng.post(f"/api/datasets/{ds['id']}/reports", json=report_body(p)).json()
    assert made["language"] == "zh-TW"  # the profile's default language: the body did not choose one
    html = eng.get(made["urls"]["html"]).text
    for text in ("Acme Quality Lab", "Acme process capability report", "QF-7.3-12", "Acme internal. Valid without signature.",
                 "圖面編號", "D-100", "PPAP 等級", "P-7"):
        assert text in html, text
    css = html.split("<style>")[1].split("</style>")[0]
    assert "#aa2255" in css and "#1f5fbf" not in css  # the accent replaces the blue of the page (data colours in the charts stay)
    assert '<img class="logo" src="data:image/png;base64,' in html
    assert html.count("<h2><span class=\"no\">21</span>") == 0 and html.count("<h2><span class=\"no\">22</span>") == 0
    assert html.count("<h2><span class=\"no\">20</span>") == 1
    assert "Acme (版次 1)" in html and "客戶專屬設定" in html
    en = eng.post(f"/api/datasets/{ds['id']}/reports", json={**report_body(p), "language": "en"}).json()
    html_en = eng.get(en["urls"]["html"]).text
    assert "Drawing number" in html_en and "Information required by Acme" in html_en and "Rev. C" in html_en and "Form QF-7.3-12" in html_en
    assert "Cw" not in html and "Cw" not in html_en


def test_report_without_a_profile_is_unchanged_and_shows_elements_21_and_22(env):
    app, admin, eng, view = env
    ds = upload(eng).json()
    made = eng.post(f"/api/datasets/{ds['id']}/reports", json={"analysis": {"lsl": 9, "usl": 11}, "meta": {"uncertainty": 0.02}}).json()
    html = eng.get(made["urls"]["html"]).text
    assert html.count("<h2><span class=\"no\">21</span>") == 1 and html.count("<h2><span class=\"no\">22</span>") == 1
    assert "#1f5fbf" in html.split("<style>")[1].split("</style>")[0] and 'class="logo"' not in html and made["language"] == "en"
    assert "profile" not in eng.get(made["urls"]["archive"]).json()


def test_the_archive_keeps_a_snapshot_so_a_later_change_does_not_touch_the_report(env):
    app, admin, eng, view = env
    p = make_profile(admin)
    ds = upload(eng).json()
    made = eng.post(f"/api/datasets/{ds['id']}/reports", json={**report_body(p), "analysis": {**report_body(p)["analysis"], "alpha": 0.05}}).json()
    html_before = eng.get(made["urls"]["html"]).text
    archive = eng.get(made["urls"]["archive"]).json()
    snap = archive["profile"]
    assert snap["name"] == "Acme" and snap["revision"] == 1 and snap["deviations"] == ["alpha"]
    assert snap["report"]["organization"] == "Acme Quality Lab"
    assert archive["request"]["target_table"]["production"]["major"] == [2.0, 1.8]  # the complete table is in the request
    assert archive["request"]["customer"] == "Acme"
    admin.put(f"/api/profiles/{p['id']}", json={"name": "Renamed", "analysis": {}, "report": {"organization": "Other Org"}})
    assert eng.get(made["urls"]["html"]).text == html_before
    wb = load_workbook(io.BytesIO(eng.get(f"{made['urls']['html']}/report.xlsx").content))
    summary = wb.worksheets[0]  # the profile's language is Chinese, so the sheet is named in Chinese
    assert "Acme Quality Lab" in str(summary["A2"].value) and "Other Org" not in str(summary["A2"].value)
    check = eng.post("/api/archive/check", content=json.dumps(archive)).json()
    assert check["integrity_ok"] and check["reproduced"]
    archive["profile"]["report"]["organization"] = "Forged Org"  # the snapshot is inside the digest
    assert not eng.post("/api/archive/check", content=json.dumps(archive)).json()["integrity_ok"]
    entry = next(e for e in admin.get("/api/audit").json()["entries"] if e["action"] == "report_created")
    assert entry["detail"]["profile"] == "Acme" and entry["detail"]["profile_revision"] == 1 and entry["detail"]["deviations"] == ["alpha"]


def test_excel_report_follows_the_profile(env):
    app, admin, eng, view = env
    p = make_profile(admin)
    ds = upload(eng).json()
    made = eng.post(f"/api/datasets/{ds['id']}/reports", json={**report_body(p), "language": "en"}).json()
    wb = load_workbook(io.BytesIO(eng.get(f"{made['urls']['html']}/report.xlsx").content))
    assert wb["Summary"]["A1"].value == "Acme process capability report"
    assert wb["Summary"]["A2"].value.startswith("Acme Quality Lab · Form QF-7.3-12 · Rev. C")
    elements = wb["Report elements"]
    nos = {elements.cell(r, 1).value for r in range(2, elements.max_row + 1)}
    assert 21 not in nos and 22 not in nos and 20 in nos and "–" in nos
    text = " ".join(str(c.value) for row in elements.iter_rows() for c in row)
    assert "Drawing number" in text and "D-100" in text and "PPAP level" in text
    annex = " ".join(str(c.value) for row in wb["Annex"].iter_rows() for c in row)
    assert "Acme (Rev. 1)" in annex


def test_text_in_a_profile_cannot_run_as_html_or_formula(env):
    app, admin, eng, view = env
    evil = {**REPORT, "organization": "<script>alert(1)</script>", "footer": "<img src=x onerror=alert(2)>", "title": "=HYPERLINK(\"x\",\"y\")",
            "extra_fields": [{"key": "k", "label_en": "<b>bold</b>", "label_zh": "", "required": False}]}
    p = make_profile(admin, "Evil", {}, evil)
    ds = upload(eng).json()
    body = report_body(p)
    body["meta"]["extra"] = {"k": "<i>v</i>"}
    made = eng.post(f"/api/datasets/{ds['id']}/reports", json={**body, "language": "en"}).json()
    html = eng.get(made["urls"]["html"]).text
    assert "<script>alert(1)</script>" not in html and "&lt;script&gt;alert(1)&lt;/script&gt;" in html
    assert "<img src=x" not in html and "<b>bold</b>" not in html and "<i>v</i>" not in html
    assert "script-src" not in eng.get(made["urls"]["html"]).headers["content-security-policy"]  # still no script at all
    wb = load_workbook(io.BytesIO(eng.get(f"{made['urls']['html']}/report.xlsx").content))
    title = wb["Summary"]["A1"]
    assert title.data_type == "s" and title.value.startswith("=HYPERLINK")


def test_deleting_a_profile_keeps_old_reports_and_unknown_ids_are_refused(env):
    app, admin, eng, view = env
    p = make_profile(admin)
    ds = upload(eng).json()
    made = eng.post(f"/api/datasets/{ds['id']}/reports", json=report_body(p)).json()
    admin.delete(f"/api/profiles/{p['id']}")
    assert "Acme Quality Lab" in eng.get(made["urls"]["html"]).text
    assert err(eng.post(f"/api/datasets/{ds['id']}/reports", json=report_body(p)))["code"] == "profile_not_found"
