"""Browser test of the real page. Skipped when Playwright or a Chromium binary is not available.

Set SPC_CHROMIUM to the path of a Chromium executable if it is not found automatically.
"""

import glob
import os
import socket
import threading
import time

import numpy as np
import pytest

playwright_sync = pytest.importorskip("playwright.sync_api")
import uvicorn  # noqa: E402

from spc.data import Dataset  # noqa: E402
from tests.conftest import PASSWORD, make_app  # noqa: E402


def find_chromium():
    candidates = [os.environ.get("SPC_CHROMIUM", "")]
    candidates += sorted(glob.glob("/opt/pw-browsers/chromium-*/chrome-linux/chrome"))
    return next((c for c in candidates if c and os.path.exists(c)), None)


@pytest.fixture(scope="module")
def app():
    return make_app(max_upload=20 * 1024 * 1024)


def sign_in(page, user="eng", password=PASSWORD, wait=True):
    page.fill("#login-user", user)
    page.fill("#login-pass", password)
    page.click("#login-form button[type=submit]")
    if wait:  # the next step must not run before the session exists (a person cannot be that fast either)
        page.wait_for_selector("#user-label", state="visible")


@pytest.fixture(scope="module")
def server(app):
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        port = s.getsockname()[1]
    srv = uvicorn.Server(uvicorn.Config(app, host="127.0.0.1", port=port, log_level="warning"))
    thread = threading.Thread(target=srv.run, daemon=True)
    thread.start()
    for _ in range(100):
        if srv.started:
            break
        time.sleep(0.05)
    yield f"http://127.0.0.1:{port}/"
    srv.should_exit = True
    thread.join(timeout=5)


@pytest.fixture(scope="module")
def browser():
    exe = find_chromium()
    with playwright_sync.sync_playwright() as p:
        try:
            b = p.chromium.launch(executable_path=exe, args=["--no-sandbox"]) if exe else p.chromium.launch()
        except Exception as exc:  # no usable browser on this machine
            pytest.skip(f"no Chromium available: {exc}")
        yield b
        b.close()


def sample(tmp_path):
    rng = np.random.default_rng(42)
    lines = ["批號,直徑"]
    for g in range(25):
        for j in range(5):
            v = 10 + rng.normal(0, 0.1) + (1.2 if (g, j) == (11, 3) else 0)
            lines.append(f"B{g + 1:02d},{v:.4f}")
    path = tmp_path / "sample.csv"
    path.write_bytes(("\n".join(lines) + "\n").encode("cp950"))
    return path


def test_full_flow_in_both_languages(server, browser, tmp_path):
    expect = playwright_sync.expect
    ctx = browser.new_context(viewport={"width": 1200, "height": 900}, locale="zh-TW")
    page = ctx.new_page()
    problems = []
    page.on("console", lambda m: problems.append(m.text) if m.type == "error" else None)
    page.on("pageerror", lambda e: problems.append(str(e)))

    page.goto(server)
    expect(page.locator("#login h2")).to_have_text("登入")  # language from the browser
    expect(page.locator("#app")).to_be_hidden()  # nothing of the program shows before the login
    sign_in(page)
    expect(page.locator("nav.tabs button[data-tab=import]")).to_have_text("1. 匯入")

    # import a Big5 file: columns are suggested from the content
    page.set_input_files("#file", str(sample(tmp_path)))
    expect(page.locator("#detected")).to_contain_text("cp950")
    expect(page.locator("#col-value")).to_have_value("直徑")
    expect(page.locator("#col-subgroup")).to_have_value("批號")
    page.click("#import-btn")
    expect(page.locator("#data-counts")).to_contain_text("共 125 筆")

    # a suspect value is only a hint. Marking needs a reason and a person.
    page.click("#suspect-btn")
    expect(page.locator("#suspect-msg")).to_contain_text("1 個可疑值")
    page.click("#suspect-select")
    page.click("#mark-btn")
    expect(page.locator("#errors")).to_contain_text("必須填寫理由")
    page.click("#errors button")
    expect(page.locator("#person-note")).to_contain_text("Eva Engineer (eng)")  # the person comes from the login
    page.fill("#reason", "typing error")
    page.click("#mark-btn")
    expect(page.locator("#data-counts")).to_contain_text("標記無效 1 筆")

    # analysis
    page.click("#to-analysis")
    expect(page.locator("#a-size-wrap")).to_be_hidden()  # the data has subgroup labels
    page.fill("#a-lsl", "9.5")
    page.fill("#a-usl", "10.5")
    page.select_option("#a-model", "A1")
    page.select_option("#a-class", "major")
    page.click("#run-btn")
    expect(page.locator("#r-names")).to_have_text("Cp 與 Cpk")
    expect(page.locator("#chart-loc svg")).to_be_visible()
    expect(page.locator("#chart-var svg")).to_be_visible()
    labels = page.locator("#chart-loc svg text.limit-label").all_text_contents()
    assert any(t.startswith("管制上限") for t in labels)

    # the language can change without losing the result
    page.select_option("#lang", "en")
    expect(page.locator("#r-names")).to_have_text("Cp and Cpk")
    expect(page.locator("#r-targets")).to_contain_text("Met, also at the lower confidence limit")

    # without a model the same data is named Pp/Ppk
    page.select_option("#a-model", "")
    page.click("#run-btn")
    expect(page.locator("#r-names")).to_have_text("Pp and Ppk")

    # the only console error allowed is the 400 that the empty-reason click caused on purpose
    assert [p for p in problems if "400" not in p] == [], problems
    ctx.close()


def test_page_does_not_overflow_on_a_phone(server, browser):
    ctx = browser.new_context(viewport={"width": 390, "height": 800}, locale="en")
    page = ctx.new_page()
    page.goto(server)
    playwright_sync.expect(page.locator("#login-form")).to_be_visible()
    assert page.evaluate("() => document.documentElement.scrollWidth <= document.documentElement.clientWidth")
    sign_in(page)
    playwright_sync.expect(page.locator("nav.tabs")).to_be_visible()
    assert page.evaluate("() => document.documentElement.scrollWidth <= document.documentElement.clientWidth")
    ctx.close()


def test_user_text_is_never_run_as_html(server, browser, tmp_path):
    """A subgroup label that looks like HTML must show as text, not run."""
    path = tmp_path / "xss.csv"
    path.write_bytes("lot,v\n<img src=x onerror=window.__pwned=1>,1.0\n<img src=x onerror=window.__pwned=1>,1.1\nb,1.2\nb,1.3\n".encode())
    ctx = browser.new_context(locale="en")
    page = ctx.new_page()
    page.goto(server)
    sign_in(page)
    page.set_input_files("#file", str(path))
    page.select_option("#col-value", "v")
    page.select_option("#col-subgroup", "lot")
    page.click("#import-btn")
    playwright_sync.expect(page.locator("#rows")).to_contain_text("<img src=x onerror")
    assert page.evaluate("() => window.__pwned === undefined")
    ctx.close()


def test_report_and_archive_check_in_the_browser(server, browser, tmp_path):
    expect = playwright_sync.expect
    ctx = browser.new_context(viewport={"width": 1200, "height": 900}, locale="en")
    page = ctx.new_page()
    problems = []
    page.on("console", lambda m: problems.append(m.text) if m.type == "error" else None)
    page.on("pageerror", lambda e: problems.append(str(e)))

    page.goto(server)
    sign_in(page)
    page.set_input_files("#file", str(sample(tmp_path)))
    page.select_option("#col-value", "直徑")
    page.select_option("#col-subgroup", "批號")
    page.click("#import-btn")
    expect(page.locator("#data-counts")).to_contain_text("125 values")

    # the sample holds one wrong value. Left in, the charts alarm and the indices cannot be called Cp/Cpk.
    page.click("#suspect-btn")
    expect(page.locator("#suspect-msg")).to_contain_text("1 suspect")
    page.click("#suspect-select")
    page.fill("#reason", "typing error")
    page.click("#mark-btn")
    expect(page.locator("#data-counts")).to_contain_text("1 marked invalid")

    page.click("#to-analysis")
    page.fill("#a-lsl", "9.5")
    page.fill("#a-usl", "10.5")
    page.select_option("#a-model", "A1")
    page.select_option("#a-class", "major")

    # the report panel is part of the result, so it only exists after an analysis
    expect(page.locator("#rp-create")).to_be_hidden()
    page.click("#run-btn")
    expect(page.locator("#r-names")).to_have_text("Cp and Cpk")
    page.fill("#rp-process", "turning")
    page.fill("#rp-machine", "CNC-07")
    page.fill("#rp-unit", "mm")
    page.fill("#rp-uncertainty", "0.0211")
    page.fill("#rp-recommendations", "keep sampling")
    page.select_option("#rp-language", "en")
    assert page.locator("#rp-msa option").count() >= 1  # "(none)" and the measurement systems
    page.click("#rp-create")
    expect(page.locator("#rp-created")).to_contain_text("was created")
    href = page.locator("#rp-open").get_attribute("href")
    assert href.startswith("/api/reports/")

    # the report opens as its own page: it needs inline styles and SVG, and must run no script
    report = ctx.new_page()
    report_problems = []
    report.on("console", lambda m: report_problems.append(m.text) if m.type == "error" else None)
    report.goto(server.rstrip("/") + href)
    expect(report.locator("h1")).to_have_text("Process study report")
    assert report.locator("svg").count() == 5
    expect(report.locator("section.el")).not_to_have_count(0)
    assert report.evaluate("() => getComputedStyle(document.querySelector('h1')).fontSize") == "24px"  # inline CSS applied (18pt)
    assert report_problems == [], report_problems
    assert "CNC-07" in report.content() and "keep sampling" in report.content()

    # download the archive and check it in the tools tab
    archive = ctx.request.get(server.rstrip("/") + page.locator("#rp-archive").get_attribute("href"))
    path = tmp_path / "archive.json"
    path.write_bytes(archive.body())
    page.click("nav.tabs button[data-tab=tools]")
    page.set_input_files("#archive-file", str(path))
    page.click("#archive-btn")
    expect(page.locator("#archive-out")).to_contain_text("Content is unchanged")
    expect(page.locator("#archive-out")).to_contain_text("gives the stored result again")

    # a changed archive is reported as changed
    import json

    changed = json.loads(path.read_text(encoding="utf-8"))
    changed["dataset"]["values"][0] += 0.5
    path.write_text(json.dumps(changed), encoding="utf-8")
    page.set_input_files("#archive-file", str(path))
    page.click("#archive-btn")
    expect(page.locator("#archive-out")).to_contain_text("was changed after the archive was made")

    assert problems == [], problems
    ctx.close()


def test_wrong_password_logout_and_a_reload_keep_the_data_private(server, browser):
    expect = playwright_sync.expect
    ctx = browser.new_context(locale="en")
    page = ctx.new_page()
    page.goto(server)
    sign_in(page, "eng", "not the password!!", wait=False)
    expect(page.locator("#errors")).to_contain_text("User name or password is wrong")
    assert page.locator("#app").is_hidden()
    page.click("#errors button")
    sign_in(page)
    expect(page.locator("#user-label")).to_contain_text("Eva Engineer (eng) · Engineer")

    page.reload()  # the session cookie keeps the login
    expect(page.locator("#user-label")).to_contain_text("Eva Engineer")
    cookie = next(c for c in ctx.cookies() if c["name"] == "spc_session")
    assert cookie["httpOnly"] and cookie["sameSite"] == "Strict"

    page.click("#logout-btn")
    expect(page.locator("#login-form")).to_be_visible()
    expect(page.locator("#app")).to_be_hidden()
    assert ctx.request.get(server.rstrip("/") + "/api/datasets").status == 401
    ctx.close()


def test_a_viewer_can_read_saved_data_but_sees_no_way_to_change_it(server, browser, app):
    expect = playwright_sync.expect
    owner = app.state.auth.list_users()[0].id
    demo = Dataset.from_values([10.0, 10.1, 9.9, 10.05, 10.02, 9.98] * 5)
    app.state.store.add(demo, owner, "demo-saved")
    ctx = browser.new_context(locale="en")
    page = ctx.new_page()
    page.goto(server)
    sign_in(page, "view")
    expect(page.locator("#tab-saved")).to_be_visible()  # a viewer starts at the saved data
    expect(page.locator("nav.tabs button[data-tab=import]")).to_be_hidden()
    expect(page.locator("nav.tabs button[data-tab=admin]")).to_be_hidden()
    expect(page.locator("#saved-datasets")).to_contain_text("demo-saved")
    expect(page.locator("#saved-datasets button", has_text="Delete")).to_have_count(0)
    page.click("#saved-datasets button:has-text('Open')")
    expect(page.locator("#data-counts")).to_contain_text("30 values")
    expect(page.locator("#mark-btn")).to_be_hidden()
    expect(page.locator("#reason")).to_be_hidden()
    ctx.close()


def test_a_new_user_must_change_the_password_before_anything_else(server, browser):
    expect = playwright_sync.expect
    ctx = browser.new_context(locale="en")
    admin = ctx.new_page()
    admin.goto(server)
    sign_in(admin, "admin")
    admin.click("nav.tabs button[data-tab=admin]")
    admin.fill("#nu-name", "newbie")
    admin.fill("#nu-display", "N. Ewbie")
    admin.select_option("#nu-role", "engineer")
    admin.fill("#nu-pass", "temporary-pass-1")
    admin.click("#nu-create")
    expect(admin.locator("#admin-msg")).to_contain_text("User newbie created")
    expect(admin.locator("#admin-users")).to_contain_text("N. Ewbie")
    admin.click("#audit-verify")
    expect(admin.locator("#audit-verdict")).to_contain_text("Chain intact")
    expect(admin.locator("#audit-log")).to_contain_text("user_created")
    ctx.close()

    ctx = browser.new_context(locale="en")
    page = ctx.new_page()
    page.goto(server)
    sign_in(page, "newbie", "temporary-pass-1")
    expect(page.locator("#password-forced")).to_be_visible()
    page.click("nav.tabs button[data-tab=tools]")  # no other tab opens
    expect(page.locator("#tab-password")).to_be_visible()
    page.fill("#pw-current", "temporary-pass-1")
    page.fill("#pw-new", "a much better password")
    page.fill("#pw-repeat", "something different!!")
    page.click("#password-form button[type=submit]")
    expect(page.locator("#errors")).to_contain_text("The two new passwords differ")
    page.click("#errors button")
    page.fill("#pw-repeat", "a much better password")
    page.click("#password-form button[type=submit]")
    expect(page.locator("#pw-done")).to_be_visible()
    page.click("nav.tabs button[data-tab=tools]")
    expect(page.locator("#tab-tools")).to_be_visible()
    ctx.close()


def test_fitted_distribution_in_the_browser(server, browser, tmp_path):
    expect = playwright_sync.expect
    rng = np.random.default_rng(3)
    lines = ["lot,v"] + [f"L{i // 5 + 1},{9 + rng.lognormal(0, 0.4):.4f}" for i in range(125)]
    path = tmp_path / "skew.csv"
    path.write_text("\n".join(lines) + "\n")
    ctx = browser.new_context(viewport={"width": 1200, "height": 900}, locale="en")
    page = ctx.new_page()
    problems = []
    page.on("pageerror", lambda e: problems.append(str(e)))
    page.goto(server)
    sign_in(page)
    page.set_input_files("#file", str(path))
    page.select_option("#col-value", "v")
    page.select_option("#col-subgroup", "lot")
    page.click("#import-btn")
    page.click("#to-analysis")
    page.fill("#a-lsl", "8.5")
    page.fill("#a-usl", "13")
    page.select_option("#a-class", "major")
    page.select_option("#a-model", "A2")

    expect(page.locator("#a-method")).to_be_hidden()  # only for a non-normal distribution
    expect(page.locator("#r-dist-block")).to_be_hidden()
    page.click("#run-btn")
    expect(page.locator("#r-normality")).to_contain_text("Normality test")
    expect(page.locator("#r-dist-block")).to_be_hidden()

    page.select_option("#a-dist", "lognormal")
    expect(page.locator("#a-method")).to_be_visible()
    page.fill("#a-boot", "40")
    page.click("#run-btn")
    expect(page.locator("#r-dist-block")).to_be_visible()
    expect(page.locator("#r-dist-line")).to_contain_text("Lognormal (chosen by you) · method .G")
    expect(page.locator("#r-names")).to_have_text("Cp and Cpk")
    expect(page.locator("#r-indices")).to_contain_text("Cpk.G")
    expect(page.locator("#r-targets")).to_contain_text("Cpk.G")
    expect(page.locator("#r-dist-boot")).to_contain_text("seed 20260701")
    assert page.locator("#r-dist-cands tr").count() == 8  # header and seven families

    page.select_option("#a-dist", "empirical")  # needs 2000 values: a clear message, no crash
    assert page.locator("#a-method option[value=Z]").evaluate("o => o.disabled")
    page.click("#run-btn")
    expect(page.locator("#errors")).to_contain_text("2000")
    page.click("#errors button")

    page.select_option("#a-dist", "auto")
    page.select_option("#a-method", "Z")
    page.select_option("#lang", "zh-TW")
    page.click("#run-btn")
    expect(page.locator("#r-dist-line")).to_contain_text("自動選擇")
    expect(page.locator("#r-indices")).to_contain_text(".Z")
    assert problems == [], problems
    ctx.close()


def test_restart_of_the_individuals_chart_in_the_browser(server, browser, tmp_path):
    expect = playwright_sync.expect
    rng = np.random.default_rng(2)
    values = np.r_[rng.normal(10, 0.1, 30), rng.normal(10.6, 0.1, 30)]  # a tool change moved the level
    path = tmp_path / "ind.csv"
    path.write_text("v\n" + "\n".join(f"{v:.4f}" for v in values) + "\n")
    ctx = browser.new_context(viewport={"width": 1200, "height": 900}, locale="en")
    page = ctx.new_page()
    problems = []
    page.on("pageerror", lambda e: problems.append(str(e)))
    page.goto(server)
    sign_in(page)
    page.set_input_files("#file", str(path))
    page.select_option("#col-value", "v")
    page.click("#import-btn")
    expect(page.locator("#restart-list")).to_have_text("No restarts.")

    page.click("#unrestart-btn")  # nothing selected
    expect(page.locator("#errors")).to_contain_text("Select")
    page.click("#errors button")
    page.locator("#rows tr:nth-child(32) input").check()  # No. 31: the first value after the tool change
    page.click("#restart-btn")
    expect(page.locator("#errors")).to_contain_text("reason")  # a reason is required
    page.click("#errors button")
    page.locator("#rows tr:nth-child(32) input").check()
    page.fill("#reason", "tool change T-07")
    page.click("#restart-btn")
    expect(page.locator("#restart-list")).to_contain_text("31")
    expect(page.locator("#data-counts")).to_contain_text("1 restart")
    expect(page.locator("#rows tr.restart-row")).to_have_count(1)
    expect(page.locator("#rows tr.restart-row .reason")).to_have_text("tool change T-07")

    page.click("#to-analysis")
    page.fill("#a-lsl", "9")
    page.fill("#a-usl", "12")
    page.fill("#a-moving", "3")
    page.click("#run-btn")
    expect(page.locator("#r-moving")).to_contain_text("moving sample size 3 and 1 restart")
    expect(page.locator("#chart-loc svg line.restart")).to_have_count(1)
    expect(page.locator("#chart-var svg line.restart")).to_have_count(1)
    assert page.locator("#chart-loc svg polyline.limit").count() == 2  # the limits are staircases
    expect(page.locator("#r-var-alarms li")).to_have_count(0)  # no alarm from the jump itself

    page.click("nav.tabs button[data-tab=data]")
    page.locator("#rows tr:nth-child(32) input").check()
    page.fill("#reason", "entered by mistake")
    page.click("#unrestart-btn")
    expect(page.locator("#restart-list")).to_have_text("No restarts.")
    assert problems == [], problems
    ctx.close()


def test_new_limits_for_a_phase_in_the_browser(server, browser, tmp_path):
    expect = playwright_sync.expect
    rng = np.random.default_rng(4)
    values = np.r_[rng.normal(10, 0.1, 40), rng.normal(10.6, 0.15, 40)]
    path = tmp_path / "phase.csv"
    path.write_text("v\n" + "\n".join(f"{v:.4f}" for v in values) + "\n")
    ctx = browser.new_context(viewport={"width": 1200, "height": 900}, locale="en")
    page = ctx.new_page()
    problems = []
    page.on("pageerror", lambda e: problems.append(str(e)))
    page.goto(server)
    sign_in(page)
    page.set_input_files("#file", str(path))
    page.select_option("#col-value", "v")
    page.click("#import-btn")
    page.fill("#reason", "new fixture F-2")
    page.locator("#rows tr:nth-child(42) input").check()  # No. 41: the first value of the second phase
    page.check("#restart-phase")
    page.click("#restart-btn")
    expect(page.locator("#restart-list")).to_contain_text("41 (new limits)")
    expect(page.locator("#rows tr.restart-row .status")).to_contain_text("new limits from here")

    page.click("#to-analysis")
    page.fill("#a-lsl", "9")
    page.fill("#a-usl", "12")
    page.click("#run-btn")
    expect(page.locator("#r-phases li")).to_have_count(2)
    expect(page.locator("#r-phases li").nth(1)).to_contain_text("Phase 2, from file row 42: 40 values")
    expect(page.locator("#chart-loc svg line.restart.phase")).to_have_count(1)
    expect(page.locator("#chart-loc svg polyline.center")).to_have_count(1)  # one centre line per phase, drawn as a staircase
    expect(page.locator("#r-loc-alarms li")).to_have_count(0)  # the new level raises no alarm against its own limits
    assert problems == [], problems
    ctx.close()


def test_excel_downloads_in_the_browser(server, browser, tmp_path):
    from openpyxl import load_workbook

    expect = playwright_sync.expect
    ctx = browser.new_context(viewport={"width": 1200, "height": 900}, locale="en")
    page = ctx.new_page()
    page.goto(server)
    sign_in(page)
    page.set_input_files("#file", str(sample(tmp_path)))
    page.select_option("#col-value", "直徑")
    page.select_option("#col-subgroup", "批號")
    page.click("#import-btn")

    with page.expect_download() as info:
        page.click("#export-xlsx-link")
    info.value.save_as(tmp_path / "data.xlsx")
    wb = load_workbook(tmp_path / "data.xlsx")
    assert wb.sheetnames == ["Data", "Log"] and wb["Data"].max_row == 126

    page.click("#suspect-btn")
    page.click("#suspect-select")
    page.fill("#reason", "typing error")
    page.click("#mark-btn")
    page.click("#to-analysis")
    page.fill("#a-lsl", "9.5")
    page.fill("#a-usl", "10.5")
    page.select_option("#a-model", "A1")
    page.select_option("#a-class", "major")
    page.click("#run-btn")
    page.fill("#rp-process", "turning")
    page.click("#rp-create")
    expect(page.locator("#rp-created")).to_contain_text("was created")
    with page.expect_download() as info:
        page.click("#rp-excel")
    info.value.save_as(tmp_path / "report.xlsx")
    report = load_workbook(tmp_path / "report.xlsx")
    assert report.sheetnames[0] == "Summary" and report["Summary"]["A9"].value == "Cpk.G"
    assert any(c.value == "turning" for row in report["Report elements"].iter_rows() for c in row)

    page.click("nav.tabs button[data-tab=saved]")
    expect(page.locator("#saved-reports a", has_text="Excel").first).to_be_visible()  # earlier tests made reports too
    ctx.close()


def test_customer_profile_from_editor_to_report(server, browser, tmp_path):
    expect = playwright_sync.expect
    admin_ctx = browser.new_context(viewport={"width": 1200, "height": 1000}, locale="en")
    admin = admin_ctx.new_page()
    problems = []
    admin.on("pageerror", lambda e: problems.append(str(e)))
    admin.goto(server)
    sign_in(admin, "admin")
    admin.click("nav.tabs button[data-tab=admin]")
    admin.click("#pf-new")
    admin.fill("#pf-name", "Globex")
    admin.fill("#pf-org", "Globex QA")
    admin.fill("#pf-form", "G-12")
    admin.fill("#pf-rev", "B")
    admin.fill("#pf-alpha", "0.01")
    admin.fill("#pf-t-production-major-p", "2.0")
    admin.click("#pf-save")  # a cell with only one of the two numbers is refused
    expect(admin.locator("#errors")).to_contain_text("both targets")
    admin.click("#errors button")
    admin.fill("#pf-t-production-major-pk", "1.8")
    admin.check("#pf-req-part_number")
    admin.click("#pf-extra-add")
    row = admin.locator("#pf-extra .row").first
    row.locator("input[type=text]").nth(0).fill("drawing_no")
    row.locator("input[type=text]").nth(1).fill("Drawing number")
    row.locator("input[type=text]").nth(2).fill("圖面編號")
    row.locator("input[type=checkbox]").check()
    admin.uncheck("#pf-el22")
    admin.click("#pf-save")
    expect(admin.locator("#pf-msg")).to_contain_text("saved, revision 1")
    expect(admin.locator("#profile-list")).to_contain_text("Globex")
    admin.fill("#pf-rev", "C")  # a change raises the revision
    admin.click("#pf-save")
    expect(admin.locator("#pf-msg")).to_contain_text("revision 2")
    admin_ctx.close()

    ctx = browser.new_context(viewport={"width": 1200, "height": 1000}, locale="en")
    page = ctx.new_page()
    page.on("pageerror", lambda e: problems.append(str(e)))
    page.goto(server)
    sign_in(page)
    page.set_input_files("#file", str(sample(tmp_path)))
    page.select_option("#col-value", "直徑")
    page.select_option("#col-subgroup", "批號")
    page.click("#import-btn")
    page.click("#suspect-btn")
    page.click("#suspect-select")
    page.fill("#reason", "typing error")
    page.click("#mark-btn")
    page.click("#to-analysis")
    page.select_option("#a-profile", label="Globex")
    expect(page.locator("#a-alpha")).to_be_disabled()  # fixed by the profile
    expect(page.locator("#a-customer")).to_have_value("Globex")
    expect(page.locator("#a-profile-line")).to_contain_text("risk α")
    expect(page.locator("#a-profile-line")).to_contain_text("target values")
    page.fill("#a-lsl", "9.5")
    page.fill("#a-usl", "10.5")
    page.select_option("#a-model", "A1")
    page.select_option("#a-class", "major")
    page.click("#run-btn")
    expect(page.locator("#r-profile")).to_contain_text("Globex, revision 2")
    # the customer's 2.0 / 1.8, raised a little for 120 values instead of 125; the draft's 1.33 would not show
    expect(page.locator("#r-targets")).to_contain_text("2.01")
    expect(page.locator("#r-targets")).to_contain_text("1.81")
    expect(page.locator("#r-targets")).not_to_contain_text("1.33")

    expect(page.locator("#rp-extra")).to_be_visible()
    expect(page.locator("label.required", has_text="Part no.")).to_have_count(1)
    expect(page.locator("label.required", has_text="Drawing number")).to_have_count(1)
    page.click("#rp-create")
    expect(page.locator("#errors")).to_contain_text("requires these fields")
    expect(page.locator("#errors")).to_contain_text("Drawing number")
    page.click("#errors button")
    page.fill("#rp-part_number", "P-7")
    page.fill("#rp-x-drawing_no", "D-100")
    page.select_option("#rp-language", "en")
    page.click("#rp-create")
    expect(page.locator("#rp-created")).to_contain_text("was created")
    report = ctx.new_page()
    report.goto(server.rstrip("/") + page.locator("#rp-open").get_attribute("href"))
    content = report.content()
    assert "Globex QA" in content and "G-12" in content and "Rev. C" in content and "D-100" in content
    expect(report.locator("h1")).to_have_text("Process study report")
    assert report.locator("section.el h2 .no", has_text="22").count() == 0  # element 22 is switched off in the profile
    assert problems == [], problems
    ctx.close()


def test_spc_monitor_at_the_line_from_set_up_to_the_action_plan(server, browser, app):
    import numpy as np

    expect = playwright_sync.expect
    users = {u.username: u for u in app.state.auth.list_users()}
    if "oper" not in users:
        app.state.auth.create_user("oper", PASSWORD, "operator", "Olga Operator")
    problems = []

    # an engineer sets the monitor up
    ctx = browser.new_context(viewport={"width": 1250, "height": 1000}, locale="en")
    eng = ctx.new_page()
    eng.on("pageerror", lambda e: problems.append(str(e)))
    eng.goto(server)
    sign_in(eng)
    eng.click("nav.tabs button[data-tab=monitor]")
    eng.click("#mon-new")
    eng.fill("#me-name", "Shaft diameter L1")
    eng.fill("#me-process", "turning")
    eng.fill("#me-characteristic", "diameter")
    eng.fill("#me-unit", "mm")
    eng.fill("#me-line", "L1")
    eng.fill("#me-mu", "10")
    eng.fill("#me-sigma", "0.1")
    eng.fill("#me-lsl", "9.5")
    eng.fill("#me-usl", "10.5")
    eng.select_option("#me-class", "major")
    eng.select_option("#me-model", "A1")
    eng.locator("#me-ocap tr[data-key=default] input").nth(0).fill("Measure again, then call the shift leader")
    eng.locator("#me-ocap tr[data-key=default] input").nth(1).fill("Shift leader")
    eng.click("#me-save")
    expect(eng.locator("#md-title")).to_have_text("Shaft diameter L1")
    expect(eng.locator("#md-sub")).to_contain_text("limits revision 1")
    eng.locator("#md-limits").locator("xpath=ancestor::details").locator("summary").click()
    expect(eng.locator("#md-limits")).to_contain_text("Limits revision 1")
    expect(eng.locator("#md-limits")).to_contain_text("parameters: mean 10, standard deviation 0.1")
    ctx.close()

    # the operator works at the line
    ctx = browser.new_context(viewport={"width": 1250, "height": 1000}, locale="en")
    op = ctx.new_page()
    op.on("pageerror", lambda e: problems.append(str(e)))
    op.goto(server)
    sign_in(op, "oper")
    expect(op.locator("#mon-new")).to_be_hidden()  # an operator does not set monitors up
    op.click("nav.tabs button[data-tab=monitor]")
    op.click("#mon-list button:has-text('Open')")
    expect(op.locator("#mon-edit")).to_be_hidden()
    expect(op.locator("#md-ack")).to_be_visible()
    expect(op.locator("#md-plan")).to_contain_text("Measure again, then call the shift leader")

    def enter(values):
        for i, v in enumerate(values):
            op.fill(f"#md-v{i}", str(v))
        op.click("#md-submit")

    enter([10.0, 10.05, 9.97, 10.04, 9.98])
    expect(op.locator("#errors")).to_contain_text("Confirm first that you know the action plan")  # not before the confirmation
    op.click("#errors button")
    op.click("#md-ack-btn")
    expect(op.locator("#md-ack")).to_be_hidden()
    enter([10.0, 10.05, 9.97, 10.04, 9.98])
    expect(op.locator("#md-result .ok-box")).to_contain_text("within the limits")
    enter([10.4, 10.5, 10.45, 10.55, 10.5])
    expect(op.locator("#md-result .alarm-box")).to_contain_text("OUT OF CONTROL")
    expect(op.locator("#md-result .alarm-box")).to_contain_text("Shift leader")
    expect(op.locator("#md-incident")).to_be_visible()
    expect(op.locator("#alert-badge")).to_have_text("1")
    expect(op.locator("#mi-line")).to_contain_text("nobody has taken it over yet")

    # the chart shows the fixed control limits and the warning limits, and never the specification
    expect(op.locator("#md-chart-loc svg line.limit")).to_have_count(2)
    expect(op.locator("#md-chart-loc svg line.warning")).to_have_count(2)
    expect(op.locator("#md-chart-loc svg circle.alarm")).to_have_count(1)
    chart_text = op.locator("#md-chart-loc").inner_text() + op.locator("#md-chart-var").inner_text()
    assert "USL" not in chart_text and "LSL" not in chart_text and "10.5" not in op.locator("#md-chart-loc svg text.limit-label").all_text_contents()

    op.click("#mi-close")  # nothing was done yet
    expect(op.locator("#errors")).to_contain_text("Log the corrective action first")
    op.click("#errors button")
    op.click("#mi-ack")
    expect(op.locator("#mi-line")).to_contain_text("taken over by Olga Operator (oper)")
    op.select_option("#mi-step", "adjust_parameters")
    op.fill("#mi-text", "Raised the feed by 2 %")
    op.click("#mi-action")
    expect(op.locator("#mi-events")).to_contain_text("Raised the feed by 2 %")
    op.fill("#mi-close-text", "Control is back")
    op.click("#mi-close")
    expect(op.locator("#errors")).to_contain_text("new sample after the last action")  # proof is needed
    op.click("#errors button")
    enter([9.96, 10.06, 10.0, 10.08, 9.97])
    expect(op.locator("#md-result")).to_contain_text("verification")
    op.click("#mi-close")
    expect(op.locator("#md-incident")).to_be_hidden()
    expect(op.locator("#alert-badge")).to_be_hidden()

    # a sample that was measured wrongly is declared invalid, with a reason
    enter([10.0, 10.05, 9.97, 10.04, 9.98])
    op.once("dialog", lambda d: d.accept("gauge not zeroed"))
    op.locator("#md-points tr").nth(1).locator("button").click()  # the newest sample is on the first row
    expect(op.locator("#md-points tr.invalid-point")).to_have_count(1)
    expect(op.locator("#md-points tr.invalid-point td").nth(6)).to_have_text("invalid")
    ctx.close()

    # the ongoing report needs some more samples
    rng = np.random.default_rng(31)
    operator = app.state.auth.get_user(next(u.id for u in app.state.auth.list_users() if u.username == "oper"))
    mid = app.state.monitors.store.list()[-1]["id"]
    for _ in range(14):
        app.state.monitors.add_point(mid, [float(v) for v in rng.normal(10, 0.1, 5)], "", {}, None, operator)
    ctx = browser.new_context(viewport={"width": 1250, "height": 1000}, locale="en")
    page = ctx.new_page()
    page.on("pageerror", lambda e: problems.append(str(e)))
    page.goto(server)
    sign_in(page)
    page.click("nav.tabs button[data-tab=monitor]")
    page.click("#mon-list button:has-text('Open')")
    page.locator("#md-ongoing-box summary").click()
    page.fill("#og-window", "40")
    page.click("#og-show")
    expect(page.locator("#og-out")).to_contain_text("Quadrant")
    expect(page.locator("#og-out")).to_contain_text("Response of the action plan: 1 incidents")
    page.click("#og-report")
    expect(page.locator("#og-out")).to_contain_text("was created")
    page.select_option("#lang", "zh-TW")
    expect(page.locator("#md-title")).to_have_text("Shaft diameter L1")
    expect(page.locator("nav.tabs button[data-tab=monitor]")).to_contain_text("現場 SPC")
    assert problems == [], problems
    ctx.close()


def test_count_chart_monitor_with_limits_that_follow_the_sample_size(server, browser, app):
    expect = playwright_sync.expect
    problems = []
    ctx = browser.new_context(viewport={"width": 1250, "height": 1000}, locale="en")
    page = ctx.new_page()
    page.on("pageerror", lambda e: problems.append(str(e)))
    page.goto(server)
    sign_in(page)
    page.click("nav.tabs button[data-tab=monitor]")
    page.click("#mon-new")
    page.fill("#me-name", "Scratches P1")
    page.fill("#me-characteristic", "scratched parts")
    page.select_option("#me-kind", "p")
    expect(page.locator("#me-specs-box")).to_be_hidden()  # counts have no specification limits
    expect(page.locator("#me-src-rate")).to_be_visible()
    page.fill("#me-n", "50")
    page.fill("#me-rate", "0.04")
    page.locator("#me-ocap tr[data-key=default] input").nth(0).fill("Stop and call the shift leader")
    page.locator("#me-ocap tr[data-key=default] input").nth(1).fill("Shift leader")
    page.click("#me-save")
    expect(page.locator("#md-title")).to_have_text("Scratches P1")
    expect(page.locator("#md-chart-var")).to_be_hidden()
    expect(page.locator("#md-ongoing-box")).to_be_hidden()
    page.click("#md-ack-btn")
    page.fill("#md-v0", "2")
    page.fill("#md-v1", "60")
    page.click("#md-submit")
    expect(page.locator("#md-result .ok-box, #md-result .warn-box")).to_be_visible()
    page.fill("#md-v0", "30")
    page.fill("#md-v1", "60")
    page.click("#md-submit")
    expect(page.locator("#md-result .alarm-box")).to_contain_text("OUT OF CONTROL")
    expect(page.locator("#md-chart-loc svg circle.alarm")).to_have_count(1)
    page.locator("#md-limits").locator("xpath=ancestor::details").locator("summary").click()
    expect(page.locator("#md-limits")).to_contain_text("follow the size of each sample")
    page.select_option("#lang", "zh-TW")
    expect(page.locator("#md-limits")).to_contain_text("隨每個樣本的大小")
    assert problems == [], problems
    ctx.close()


def test_acceptance_and_pre_control_monitors_in_the_browser(server, browser, app):
    expect = playwright_sync.expect
    problems = []
    ctx = browser.new_context(viewport={"width": 1250, "height": 1000}, locale="en")
    page = ctx.new_page()
    page.on("pageerror", lambda e: problems.append(str(e)))
    page.goto(server)
    sign_in(page)
    page.click("nav.tabs button[data-tab=monitor]")

    page.click("#mon-new")
    page.fill("#me-name", "Boring tool B1")
    page.fill("#me-characteristic", "bore")
    page.select_option("#me-kind", "acc-xbar")
    expect(page.locator("#me-accept-box")).to_be_visible()
    expect(page.locator("#me-warn-label")).to_be_hidden()
    expect(page.locator("#me-mu-label")).to_be_hidden()
    page.fill("#me-lsl", "9")
    page.fill("#me-usl", "11")
    page.fill("#me-sigma", "0.1")
    page.locator("#me-ocap tr[data-key=default] input").nth(0).fill("Change the tool")
    page.locator("#me-ocap tr[data-key=default] input").nth(1).fill("Setter")
    page.click("#me-save")
    expect(page.locator("#md-title")).to_have_text("Boring tool B1")
    page.click("#md-ack-btn")
    for i, v in enumerate([10.3, 10.35, 10.3, 10.4, 10.3]):
        page.fill(f"#md-v{i}", str(v))
    page.click("#md-submit")
    expect(page.locator("#md-result .ok-box")).to_be_visible()
    for i, v in enumerate([10.8, 10.82, 10.8, 10.85, 10.8]):
        page.fill(f"#md-v{i}", str(v))
    page.click("#md-submit")
    expect(page.locator("#md-result .alarm-box")).to_contain_text("OUT OF CONTROL")
    expect(page.locator("#md-chart-var")).to_be_visible()
    page.locator("#md-limits").locator("xpath=ancestor::details").locator("summary").click()
    expect(page.locator("#md-limits")).to_contain_text("Acceptance chart: factor k")
    page.click("#mon-back")

    page.click("#mon-new")
    page.fill("#me-name", "Start-up S1")
    page.fill("#me-characteristic", "bore")
    page.select_option("#me-kind", "pre")
    expect(page.locator("#me-src-tolerance")).to_be_visible()
    page.fill("#me-lsl", "9.5")
    page.fill("#me-usl", "10.5")
    page.locator("#me-ocap tr[data-key=default] input").nth(0).fill("Stop")
    page.locator("#me-ocap tr[data-key=default] input").nth(1).fill("Setter")
    page.click("#me-save")
    expect(page.locator("#md-title")).to_have_text("Start-up S1")
    expect(page.locator("#md-chart-var")).to_be_hidden()
    expect(page.locator("#md-ongoing-box")).to_be_hidden()
    expect(page.locator("#md-pre-note")).to_be_visible()
    page.click("#md-ack-btn")
    for a, b in ([10.0, 10.1], [10.0, 10.05], [10.1, 9.95]):
        page.fill("#md-v0", str(a)); page.fill("#md-v1", str(b)); page.click("#md-submit")
        expect(page.locator("#md-v0")).to_have_value("", timeout=20000)  # the form is cleared when the answer is in: the next values must not be typed before
        expect(page.locator("#md-result .ok-box")).to_be_visible()
    expect(page.locator("#md-qual")).to_contain_text("Released", timeout=20000)
    page.fill("#md-v0", "10.4"); page.fill("#md-v1", "10.45"); page.click("#md-submit")
    expect(page.locator("#md-result .alarm-box")).to_contain_text("yellow zone on the same side")
    expect(page.locator("#md-qual")).to_contain_text("Not yet released")
    page.select_option("#lang", "zh-TW")
    expect(page.locator("#md-qual")).to_contain_text("尚未放行")
    assert problems == [], problems
    ctx.close()


def test_cusum_and_ewma_monitors_in_the_browser(server, browser, app):
    expect = playwright_sync.expect
    problems = []
    ctx = browser.new_context(viewport={"width": 1250, "height": 1000}, locale="en")
    page = ctx.new_page()
    page.on("pageerror", lambda e: problems.append(str(e)))
    page.goto(server)
    sign_in(page)
    page.click("nav.tabs button[data-tab=monitor]")

    def create(name, kind):
        page.click("#mon-new")
        page.fill("#me-name", name)
        page.fill("#me-characteristic", "fill weight")
        page.select_option("#me-kind", kind)
        expect(page.locator("#me-seq-box")).to_be_visible()
        expect(page.locator("#me-warn-label")).to_be_hidden()
        page.fill("#me-mu", "100")
        page.fill("#me-sigma", "1")
        page.locator("#me-ocap tr[data-key=default] input").nth(0).fill("Check the filler")
        page.locator("#me-ocap tr[data-key=default] input").nth(1).fill("Setter")

    create("Filler CUSUM", "cusum")
    expect(page.locator("#me-k-label")).to_be_visible()
    expect(page.locator("#me-lambda-label")).to_be_hidden()
    page.click("#me-save")
    expect(page.locator("#md-title")).to_have_text("Filler CUSUM")
    expect(page.locator("#md-seq-note")).to_be_visible()
    expect(page.locator("#md-chart-var")).to_be_hidden()
    page.click("#md-ack-btn")
    for i in range(10):  # a shift of one standard deviation adds 0.5 to the upper sum each time: the tenth sample passes h = 4.77
        page.fill("#md-v0", "101")
        page.click("#md-submit")
        expect(page.locator("#md-points tr")).to_have_count(i + 2)
    expect(page.locator("#md-result .alarm-box")).to_contain_text("Upward shift")
    expect(page.locator("#md-chart-loc svg polyline.series2")).to_have_count(1)
    page.locator("#md-limits").locator("xpath=ancestor::details").locator("summary").click()
    expect(page.locator("#md-limits")).to_contain_text("decision interval h = 4.77")
    expect(page.locator("#md-limits")).to_contain_text("Average run length by shift")
    page.click("#mon-back")

    create("Filler EWMA", "ewma")
    expect(page.locator("#me-lambda-label")).to_be_visible()
    expect(page.locator("#me-k-label")).to_be_hidden()
    page.fill("#me-lambda", "0.2")
    page.click("#me-save")
    expect(page.locator("#md-title")).to_have_text("Filler EWMA")
    page.click("#md-ack-btn")
    for i in range(4):  # z = 1.5 (1 - 0.8^i) first passes the narrow limits of the start at the fourth sample
        page.fill("#md-v0", "101.5")
        page.click("#md-submit")
        expect(page.locator("#md-points tr")).to_have_count(i + 2)
    expect(page.locator("#md-result .alarm-box")).to_contain_text("Upward shift")
    expect(page.locator("#md-chart-loc svg circle.alarm")).to_have_count(1)
    page.select_option("#lang", "zh-TW")
    expect(page.locator("#md-seq-note")).to_contain_text("有記憶的管制圖")
    assert problems == [], problems
    ctx.close()


def test_short_run_and_multivariate_monitors_in_the_browser(server, browser, app):
    expect = playwright_sync.expect
    problems = []
    ctx = browser.new_context(viewport={"width": 1250, "height": 1000}, locale="en")
    page = ctx.new_page()
    page.on("pageerror", lambda e: problems.append(str(e)))
    page.goto(server)
    sign_in(page)
    page.click("nav.tabs button[data-tab=monitor]")

    def start(name, kind):
        page.click("#mon-new")
        page.fill("#me-name", name)
        page.fill("#me-characteristic", "dimensions")
        page.select_option("#me-kind", kind)
        page.locator("#me-ocap tr[data-key=default] input").nth(0).fill("Check the setup")
        page.locator("#me-ocap tr[data-key=default] input").nth(1).fill("Setter")

    start("Job shop", "zmr")
    expect(page.locator("#me-specs-box")).to_be_hidden()
    expect(page.locator("#me-src-parts")).to_be_visible()
    page.fill("#me-parts", "A-100; 10.0; 0.1\nB-200; 55.0; 0.5")
    page.click("#me-save")
    expect(page.locator("#md-title")).to_have_text("Job shop")
    expect(page.locator("#md-zmr-note")).to_be_visible()
    expect(page.locator("#md-ongoing-box")).to_be_hidden()
    page.click("#md-ack-btn")
    page.select_option("#md-part", "A-100")
    page.fill("#md-v0", "10.05")
    page.click("#md-submit")
    expect(page.locator("#md-result .ok-box")).to_be_visible()
    page.select_option("#md-part", "B-200")
    page.fill("#md-v0", "57")
    page.click("#md-submit")
    expect(page.locator("#md-result .alarm-box")).to_contain_text("OUT OF CONTROL")
    expect(page.locator("#md-chart-var")).to_be_visible()
    page.locator("#md-limits").locator("xpath=ancestor::details").locator("summary").click()
    expect(page.locator("#md-limits")).to_contain_text("B-200")
    page.click("#mon-back")

    start("Housing", "t2")
    expect(page.locator("#me-mv-box")).to_be_visible()
    expect(page.locator("#me-warn-label")).to_be_hidden()
    page.fill("#me-mv-names", "length, width")
    page.fill("#me-mv-mu", "10, 5")
    page.fill("#me-mv-cov", "0.04 0.018\n0.018 0.0225")
    page.click("#me-save")
    expect(page.locator("#md-title")).to_have_text("Housing")
    expect(page.locator("#md-mv-note")).to_be_visible()
    expect(page.locator("#md-chart-var")).to_be_hidden()
    expect(page.locator("#md-values input")).to_have_count(2)
    page.click("#md-ack-btn")
    page.fill("#md-v0", "10.02"); page.fill("#md-v1", "5.01"); page.click("#md-submit")
    expect(page.locator("#md-result .ok-box")).to_be_visible()
    page.fill("#md-v0", "10.4"); page.fill("#md-v1", "4.7"); page.click("#md-submit")  # each value alone within 3 sigma
    expect(page.locator("#md-result .alarm-box")).to_contain_text("Which characteristics carry the signal")
    expect(page.locator("#md-result .alarm-box")).to_contain_text("length")
    page.select_option("#lang", "zh-TW")
    expect(page.locator("#md-mv-note")).to_contain_text("多個特性")
    assert problems == [], problems
    ctx.close()


def test_extended_limits_and_pearson_monitors_in_the_browser(server, browser, app):
    expect = playwright_sync.expect
    problems = []
    ctx = browser.new_context(viewport={"width": 1250, "height": 1000}, locale="en")
    page = ctx.new_page()
    page.on("pageerror", lambda e: problems.append(str(e)))
    page.goto(server)
    sign_in(page)
    page.click("nav.tabs button[data-tab=monitor]")

    def start(name, kind):
        page.click("#mon-new")
        page.fill("#me-name", name)
        page.fill("#me-characteristic", "diameter")
        page.select_option("#me-kind", kind)
        page.locator("#me-ocap tr[data-key=default] input").nth(0).fill("Check the tool")
        page.locator("#me-ocap tr[data-key=default] input").nth(1).fill("Setter")

    start("Tool wear", "ext-xbar")
    expect(page.locator("#me-warn-label")).to_be_hidden()
    expect(page.locator("#me-sigma-out")).to_be_visible()
    expect(page.locator("#me-skew")).to_be_hidden()
    assert page.input_value("#me-n") == "5"
    page.fill("#me-mu", "10")
    page.fill("#me-sigma", "0.1")
    page.fill("#me-sigma-out", "0.06")
    page.click("#me-save")
    expect(page.locator("#md-title")).to_have_text("Tool wear")
    page.click("#md-ack-btn")
    for i, v in enumerate([10.15, 10.2, 10.17, 10.18, 10.2]):  # beyond a plain 3 sigma limit, inside the extended one
        page.fill(f"#md-v{i}", str(v))
    page.click("#md-submit")
    expect(page.locator("#md-result .ok-box")).to_be_visible()
    expect(page.locator("#md-chart-var")).to_be_visible()
    page.locator("#md-limits").locator("xpath=ancestor::details").locator("summary").click()
    expect(page.locator("#md-limits")).to_contain_text("Extended limits (Analysis of variance")
    page.click("#mon-back")

    start("Runout", "pearson")
    assert page.input_value("#me-n") == "1"
    expect(page.locator("#me-skew")).to_be_visible()
    expect(page.locator("#me-sigma-out")).to_be_hidden()
    page.fill("#me-mu", "4")
    page.fill("#me-sigma", "2")
    page.fill("#me-skew", "1")
    page.fill("#me-kurt", "4.5")
    page.click("#me-save")
    expect(page.locator("#md-title")).to_have_text("Runout")
    page.click("#md-ack-btn")
    page.fill("#md-v0", "9.5")  # 2.75 sigma above: inside the Pearson limit
    page.click("#md-submit")
    expect(page.locator("#md-result .ok-box")).to_be_visible()
    page.fill("#md-v0", "13")
    page.click("#md-submit")
    expect(page.locator("#md-result .alarm-box")).to_contain_text("OUT OF CONTROL")
    page.locator("#md-limits").locator("xpath=ancestor::details").locator("summary").click()
    expect(page.locator("#md-limits")).to_contain_text("curve III")
    page.select_option("#lang", "zh-TW")
    expect(page.locator("#md-limits")).to_contain_text("偏度")
    assert problems == [], problems
    ctx.close()


def test_autocorrelated_and_multistream_monitors_in_the_browser(server, browser, app):
    import numpy as np

    expect = playwright_sync.expect
    problems = []
    ctx = browser.new_context(viewport={"width": 1250, "height": 1000}, locale="en")
    page = ctx.new_page()
    page.on("pageerror", lambda e: problems.append(str(e)))
    page.goto(server)
    sign_in(page)
    page.click("nav.tabs button[data-tab=monitor]")

    def start(name, kind):
        page.click("#mon-new")
        page.fill("#me-name", name)
        page.fill("#me-characteristic", "temperature")
        page.select_option("#me-kind", kind)
        page.locator("#me-ocap tr[data-key=default] input").nth(0).fill("Check the controller")
        page.locator("#me-ocap tr[data-key=default] input").nth(1).fill("Setter")

    start("Furnace", "ar")
    expect(page.locator("#me-specs-box")).to_be_hidden()
    expect(page.locator("#me-phi")).to_be_visible()
    assert page.input_value("#me-n") == "1"
    page.fill("#me-mu", "100")
    page.fill("#me-sigma", "0.5")
    page.fill("#me-phi", "0.8")
    page.click("#me-save")
    expect(page.locator("#md-title")).to_have_text("Furnace")
    expect(page.locator("#md-dep-note")).to_be_visible()
    expect(page.locator("#md-ongoing-box")).to_be_hidden()
    page.click("#md-ack-btn")
    page.fill("#md-v0", "100.2")
    page.click("#md-submit")
    expect(page.locator("#md-result .ok-box")).to_be_visible()
    page.fill("#md-v0", "103")  # far from the forecast 100.16
    page.click("#md-submit")
    expect(page.locator("#md-result .alarm-box")).to_contain_text("OUT OF CONTROL")
    page.locator("#md-limits").locator("xpath=ancestor::details").locator("summary").click()
    expect(page.locator("#md-limits")).to_contain_text("Autoregressive model of order 1")
    page.click("#mon-back")

    rng = np.random.default_rng(5)
    rows = 10.0 + np.array([0.3, -0.1, 0.0, -0.2]) + rng.normal(0, 0.05, (30, 1)) + rng.normal(0, 0.1, (30, 4))
    start("Four spindles", "multistream")
    assert page.input_value("#me-n") == "4"
    page.select_option("#me-src-type", "observations")
    expect(page.locator("#me-src-rows")).to_be_visible()
    page.fill("#me-ms-names", "S1, S2, S3, S4")
    page.fill("#me-mv-rows", "\n".join(" ".join(f"{v:.4f}" for v in r) for r in rows))
    page.click("#me-save")
    expect(page.locator("#md-title")).to_have_text("Four spindles")
    expect(page.locator("#md-values input")).to_have_count(4)
    expect(page.locator("#md-chart-var")).to_be_visible()
    page.click("#md-ack-btn")
    for i, v in enumerate([10.3, 9.9, 10.0, 9.8]):
        page.fill(f"#md-v{i}", str(v))
    page.click("#md-submit")
    expect(page.locator("#md-result .ok-box")).to_be_visible()
    for i, v in enumerate([10.3, 10.8, 10.0, 9.8]):  # spindle 2 wanders off
        page.fill(f"#md-v{i}", str(v))
    page.click("#md-submit")
    expect(page.locator("#md-result .alarm-box")).to_contain_text("S2")
    page.select_option("#lang", "zh-TW")
    expect(page.locator("#md-dep-note")).to_contain_text("多個串流")
    assert problems == [], problems
    ctx.close()


def test_trend_monitor_with_a_cycle_marked_by_the_operator_in_the_browser(server, browser, app):
    expect = playwright_sync.expect
    problems = []
    ctx = browser.new_context(viewport={"width": 1250, "height": 1000}, locale="en")
    page = ctx.new_page()
    page.on("pageerror", lambda e: problems.append(str(e)))
    page.goto(server)
    sign_in(page)
    page.click("nav.tabs button[data-tab=monitor]")
    page.click("#mon-new")
    page.fill("#me-name", "Honing")
    page.fill("#me-characteristic", "bore diameter")
    page.select_option("#me-kind", "trend")
    expect(page.locator("#me-specs-box")).to_be_hidden()
    expect(page.locator("#me-intercept")).to_be_visible()
    expect(page.locator("#me-mu-label")).to_be_hidden()
    assert page.input_value("#me-n") == "1"
    page.fill("#me-intercept", "10")
    page.fill("#me-slope", "-0.01")
    page.fill("#me-sigma", "0.01")
    page.locator("#me-ocap tr[data-key=default] input").nth(0).fill("Measure again")
    page.locator("#me-ocap tr[data-key=default] input").nth(1).fill("Setter")
    page.click("#me-save")
    expect(page.locator("#md-title")).to_have_text("Honing")
    expect(page.locator("#md-cycle-info")).to_contain_text("first sample starts cycle 1")
    page.click("#md-ack-btn")
    for i, v in enumerate([10.004, 9.986, 9.983, 9.968]):  # near the line 10 - 0.01 x position (not exactly on it: two equal residuals would be a moving range of 0)
        page.fill("#md-v0", str(v))
        page.click("#md-submit")
        expect(page.locator("#md-result .ok-box")).to_contain_text(f"Sample {i + 1}: within the limits")  # wait for the answer to this sample, not the one before
        expect(page.locator("#md-result .ok-box")).to_contain_text(f"position {i}")
    expect(page.locator("#md-cycle-info")).to_contain_text("next sample is number 5 of the cycle")
    # the tool is changed: without saying so, the value 10.00 is far above the line (9.96 expected)
    page.fill("#md-v0", "10.0")
    page.click("#md-submit")
    expect(page.locator("#md-result .alarm-box")).to_contain_text("OUT OF CONTROL")
    expect(page.locator("#md-result .alarm-box")).to_contain_text("Sample 5")
    # a new cycle needs a note
    page.check("#md-cycle")
    expect(page.locator("#md-cycle-note")).to_be_visible()
    page.fill("#md-v0", "10.0")
    page.click("#md-submit")
    expect(page.locator("#errors")).to_contain_text("note of what happened")
    page.click("#errors button")
    page.fill("#md-cycle-note", "Stone 8 fitted")
    page.click("#md-submit")
    expect(page.locator("#md-result")).to_contain_text("Sample 6", timeout=20000)
    expect(page.locator("#md-result")).to_contain_text("Cycle 2, position 0")
    expect(page.locator("#md-cycle")).not_to_be_checked()
    expect(page.locator("#md-cycle-info")).to_contain_text("Cycle 2: the next sample is number 2 of the cycle")
    expect(page.locator("#md-points")).to_contain_text("↻ Stone 8 fitted")
    page.locator("#md-limits").locator("xpath=ancestor::details").locator("summary").click()
    expect(page.locator("#md-limits")).to_contain_text("Line of the cycle")
    page.select_option("#lang", "zh-TW")
    expect(page.locator("#md-dep-note")).to_contain_text("週期由人宣告")
    assert problems == [], problems
    ctx.close()


def test_mcusum_monitor_in_the_browser(server, browser, app):
    expect = playwright_sync.expect
    problems = []
    ctx = browser.new_context(viewport={"width": 1250, "height": 1000}, locale="en")
    page = ctx.new_page()
    page.on("pageerror", lambda e: problems.append(str(e)))
    page.goto(server)
    sign_in(page)
    page.click("nav.tabs button[data-tab=monitor]")
    page.click("#mon-new")
    page.fill("#me-name", "Pair MCUSUM")
    page.fill("#me-characteristic", "pair")
    page.select_option("#me-kind", "mcusum")
    expect(page.locator("#me-k-label")).to_be_visible()
    expect(page.locator("#me-fir-label")).to_be_hidden()
    expect(page.locator("#me-lambda-label")).to_be_hidden()
    expect(page.locator("#me-mv-box")).to_be_visible()
    page.fill("#me-mv-names", "length, width")
    page.fill("#me-mv-mu", "10, 5")
    page.fill("#me-mv-cov", "0.04 0.012\n0.012 0.0225")
    page.locator("#me-ocap tr[data-key=default] input").nth(0).fill("Check the setup")
    page.locator("#me-ocap tr[data-key=default] input").nth(1).fill("Setter")
    page.click("#me-save")
    expect(page.locator("#md-title")).to_have_text("Pair MCUSUM")
    expect(page.locator("#md-values input")).to_have_count(2)
    page.click("#md-ack-btn")
    for i in range(12):
        page.fill("#md-v0", "10.25")
        page.fill("#md-v1", "5.15")
        page.click("#md-submit")
        expect(page.locator("#md-points tr")).to_have_count(i + 2)
        if page.locator("#md-result .alarm-box").count():
            break
    expect(page.locator("#md-result .alarm-box")).to_contain_text("OUT OF CONTROL")
    page.locator("#md-limits").locator("xpath=ancestor::details").locator("summary").click()
    expect(page.locator("#md-limits")).to_contain_text("MCUSUM (Crosier)")
    page.select_option("#lang", "zh-TW")
    expect(page.locator("#md-limits")).to_contain_text("決策界限")
    assert problems == [], problems
    ctx.close()


def proven_system(app, user, name, tolerance=1.0):
    """A measurement system with a gauge R&R, a type 1 study and a stability check: its gate is open."""
    import datetime

    rng = np.random.default_rng(2)
    svc = app.state.msa
    sid = svc.create({"name": name, "characteristic": "bore", "unit": "mm", "resolution": 0.001, "tolerance": tolerance}, user)["system"]["id"]
    day = datetime.date.today().isoformat()
    data = (10 + rng.normal(0, 0.2, (10, 1, 1)) + rng.normal(0, 0.003, (10, 3, 3))).tolist()
    svc.add_study(sid, "grr", day, "", {"data": data}, user)
    svc.add_study(sid, "stability", day, "", {"values": [float(v) for v in 10 + rng.normal(0, 0.005, 25)]}, user)
    svc.add_study(sid, "type1", day, "", {"reference": 10.0, "values": [float(v) for v in 10 + rng.normal(0, 0.004, 50)]}, user)
    return sid


def test_machine_study_checklist_in_the_browser(server, browser, app):
    import numpy as np

    from spc.data import Dataset
    from spc.study import checklist as cl

    expect = playwright_sync.expect
    problems = []
    rng = np.random.default_rng(8)
    times = [f"2026-03-02T08:{i // 60:02d}:{i % 60:02d}" for i in range(50)]
    admin = next(u for u in app.state.auth.list_users() if u.username == "admin")
    app.state.store.add(Dataset.from_values([float(v) for v in rng.normal(10.0, 0.02, 50)], timestamp=times), admin.id, "Grinder run 1")
    proven_system(app, admin, "Study gauge")
    ctx = browser.new_context(viewport={"width": 1250, "height": 1000}, locale="en")
    page = ctx.new_page()
    page.on("pageerror", lambda e: problems.append(str(e)))
    page.goto(server)
    sign_in(page)
    page.click("nav.tabs button[data-tab=study]")
    page.click("#st-new")
    page.fill("#se-name", "Grinder 4 bore")
    page.fill("#se-machine", "Grinder 4")
    page.fill("#se-characteristic", "bore diameter")
    page.fill("#se-station", "spindle 1")
    page.select_option("#se-dataset", label="Grinder run 1 (50)")
    page.select_option("#se-msa", label="Study gauge (gate open)")
    page.fill("#se-lsl", "9.9")
    page.fill("#se-usl", "10.1")
    page.fill("#se-pre-one", "10.01")
    page.click("#se-save")
    expect(page.locator("#st-title")).to_have_text("Grinder 4 bore")
    expect(page.locator("#st-verdict")).to_contain_text("open or failed")
    expect(page.locator("#st-close")).to_be_disabled()
    rows = page.locator("#st-items tr")
    expect(rows).to_have_count(len(cl.ITEMS) + 1)
    expect(page.locator("#st-items tr", has_text="Number of parts")).to_contain_text("50 parts (at least 50)")
    expect(page.locator("#st-items tr", has_text="Pre-production run, 1 part")).to_contain_text("lies in the area")
    first = page.locator("#st-items tr", has_text="HuMan")
    first.locator("select").select_option("not_applicable")
    first.locator("button").click()
    expect(page.locator("#errors")).to_contain_text("A note is required")
    page.click("#errors button")
    first.locator("input").fill("automatic line, no operator")
    first.locator("select").select_option("not_applicable")
    first.locator("button").click()
    expect(page.locator("#st-items tr", has_text="HuMan")).to_contain_text("does not apply")
    # the rest is confirmed through the service, the browser then closes the study
    study_id = app.state.studies.list()[0]["id"]
    user = app.state.auth.get_user(next(u.id for u in app.state.auth.list_users() if u.username == "admin"))
    for item in cl.ITEMS:
        if not item.auto and item.key != "human":
            app.state.studies.set_item(study_id, item.key, "ok", "", user)
    page.click("#st-back")
    page.click("#st-list button:has-text('Open')")
    expect(page.locator("#st-verdict")).to_contain_text("can be closed")
    page.fill("#st-close-reason", "all conditions met")
    page.click("#st-close")
    expect(page.locator("#st-closed-line")).to_contain_text("all conditions met")
    expect(page.locator("#st-items select").first).to_be_disabled()
    page.select_option("#lang", "zh-TW")
    expect(page.locator("#st-items tr", has_text="零件數")).to_contain_text("50 件")
    page.click("#st-back")
    expect(page.locator("#st-list")).to_contain_text("已於")
    assert problems == [], problems
    ctx.close()


def test_time_model_suggestion_in_the_browser(server, browser, tmp_path):
    expect = playwright_sync.expect
    rng = np.random.default_rng(4)
    rows = ["lot,v"]
    for i in range(40):  # tool wear: the location follows a trend, the variation is constant
        rows += [f"L{i + 1},{10 + 0.2 * i + rng.normal(0, 1):.4f}" for _ in range(5)]
    path = tmp_path / "wear.csv"
    path.write_text("\n".join(rows) + "\n")
    ctx = browser.new_context(viewport={"width": 1200, "height": 1100}, locale="en")
    page = ctx.new_page()
    problems = []
    page.on("pageerror", lambda e: problems.append(str(e)))
    page.goto(server)
    sign_in(page)
    page.set_input_files("#file", str(path))
    page.select_option("#col-value", "v")
    page.select_option("#col-subgroup", "lot")
    page.click("#import-btn")
    page.click("#to-analysis")
    assert page.input_value("#a-model") == ""
    page.locator("#a-model-assist summary").click()
    page.locator("input[data-hint=multi_stream]").check()
    page.click("#a-model-suggest")
    box = page.locator("#a-model-suggestion")
    expect(box).to_contain_text("Suggestion: C3")
    expect(box).to_contain_text("the location follows a trend")
    expect(box).to_contain_text("Based on 40 subgroups of 5 values")
    expect(box).to_contain_text("points to D")  # the hint that does not fit is said
    expect(box).to_contain_text("not in statistical control")
    page.click("#a-model-suggestion button")
    assert page.input_value("#a-model") == "C3"
    page.select_option("#lang", "zh-TW")
    expect(box).to_contain_text("建議：")
    expect(box).to_contain_text("位置隨趨勢改變")
    assert problems == [], problems
    ctx.close()


def test_measurement_system_and_the_gate_in_the_browser(server, browser, app):
    expect = playwright_sync.expect
    problems = []
    rng = np.random.default_rng(12)
    ctx = browser.new_context(viewport={"width": 1250, "height": 1100}, locale="en")
    page = ctx.new_page()
    page.on("pageerror", lambda e: problems.append(str(e)))
    page.goto(server)
    sign_in(page)
    page.click("nav.tabs button[data-tab=msa]")
    page.click("#ms-new")
    page.fill("#mse-name", "Air gauge 7")
    page.fill("#mse-characteristic", "bore diameter")
    page.fill("#mse-unit", "mm")
    page.fill("#mse-resolution", "0.001")
    page.fill("#mse-tolerance", "1")
    page.locator("#ms-editor summary").click()
    assert page.input_value("#msp-validity_months") == "12" and page.input_value("#msp-guard_band_risk") == "5"
    page.click("#mse-save")
    expect(page.locator("#ms-title")).to_have_text("Air gauge 7")
    expect(page.locator("#ms-status")).to_contain_text("blocked")
    expect(page.locator("#ms-checks tr", has_text="Gauge R&R").first).to_contain_text("no gauge R&R study")

    def add(kind, text, reference=None):
        page.select_option("#mss-kind", kind)
        if reference is not None:
            page.fill("#mss-reference", str(reference))
        page.fill("#mss-data", text)
        page.click("#mss-save")

    grr = (10 + rng.normal(0, 0.2, (10, 1, 1)) + rng.normal(0, 0.003, (10, 3, 3)))
    add("grr", "\n".join(" ; ".join(" ".join(f"{v:.4f}" for v in op) for op in part) for part in grr))
    expect(page.locator("#ms-studies tr", has_text="Gauge R&R")).to_contain_text("capable")
    add("stability", " ".join(f"{v:.4f}" for v in 10 + rng.normal(0, 0.005, 25)))
    expect(page.locator("#ms-status")).to_contain_text("The gate is open")
    expect(page.locator("#ms-uncertainty")).to_contain_text("Expanded uncertainty U =")
    # a bad type 1 study stops the gate; a waiver with a reason lets it pass, flagged
    add("type1", " ".join(f"{v:.4f}" for v in 10 + rng.normal(0, 0.08, 40)), reference=10.0)
    expect(page.locator("#ms-status")).to_contain_text("blocked")
    expect(page.locator("#ms-checks tr", has_text="Type 1 study")).to_contain_text("below the limit")
    page.once("dialog", lambda d: d.accept("customer agreed to a type 1 later"))
    page.locator("#ms-checks tr", has_text="Type 1 study").locator("button").click()
    expect(page.locator("#ms-status")).to_contain_text("The gate is open")
    expect(page.locator("#ms-checks tr", has_text="Type 1 study")).to_contain_text("waived")
    expect(page.locator("#ms-checks tr", has_text="Type 1 study")).to_contain_text("customer agreed")
    # a monitor tied to a blocked system takes no sample
    other = app.state.msa.create({"name": "No proof gauge", "resolution": 0.001, "tolerance": 1.0},
                                 app.state.auth.get_user(next(u.id for u in app.state.auth.list_users() if u.username == "admin")))["system"]["id"]
    page.click("nav.tabs button[data-tab=monitor]")
    page.click("#mon-new")
    page.fill("#me-name", "Bore monitor")
    page.fill("#me-characteristic", "bore")
    page.select_option("#me-msa", label="No proof gauge (gate blocked)")
    page.fill("#me-mu", "10")
    page.fill("#me-sigma", "0.1")
    page.locator("#me-ocap tr[data-key=default] input").nth(0).fill("Check the setup")
    page.locator("#me-ocap tr[data-key=default] input").nth(1).fill("Setter")
    page.click("#me-save")
    expect(page.locator("#md-msa")).to_contain_text("gate blocked")
    page.click("#md-ack-btn")
    for i, v in enumerate([10.0, 10.01, 9.99, 10.02, 10.0]):
        page.fill(f"#md-v{i}", str(v))
    page.click("#md-submit")
    expect(page.locator("#errors")).to_contain_text("is not proven")
    page.select_option("#lang", "zh-TW")
    expect(page.locator("#md-msa")).to_contain_text("閘門封鎖")
    assert problems == [], problems
    ctx.close()


def test_control_plan_and_roles_in_the_browser(server, browser):
    expect = playwright_sync.expect
    problems = []
    ctx = browser.new_context(viewport={"width": 1250, "height": 1100}, locale="en")
    page = ctx.new_page()
    page.on("pageerror", lambda e: problems.append(str(e)))
    page.goto(server)
    sign_in(page)
    page.click("nav.tabs button[data-tab=roles]")
    matrix = page.locator("#rl-matrix")
    expect(matrix.locator("tr", has_text="Statistics")).to_be_visible()
    expect(matrix.locator("tr").first).to_contain_text("Line operator")
    page.locator("#rl-people tr", has_text="Eva Engineer").locator("button").click()
    page.locator("#rl-person-roles label", has_text="Quality planning").locator("input").check()
    page.click("#rl-person-save-roles")
    expect(page.locator("#rl-people tr", has_text="Eva Engineer")).to_contain_text("Quality planning ⚠")  # not qualified yet
    answers = ["2", "trained and tested in 2026"]  # the prompts of the page are answered in this order
    page.on("dialog", lambda d: d.accept(answers.pop(0)))
    page.locator("#rl-comp tr", has_text="Quality management system").locator("button").click()
    expect(page.locator("#rl-comp tr", has_text="Quality management system")).to_contain_text("trained and tested in 2026")

    page.click("nav.tabs button[data-tab=plan]")
    page.click("#pl-new")
    page.fill("#ple-name", "Housing line 1")
    page.fill("#ple-part", "H-100")
    page.click("#ple-add-line")
    page.fill("#pll-step", "Turning")
    page.fill("#pll-characteristic", "bore")
    page.fill("#pll-lsl", "9.9")
    page.fill("#pll-usl", "10.1")
    page.select_option("#pll-control", "other")
    page.fill("#pll-sample_size", "5")
    page.fill("#pll-frequency", "every hour")
    page.fill("#pll-method", "air gauge")
    page.locator("#pll-responsible label", has_text="Line operator").locator("input").check()
    page.click("#pll-save")
    page.click("#ple-save")
    expect(page.locator("#pl-title")).to_have_text("Housing line 1")
    expect(page.locator("#pl-lines tr", has_text="Turning")).to_contain_text("no reaction plan")
    expect(page.locator("#pl-verdict")).to_contain_text("Not releasable")
    expect(page.locator("#pl-release")).to_be_disabled()
    page.click("#pl-edit")
    page.locator("#ple-lines tr", has_text="Turning").locator("button", has_text="Edit").click()
    page.fill("#pll-reaction", "sort the lot and call the supervisor")
    page.click("#pll-save")
    page.click("#ple-save")
    expect(page.locator("#pl-lines tr", has_text="Turning")).not_to_contain_text("no reaction plan")
    # an approval needs the role: the engineer holds only quality planning
    answers.append("ok")
    page.locator("#pl-approvals tr", has_text="Product developer").locator("button").click()
    expect(page.locator("#errors")).to_contain_text("do not hold this SPC role")
    page.click("#errors button")
    answers.append("plan reviewed")
    page.locator("#pl-approvals tr", has_text="Quality planning").locator("button").click()
    expect(page.locator("#pl-approvals tr", has_text="Quality planning")).to_contain_text("plan reviewed")
    expect(page.locator("#pl-verdict")).to_contain_text("Not releasable")  # three approvals are missing
    page.select_option("#lang", "zh-TW")
    expect(page.locator("#pl-verdict")).to_contain_text("無法發行")
    expect(page.locator("#pl-lines")).to_contain_text("沒有合格人員")
    assert problems == [], problems
    ctx.close()


def test_validation_run_and_reference_case_in_the_browser(server, browser):
    expect = playwright_sync.expect
    problems = []
    ctx = browser.new_context(viewport={"width": 1250, "height": 1100}, locale="en")
    page = ctx.new_page()
    page.on("pageerror", lambda e: problems.append(str(e)))
    page.set_default_timeout(40000)  # a run executes about six hundred checks and takes some seconds
    page.goto(server)
    sign_in(page)
    page.click("nav.tabs button[data-tab=validation]")
    expect(page.locator("#va-runs")).to_contain_text("No runs yet")
    expect(page.locator("#va-iso tr", has_text="Model C2").first).to_contain_text("not run yet")
    page.click("#va-run")
    row = page.locator("#va-runs tr", has_text="Verified, not validated")
    expect(row).to_be_visible(timeout=40000)
    expect(page.locator("#va-iso tr", has_text="Model C2").first).to_contain_text("known differences of the standard", timeout=40000)
    # a reference case: five values with a mean the user knows
    expect(page.locator("#va-iso tr")).to_have_count(12)
    page.click("#va-new")
    page.fill("#vae-name", "Textbook example")
    page.fill("#vae-source", "hand calculation 2026-02")
    page.fill("#vae-values", "9.8 10.0 10.2 9.9 10.1 10.0 9.9 10.1 10.0 10.0")
    page.fill("#vae-lsl", "9.0")
    page.fill("#vae-usl", "11.0")
    page.fill("#vae-expected", "indices.mean 10.0 1e-9\nindices.n 10 0")
    page.click("#vae-save")
    expect(page.locator("#va-cases")).to_contain_text("Textbook example")
    page.click("#va-run")
    expect(page.locator("#va-runs tr").nth(1)).to_contain_text("Passed", timeout=40000)
    page.select_option("#lang", "zh-TW")
    expect(page.locator("#va-runs")).to_contain_text("通過")
    link = page.locator("#va-runs tr").nth(1).locator("a").nth(1)
    href = link.get_attribute("href")
    assert href.endswith("report?lang=zh-TW")
    assert problems == [], problems
    ctx.close()


def test_equipment_link_in_the_browser(server, browser, app):
    expect = playwright_sync.expect
    eng = next(u for u in app.state.auth.list_users() if u.username == "eng")
    config = {"name": "Link chart", "characteristic": "bore", "kind": "xbar-s", "n": 3, "ocap": {"default": {"operator_action": "Measure again"}}}
    mid = app.state.monitors.create(config, {"type": "parameters", "mu": 10.0, "sigma": 0.1}, eng)["id"]
    app.state.monitors.acknowledge_ocap(mid, eng)
    app.state.equipment.prober = lambda link: {"ok": True, "error": None, "nodes": [
        {"node_id": n["node_id"], "ok": True, "value": 10.0, "type": "Double", "quality": "good", "source_time": "2026-10-01T08:00:00Z", "error": None} for n in link["nodes"]]}
    problems = []
    ctx = browser.new_context(viewport={"width": 1250, "height": 1100}, locale="en")
    page = ctx.new_page()
    page.on("pageerror", lambda e: problems.append(str(e)))
    page.goto(server)
    sign_in(page)
    page.click("nav.tabs button[data-tab=equipment]")
    expect(page.locator("#eq-list")).to_contain_text("No links yet")
    page.click("#eq-new")
    page.fill("#eqe-name", "Grinder 4")
    page.fill("#eqe-endpoint", "opc.tcp://127.0.0.1:4840/grinder")
    page.fill("#eqe-nodes", f"ns=2;s=Bore {mid} 0.001 0")
    page.click("#eqe-save")
    expect(page.locator("#eq-title")).to_have_text("Grinder 4")
    expect(page.locator("#eq-test-table")).to_contain_text("not tested")
    page.click("#eq-enable")
    expect(page.locator("#errors")).to_contain_text("Run a test read first")
    page.click("#errors button")
    page.click("#eq-test")
    expect(page.locator("#eq-test-table")).to_contain_text("Tested")
    expect(page.locator("#eq-test-table")).to_contain_text("Double")
    page.click("#eq-enable")
    expect(page.locator("#eq-state")).to_contain_text("client not running")
    expect(page.locator("#eq-enable")).to_have_text("Disable")
    page.select_option("#lang", "zh-TW")
    expect(page.locator("#eq-state")).to_contain_text("用戶端未執行")
    page.click("#eq-back")
    expect(page.locator("#eq-list")).to_contain_text("Grinder 4")
    assert problems == [], problems
    ctx.close()


def test_state_tests_in_the_browser(server, browser, tmp_path):
    expect = playwright_sync.expect
    rng = np.random.default_rng(12)
    rows = ["machine,v"]
    for m, mu, sd in (("M1", 10.0, 0.2), ("M2", 10.6, 0.2), ("M3", 10.0, 0.7)):
        rows += [f"{m},{rng.normal(mu, sd):.4f}" for _ in range(30)]
    path = tmp_path / "states.csv"
    path.write_text("\n".join(rows) + "\n")
    ctx = browser.new_context(viewport={"width": 1200, "height": 1100}, locale="en")
    page = ctx.new_page()
    problems = []
    page.on("pageerror", lambda e: problems.append(str(e)))
    page.goto(server)
    sign_in(page)
    page.set_input_files("#file", str(path))
    page.select_option("#col-value", "v")
    page.select_option("#col-subgroup", "machine")
    page.click("#import-btn")
    page.click("#to-analysis")
    page.click("#a-states-run")
    box = page.locator("#a-states-result")
    expect(box).to_contain_text("Bartlett, equal variances")
    expect(box).to_contain_text("the states differ")
    expect(box.locator("tr", has_text="M2")).to_be_visible()
    # Pm and Pmk of the states, ISO 22514-8
    page.locator("#a-ms summary").click()
    page.click("#ams-run")
    expect(page.locator("#errors")).to_contain_text("Enter the lower and the upper specification limit")
    page.click("#errors button")
    page.fill("#a-lsl", "9")
    page.fill("#a-usl", "12")
    page.click("#ams-run")
    res = page.locator("#ams-result")
    expect(res).to_contain_text("Type of global dispersion")
    expect(res).to_contain_text("Locations, Fisher")
    expect(res.locator("tr", has_text="M3")).to_be_visible()
    expect(res).to_contain_text("Pm =")
    page.select_option("#ams-locations", "true")
    page.select_option("#ams-widths", "true")
    page.click("#ams-run")
    expect(res).to_contain_text("one dispersion for all states")
    page.select_option("#lang", "zh-TW")
    expect(box).to_contain_text("各狀態有差異")
    expect(res).to_contain_text("所有狀態共用一個分散")
    assert problems == [], problems
    ctx.close()


def test_multistate_result_in_the_report_in_the_browser(server, browser, tmp_path):
    expect = playwright_sync.expect
    rng = np.random.default_rng(14)
    rows = ["machine,v"]
    for m, mu, sd in (("M1", 10.0, 0.2), ("M2", 10.6, 0.2), ("M3", 10.2, 0.25)):
        rows += [f"{m},{rng.normal(mu, sd):.3f}" for _ in range(25)]
    path = tmp_path / "states.csv"
    path.write_text("\n".join(rows) + "\n")
    ctx = browser.new_context(viewport={"width": 1200, "height": 1100}, locale="en")
    page = ctx.new_page()
    problems = []
    page.on("pageerror", lambda e: problems.append(str(e)))
    page.goto(server)
    sign_in(page)
    page.set_input_files("#file", str(path))
    page.select_option("#col-value", "v")
    page.select_option("#col-subgroup", "machine")
    page.click("#import-btn")
    page.click("#to-analysis")
    page.fill("#a-lsl", "9")
    page.fill("#a-usl", "12")
    page.locator("#a-ms summary").click()
    page.select_option("#ams-locations", "false")
    page.click("#run-btn")
    page.check("#rp-ms")
    page.click("#rp-create")
    expect(page.locator("#rp-created")).to_contain_text("was created")
    href = page.locator("#rp-open").get_attribute("href")
    report = ctx.new_page()
    report.goto(server.rstrip("/") + href)
    expect(report.locator("h2", has_text="Annex D")).to_be_visible()
    content = report.content()
    assert "Decided by the analyst: the locations of the local dispersions are different" in content and "Pm =" in content and "M2" in content
    assert problems == [], problems
    ctx.close()


def test_external_signature_and_trusted_signer_in_the_browser(server, browser, app):
    pytest.importorskip("cryptography")
    import base64

    from cryptography.hazmat.primitives import serialization
    from cryptography.hazmat.primitives.asymmetric import ed25519

    from spc.signing import core
    from tests.test_api import REPORT_BODY, upload
    from tests.conftest import logged_in_client

    expect = playwright_sync.expect
    api_client = logged_in_client(app, "eng")
    out = api_client.post(f"/api/datasets/{upload(api_client).json()['id']}/reports", json=REPORT_BODY).json()
    private = ed25519.Ed25519PrivateKey.generate()
    pem = private.public_key().public_bytes(serialization.Encoding.PEM, serialization.PublicFormat.SubjectPublicKeyInfo).decode()
    raw, _ = core.sign(private, out["digest"])

    ctx = browser.new_context(viewport={"width": 1300, "height": 900}, locale="en")
    page = ctx.new_page()
    problems = []
    page.on("pageerror", lambda e: problems.append(str(e)))
    page.goto(server)
    sign_in(page, "admin")
    page.click("nav.tabs button[data-tab=admin]")
    page.fill("#sn-name", "QA lead")
    page.fill("#sn-key", pem)
    page.click("#sn-add")
    expect(page.locator("#signer-list")).to_contain_text("QA lead", timeout=20000)
    page.click("nav.tabs button[data-tab=saved]")
    page.locator(f"#saved-reports button[data-report='{out['id'] if 'id' in out else out['report_id']}']").click()
    expect(page.locator("#sig-message")).to_contain_text("spc-archive-v1", timeout=20000)
    page.fill("#sig-value", "AAAA")
    page.fill("#sig-key", pem)
    page.click("#sig-submit")
    expect(page.locator("#errors")).to_contain_text("The signature is neither", timeout=20000)
    page.fill("#sig-value", base64.b64encode(raw).decode())
    page.fill("#sig-note", "released")
    page.click("#sig-submit")
    expect(page.locator("#sig-list")).to_contain_text("Valid, trusted", timeout=20000)
    expect(page.locator("#sig-list")).to_contain_text("QA lead")
    expect(page.locator("#saved-reports")).to_contain_text("1")
    assert problems == []
    ctx.close()


def test_special_cases_multistage_and_gdt_in_the_browser(server, browser, tmp_path):
    expect = playwright_sync.expect
    ctx = browser.new_context(viewport={"width": 1300, "height": 1000}, locale="en")
    page = ctx.new_page()
    problems = []
    page.on("pageerror", lambda e: problems.append(str(e)))
    page.goto(server)
    sign_in(page)
    page.click("nav.tabs button[data-tab=special]")
    expect(page.locator("#sp-comb-none")).to_be_visible()

    # the scope of the draft's example 2: 3 machines, 7 pallets
    page.click("#sp-scope-btn")
    expect(page.locator("#sp-scope-out")).to_contain_text("21 combinations", timeout=20000)
    expect(page.locator("#sp-scope-out")).to_contain_text("1050")

    # combinations of an imported data set: the tag "machine" is the factor
    page.click("nav.tabs button[data-tab=import]")
    csv = tmp_path / "combo.csv"
    rng = np.random.default_rng(8)
    lines = ["pallet,machine,diameter"]
    for i in range(60):
        pallet, machine = f"P{i % 3 + 1}", f"M{i % 2 + 1}"
        lines.append(f"{pallet},{machine},{10 + rng.normal(0, 0.05) + (0.4 if pallet == 'P3' else 0):.4f}")
    csv.write_text("\n".join(lines), encoding="utf-8")
    page.set_input_files("#file", str(csv))
    page.select_option("#col-value", "diameter")
    page.locator("#col-tags input[value=pallet]").check()
    page.click("#import-btn")
    expect(page.locator("#data-counts")).to_contain_text("60 values", timeout=20000)
    page.click("nav.tabs button[data-tab=special]")
    expect(page.locator("#sp-comb-form")).to_be_visible()
    page.locator("#sp-factors input[value=pallet]").check()
    page.fill("#sp-lsl", "9.5")
    page.fill("#sp-usl", "10.8")
    page.click("#sp-comb-btn")
    expect(page.locator("#sp-comb-out")).to_contain_text("60 values together", timeout=20000)
    expect(page.locator("#sp-comb-out")).to_contain_text("deviates")
    expect(page.locator("#sp-comb-out")).to_contain_text("Complete: 3 of 3 combinations")

    # GD&T: the draft's bore
    rng = np.random.default_rng(4)
    rows = [f"{a:.4f} {b:.4f} {c:.4f}" for a, b, c in zip(rng.normal(20.1, 0.03, 80), rng.normal(0, 0.03, 80), rng.normal(0, 0.03, 80))]
    page.fill("#gd-data", "\n".join(rows))
    page.fill("#gd-boot", "0")
    page.click("#gd-run")
    expect(page.locator("#gd-out")).to_contain_text("Virtual size 19.8", timeout=20000)
    expect(page.locator("#gd-out")).to_contain_text("Ppk =")
    expect(page.locator("#gd-out")).to_contain_text("draft prints")
    page.fill("#gd-data", "20.1\n20.1 0.1")
    page.click("#gd-run")
    expect(page.locator("#errors")).to_contain_text("2 numbers", timeout=20000)
    assert problems == []
    ctx.close()


def test_multivariate_nested_and_trend_in_the_browser(server, browser, tmp_path):
    expect = playwright_sync.expect
    ctx = browser.new_context(viewport={"width": 1300, "height": 1000}, locale="en")
    page = ctx.new_page()
    problems = []
    page.on("pageerror", lambda e: problems.append(str(e)))
    page.goto(server)
    sign_in(page)

    # lots > parts > measurements, with the lot as the large source; the same file has a wear trend inside the lots
    rng = np.random.default_rng(12)
    lines = ["lot,part,value"]
    k = 0
    for a in range(5):
        ea = rng.normal(0, 1.0)
        for b in range(3):
            for c in range(4):
                lines.append(f"L{a + 1},P{b + 1},{10 + ea + 0.02 * (k % 20) + rng.normal(0, 0.1):.4f}")
                k += 1
    csv = tmp_path / "nested.csv"
    csv.write_text("\n".join(lines), encoding="utf-8")
    page.click("nav.tabs button[data-tab=import]")
    page.set_input_files("#file", str(csv))
    page.select_option("#col-value", "value")
    page.select_option("#col-subgroup", "")  # the lot is a tag here, not the subgroup label
    page.locator("#col-tags input[value=lot]").check()
    page.locator("#col-tags input[value=part]").check()
    page.click("#import-btn")
    expect(page.locator("#data-counts")).to_contain_text("60 values", timeout=20000)
    page.click("nav.tabs button[data-tab=special]")

    # multivariate
    page.fill("#mv-limits", "9 11 x\n4 6 y")
    pts = rng.multivariate_normal([10.0, 5.0], [[0.01, 0.004], [0.004, 0.02]], size=80)
    page.fill("#mv-data", "\n".join(f"{a:.4f} {b:.4f}" for a, b in pts))
    page.click("#mv-run")
    expect(page.locator("#mv-out")).to_contain_text("Pm =", timeout=20000)
    expect(page.locator("#mv-out")).to_contain_text("Case 1")
    page.fill("#mv-data", "1 2 3\n4 5 6")
    page.click("#mv-run")
    expect(page.locator("#errors")).to_contain_text("one number per limit line", timeout=20000)

    # nested: the levels are filled in from the tags
    expect(page.locator("#nest-levels")).to_have_value("lot, part")
    page.click("#nest-run")
    expect(page.locator("#nest-out")).to_contain_text("The largest source: lot", timeout=20000)
    expect(page.locator("#nest-out")).to_contain_text("within the last level")

    # trend
    page.fill("#tr-cycle", "20")
    page.fill("#tr-lsl", "8")
    page.fill("#tr-usl", "13")
    page.click("#tr-run")
    expect(page.locator("#tr-out")).to_contain_text("Slope per sample", timeout=20000)
    expect(page.locator("#tr-out svg")).to_be_visible()
    expect(page.locator("#tr-out")).to_contain_text("never capability indices")

    # the results go into the report as annex E
    page.click("#to-analysis") if page.locator("#to-analysis").is_visible() else page.click("nav.tabs button[data-tab=analysis]")
    page.fill("#a-lsl", "8")
    page.fill("#a-usl", "13")
    page.select_option("#a-model", "C3")
    expect(page.locator("#a-model-rec")).to_contain_text("Table 10-2 of the draft for model C3", timeout=20000)
    expect(page.locator("#a-model-rec")).to_contain_text("acceptance chart")
    expect(page.locator("#a-model-rec")).to_contain_text("Sample size: smaller; sampling frequency: higher")
    page.select_option("#a-model", "A1")
    expect(page.locator("#a-model-rec")).to_contain_text("Shewhart chart", timeout=20000)
    page.select_option("#a-class", "major")
    page.click("#run-btn")
    expect(page.locator("#rp-special-wrap")).to_be_visible(timeout=40000)
    page.locator("#rp-special input[value=trend]").check()
    page.locator("#rp-special input[value=nested]").check()
    page.locator("#rp-special input[value=multivariate]").check()
    page.click("#rp-create")
    expect(page.locator("#rp-created")).to_contain_text("created", timeout=40000)
    href = page.locator("#rp-open").get_attribute("href")
    html = page.request.get(server.rstrip("/") + href).text()
    assert "Annex E" in html and "regression control chart" in html and "nested variance components" in html and "Multi-stage machining: scope" not in html
    assert problems == []
    ctx.close()


def test_chart_selection_guide_in_the_browser(server, browser):
    expect = playwright_sync.expect
    ctx = browser.new_context(viewport={"width": 1300, "height": 1000}, locale="en")
    page = ctx.new_page()
    problems = []
    page.on("pageerror", lambda e: problems.append(str(e)))
    page.goto(server)
    sign_in(page)
    page.click("nav.tabs button[data-tab=tools]")
    expect(page.locator("#guide-question")).to_contain_text("What kind of data is charted?", timeout=20000)
    page.click("#guide-question button[data-option=variable]")
    page.click("#guide-question button[data-option=one]")
    page.click("#guide-question button[data-option=no]")
    expect(page.locator("#guide-question")).to_contain_text("Which situation", timeout=20000)
    page.click("#guide-question button[data-option=rare_events]")
    expect(page.locator("#guide-result")).to_contain_text("G chart", timeout=20000)
    expect(page.locator("#guide-result")).to_contain_text("tool: Charts for special situations")
    page.click("#guide-back")
    page.click("#guide-question button[data-option=autocorrelation]")
    expect(page.locator("#guide-result")).to_contain_text("CUSUM chart", timeout=20000)
    expect(page.locator("#guide-result")).to_contain_text("Regression control chart")
    page.click("#guide-restart")
    page.click("#guide-question button[data-option=attribute]")
    page.click("#guide-question button[data-option=defects]")
    page.click("#guide-question button[data-option=yes]")
    expect(page.locator("#guide-result")).to_contain_text("c chart", timeout=20000)
    expect(page.locator("#guide-result")).to_contain_text("Laney")
    assert problems == []
    ctx.close()


def test_attribute_measurement_system_in_the_browser(server, browser):
    expect = playwright_sync.expect
    ctx = browser.new_context(viewport={"width": 1300, "height": 1000}, locale="en")
    page = ctx.new_page()
    problems = []
    page.on("pageerror", lambda e: problems.append(str(e)))
    page.goto(server)
    sign_in(page)
    page.click("nav.tabs button[data-tab=msa]")
    page.click("#ms-new")
    page.fill("#mse-name", "Visual check of the thread")
    page.select_option("#mse-kind", "attribute")
    expect(page.locator("#mse-resolution")).to_be_hidden()  # the fields of a variable system are not shown
    page.click("#ms-editor summary")  # the criteria of the gate are in a closed section
    expect(page.locator("#msp-attr_eff_pass")).to_be_visible()
    expect(page.locator("#msp-grr_pass")).to_be_hidden()
    page.click("#mse-save")
    expect(page.locator("#ms-title")).to_have_text("Visual check of the thread", timeout=20000)
    expect(page.locator("#ms-status")).to_contain_text("blocked")
    expect(page.locator("#ms-checks")).to_contain_text("Attribute agreement study")
    expect(page.locator("#ms-checks")).not_to_contain_text("Gauge R&R")  # a check of the other kind of system
    expect(page.locator("#mss-kind")).to_have_value("attribute")

    def table(miss_parts=()):
        rows = ["ref A:1 A:2 A:3 B:1 B:2 B:3"]
        for i in range(40):
            ref = "ok" if i < 20 else "ng"
            a = "ok" if i in miss_parts else ref
            rows.append(f"{ref} {a} {a} {a} {ref} {ref} {ref}")
        return "\n".join(rows)

    page.fill("#mss-data", table().replace("ok", "maybe", 1))
    page.click("#mss-save")
    expect(page.locator("#errors")).to_contain_text("is not a decision", timeout=20000)
    page.fill("#mss-data", table())
    page.click("#mss-save")
    expect(page.locator("#ms-status")).to_contain_text("open", timeout=20000)
    expect(page.locator("#ms-checks")).to_contain_text("capable")
    expect(page.locator("#ms-attr-detail")).to_contain_text("Attribute study 1")
    expect(page.locator("#ms-attr-detail")).to_contain_text("Only 40 parts")
    assert problems == []
    ctx.close()


def test_charts_for_special_situations_and_the_laney_option_in_the_browser(server, browser):
    expect = playwright_sync.expect
    ctx = browser.new_context(viewport={"width": 1300, "height": 1000}, locale="en")
    page = ctx.new_page()
    problems = []
    page.on("pageerror", lambda e: problems.append(str(e)))
    page.goto(server)
    sign_in(page)
    rng = np.random.default_rng(31)
    sizes = rng.integers(100, 200, 40)
    counts = rng.binomial(sizes, np.clip(0.05 + rng.normal(0, 0.02, 40), 0.002, 0.5))

    page.click("nav.tabs button[data-tab=tools]")
    page.select_option("#sc-kind", "laney-p")
    page.fill("#sc-data", "\n".join(f"{c} {n}" for c, n in zip(counts, sizes)))
    page.click("#sc-run")
    expect(page.locator("#sc-out")).to_contain_text("Sigma of the z values", timeout=20000)
    expect(page.locator("#sc-out svg")).to_be_visible()
    expect(page.locator("#sc-out")).to_contain_text("dashed lines")
    page.select_option("#sc-kind", "g")
    expect(page.locator("#sc-reference")).to_be_visible()
    page.fill("#sc-data", "\n".join(str(int(v)) for v in rng.geometric(0.03, 50) - 1))
    page.click("#sc-run")
    expect(page.locator("#sc-out")).to_contain_text("p of the geometric distribution", timeout=20000)
    page.select_option("#sc-kind", "delta-target")
    page.fill("#sc-data", "A 10.1\nB 25.0")
    page.fill("#sc-targets", "A 10")
    page.click("#sc-run")
    expect(page.locator("#errors")).to_contain_text("at least 20", timeout=20000)
    page.select_option("#sc-kind", "uwma")
    page.fill("#sc-data", "1\nx")
    page.click("#sc-run")
    expect(page.locator("#errors")).to_contain_text("is not a number", timeout=20000)

    # the Laney option of a monitor: only for p and u, and then only reference counts as the source
    page.click("nav.tabs button[data-tab=monitor]")
    page.click("#mon-new")
    page.select_option("#me-kind", "xbar-s")
    expect(page.locator("#me-laney-label")).to_be_hidden()
    page.select_option("#me-kind", "p")
    expect(page.locator("#me-laney-label")).to_be_visible()
    page.check("#me-laney")
    expect(page.locator("#me-src-type option[value=rate]")).to_be_hidden()
    assert page.locator("#me-src-type").input_value() == "counts"
    assert problems == []
    ctx.close()


def test_process_characterization_in_the_browser(server, browser):
    expect = playwright_sync.expect
    ctx = browser.new_context(viewport={"width": 1300, "height": 1000}, locale="en")
    page = ctx.new_page()
    problems = []
    page.on("pageerror", lambda e: problems.append(str(e)))
    page.goto(server)
    sign_in(page)
    page.click("nav.tabs button[data-tab=tools]")
    page.select_option("#doe-mode", "factorial")
    A = [-1] * 3 + [1] * 3 + [-1] * 3 + [1] * 3
    B = [-1] * 6 + [1] * 6
    y = [28, 25, 27, 36, 32, 32, 18, 19, 23, 31, 30, 29]
    page.fill("#doe-data", "y A B\n" + "\n".join(f"{v} {a} {b}" for v, a, b in zip(y, A, B)))
    page.click("#doe-run")
    expect(page.locator("#doe-out")).to_contain_text("8 degrees of freedom", timeout=20000)
    expect(page.locator("#doe-out table")).to_contain_text("A:B")
    page.select_option("#doe-mode", "regression")
    page.fill("#doe-data", "y x\n1 1\n2 2\nx 3\n4 4\n5 5")
    page.click("#doe-run")
    expect(page.locator("#errors")).to_contain_text("Give a header row", timeout=20000)
    page.fill("#doe-data", "y x\n1.1 1\n1.9 2\n3.2 3\n3.8 4\n5.1 5")
    page.click("#doe-run")
    expect(page.locator("#doe-out")).to_contain_text("adjusted R²", timeout=20000)
    assert problems == []
    ctx.close()


def test_fraction_surface_and_plan_in_the_browser(server, browser):
    expect = playwright_sync.expect
    ctx = browser.new_context(viewport={"width": 1300, "height": 1000}, locale="en")
    page = ctx.new_page()
    problems = []
    page.on("pageerror", lambda e: problems.append(str(e)))
    page.goto(server)
    sign_in(page)
    page.click("nav.tabs button[data-tab=tools]")
    # the plan of a 2^(4-1): defining relation, resolution and the table of runs
    page.locator("#doe-card details summary").click()
    page.select_option("#dp-kind", "fractional")
    page.fill("#dp-k", "4")
    page.click("#dp-run")
    expect(page.locator("#dp-out")).to_contain_text("I = ABCD", timeout=20000)
    expect(page.locator("#dp-out")).to_contain_text("resolution IV")
    expect(page.locator("#dp-out table tr")).to_have_count(9)
    page.select_option("#dp-kind", "central_composite")
    expect(page.locator("#dp-axial-wrap")).to_be_visible()
    expect(page.locator("#dp-gen-wrap")).to_be_hidden()
    page.fill("#dp-k", "2")
    page.fill("#dp-seed", "5")
    page.click("#dp-run")
    expect(page.locator("#dp-out")).to_contain_text("13 runs", timeout=20000)
    expect(page.locator("#dp-out")).to_contain_text("random (seed 5)")
    # the half fraction of the filtration example (Montgomery 8.1): the aliased effects
    page.select_option("#doe-mode", "fractional")
    runs = {"": 45, "ab": 65, "ac": 60, "bc": 80, "ad": 100, "bd": 45, "cd": 75, "abcd": 96}
    rows = ["y A B C D"]
    for key, v in runs.items():
        rows.append(f"{v} " + " ".join("1" if c in key else "-1" for c in "abcd"))
    page.fill("#doe-data", "\n".join(rows))
    page.click("#doe-run")
    expect(page.locator("#doe-out")).to_contain_text("I = A:B:C:D", timeout=20000)
    expect(page.locator("#doe-out")).to_contain_text("resolution IV")
    ac = page.locator("#doe-out table tr:has(td:text-is('A:C'))")
    expect(ac).to_contain_text("-18.5")
    expect(ac).to_contain_text("+ B:D")
    # the response surface of the process yield (Montgomery 11.2)
    page.select_option("#doe-mode", "surface")
    x1 = [-1, 1, -1, 1, -1.414, 1.414, 0, 0, 0, 0, 0, 0, 0]
    x2 = [-1, -1, 1, 1, 0, 0, -1.414, 1.414, 0, 0, 0, 0, 0]
    y = [76.5, 78.0, 77.0, 79.5, 75.6, 78.4, 77.0, 78.5, 79.9, 80.3, 80.0, 79.7, 79.8]
    page.fill("#doe-data", "y x1 x2\n" + "\n".join(f"{v} {a} {b}" for v, a, b in zip(y, x1, x2)))
    page.click("#doe-run")
    expect(page.locator("#doe-out")).to_contain_text("has a maximum", timeout=20000)
    expect(page.locator("#doe-out")).to_contain_text("x1 = 0.3892")
    expect(page.locator("#doe-out")).to_contain_text("80.212")
    expect(page.locator("#doe-out")).to_contain_text("Lack of fit")
    assert problems == []
    ctx.close()


def test_linearity_nested_grr_and_budget_in_the_browser(server, browser):
    expect = playwright_sync.expect
    ctx = browser.new_context(viewport={"width": 1300, "height": 1000}, locale="en")
    page = ctx.new_page()
    problems = []
    page.on("pageerror", lambda e: problems.append(str(e)))
    page.goto(server)
    sign_in(page)
    page.click("nav.tabs button[data-tab=msa]")
    page.click("#ms-new")
    page.fill("#mse-name", "Tensile tester")
    page.fill("#mse-resolution", "0.01")
    page.fill("#mse-tolerance", "6")
    page.click("#mse-save")
    expect(page.locator("#ms-title")).to_have_text("Tensile tester", timeout=20000)
    expect(page.locator("#ms-checks")).to_contain_text("Linearity and bias")
    rng = np.random.default_rng(12)

    # linearity: the bias grows with the size
    page.select_option("#mss-kind", "linearity")
    expect(page.locator("#mss-pv-label")).to_be_visible()
    page.fill("#mss-pv", "6")
    page.fill("#mss-data", "2 3")
    page.click("#mss-save")
    expect(page.locator("#errors")).to_contain_text("at least 2 readings", timeout=20000)
    rows = [f"{x} " + " ".join(f"{x * 1.05 + rng.normal(0, 0.1):.3f}" for _ in range(10)) for x in (2, 4, 6, 8, 10)]
    page.fill("#mss-data", "\n".join(rows))
    page.click("#mss-save")
    expect(page.locator("#ms-more-detail")).to_contain_text("Linearity and bias, study 1", timeout=20000)
    expect(page.locator("#ms-more-detail")).to_contain_text("leaves the confidence band")
    expect(page.locator("#ms-checks")).to_contain_text("leaves the confidence band")

    # nested gauge R&R
    page.select_option("#mss-kind", "grr_nested")
    expect(page.locator("#mss-pv-label")).to_be_hidden()
    data = "\n".join(" ; ".join(" ".join(f"{10 + o * 0.01 + p * 0.5 + rng.normal(0, 0.05):.3f}" for _ in range(2)) for p in range(4)) for o in range(3))
    page.fill("#mss-data", data)
    page.click("#mss-save")
    expect(page.locator("#ms-more-detail")).to_contain_text("Nested gauge R&R, study 2", timeout=20000)
    expect(page.locator("#ms-checks")).to_contain_text("Gauge R&R")

    # uncertainty budget
    page.select_option("#mss-kind", "budget")
    page.fill("#mss-data", "calibration ; B ; 0.1 ; normal_k2\nresolution ; B ; 0.01 ; rectangular\nrepeatability ; A ; 0.03 ; standard ; 1 ; 19")
    page.click("#mss-save")
    expect(page.locator("#ms-more-detail")).to_contain_text("Uncertainty budget, study 3", timeout=20000)
    expect(page.locator("#ms-more-detail")).to_contain_text("effective degrees of freedom")
    expect(page.locator("#ms-checks")).to_contain_text("2U/T")
    page.fill("#mss-data", "only a name")
    page.click("#mss-save")
    expect(page.locator("#errors")).to_contain_text("not understood", timeout=20000)
    assert problems == []
    ctx.close()


def test_standardised_option_of_a_p_monitor_in_the_browser(server, browser):
    expect = playwright_sync.expect
    ctx = browser.new_context(viewport={"width": 1300, "height": 1000}, locale="en")
    page = ctx.new_page()
    problems = []
    page.on("pageerror", lambda e: problems.append(str(e)))
    page.goto(server)
    sign_in(page)
    page.click("nav.tabs button[data-tab=monitor]")
    page.click("#mon-new")
    page.select_option("#me-kind", "xbar-s")
    expect(page.locator("#me-std-label")).to_be_hidden()
    page.select_option("#me-kind", "p")
    expect(page.locator("#me-std-label")).to_be_visible()
    page.check("#me-laney")
    page.check("#me-std")  # one of the two options: the later choice wins
    assert not page.locator("#me-laney").is_checked() and page.locator("#me-std").is_checked()
    page.check("#me-laney")
    assert page.locator("#me-laney").is_checked() and not page.locator("#me-std").is_checked()
    assert problems == []
    ctx.close()


def test_lot_release_in_the_browser(server, browser):
    expect = playwright_sync.expect
    ctx = browser.new_context(viewport={"width": 1300, "height": 1000}, locale="en")
    page = ctx.new_page()
    problems = []
    page.on("pageerror", lambda e: problems.append(str(e)))
    page.goto(server)
    sign_in(page)
    page.click("nav.tabs button[data-tab=lots]")
    page.fill("#lot-no", "L-2026-001")
    page.fill("#lot-quantity", "200")
    page.fill("#lot-product", "Bracket")
    page.click("#lot-create")
    expect(page.locator("#lot-title")).to_contain_text("L-2026-001: 200 parts", timeout=20000)
    expect(page.locator("#lot-evidence")).to_contain_text("not linked to a monitor")
    page.select_option("#lotd-decision", "release")
    page.click("#lotd-save")
    expect(page.locator("#errors")).to_contain_text("reason", timeout=20000)  # a lot without a monitor is released with a reason
    # sorting: all parts are inspected and the counts must add up
    page.select_option("#lotd-decision", "sort")
    expect(page.locator("#lotd-good")).to_be_visible()
    page.fill("#lotd-good", "180")
    page.fill("#lotd-rejected", "10")
    page.click("#lotd-save")
    expect(page.locator("#errors")).to_contain_text("add up", timeout=20000)
    page.fill("#lotd-good", "185")
    page.fill("#lotd-rejected", "15")
    page.select_option("#lotd-rejected-to", "rework")
    page.click("#lotd-save")
    expect(page.locator("#lot-status")).to_have_text("decided", timeout=20000)
    expect(page.locator("#lot-decision-box")).to_contain_text("Sorted: 200 inspected, 185 good, 15 rejected")
    expect(page.locator("#lot-decide-card")).to_be_hidden()
    page.once("dialog", lambda d: d.accept("Counted wrongly"))
    page.click("#lot-reopen")
    expect(page.locator("#lot-status")).to_have_text("awaiting a decision", timeout=20000)
    expect(page.locator("#lot-history")).to_contain_text("Reopened")
    page.click("#lot-back")
    expect(page.locator("#lot-list")).to_contain_text("L-2026-001", timeout=20000)
    assert problems == []
    ctx.close()


def test_random_plan_and_audit_anchor_in_the_browser(server, browser):
    expect = playwright_sync.expect
    ctx = browser.new_context(viewport={"width": 1300, "height": 1000}, locale="en")
    page = ctx.new_page()
    problems = []
    page.on("pageerror", lambda e: problems.append(str(e)))
    page.goto(server)
    sign_in(page, user="admin")
    page.click("nav.tabs button[data-tab=tools]")
    page.fill("#rpl-k", "12")
    page.fill("#rpl-n", "5")
    page.fill("#rpl-seed", "42")
    page.fill("#rpl-factors", "tool: A B C\nshift: 1 2")
    page.click("#rpl-run")
    expect(page.locator("#rpl-out")).to_contain_text("12 subgroups of 5 parts (60 parts), seed 42", timeout=20000)
    expect(page.locator("#rpl-out table tr")).to_have_count(13)
    expect(page.locator("#rpl-out")).to_contain_text("tool: A × 4, B × 4, C × 4")
    page.fill("#rpl-factors", "tool")
    page.click("#rpl-run")
    expect(page.locator("#errors")).to_contain_text("at least one level", timeout=20000)
    page.click("nav.tabs button[data-tab=admin]")
    page.click("#audit-verify")
    expect(page.locator("#audit-verdict")).not_to_be_empty(timeout=20000)
    text = page.locator("#audit-verdict").inner_text()
    import re
    n, h = re.search(r"(\d+)", text).group(1), re.search(r"\b([0-9a-f]{64})\b", text).group(1)
    page.fill("#anchor-entries", n)
    page.fill("#anchor-hash", h)
    page.click("#anchor-check")
    expect(page.locator("#anchor-verdict")).to_contain_text("still holds the entry you kept", timeout=20000)
    page.fill("#anchor-hash", "0" * 64)
    page.click("#anchor-check")
    expect(page.locator("#anchor-verdict")).to_contain_text("not the one you kept", timeout=20000)
    page.fill("#anchor-hash", "xyz")
    page.click("#anchor-check")
    expect(page.locator("#errors")).to_contain_text("64 characters", timeout=20000)
    assert problems == []
    ctx.close()


def test_iso_22514_7_studies_in_the_browser(server, browser):
    from tests.test_iso22514_7 import A1_X, A1_Y, A4, FIG6, figure_6_results

    expect = playwright_sync.expect
    ctx = browser.new_context(viewport={"width": 1300, "height": 1000}, locale="en")
    page = ctx.new_page()
    problems = []
    page.on("pageerror", lambda e: problems.append(str(e)))
    page.goto(server)
    sign_in(page)
    page.click("nav.tabs button[data-tab=msa]")
    page.click("#ms-new")
    page.fill("#mse-name", "Optical comparator")
    page.fill("#mse-resolution", "0.005")
    page.fill("#mse-tolerance", "9")
    page.click("#mse-save")
    expect(page.locator("#ms-title")).to_have_text("Optical comparator", timeout=20000)
    expect(page.locator("#ms-checks")).to_contain_text("ISO 22514-7 study")

    # the worked example of the standard (annex A)
    page.select_option("#mss-kind", "iso_study")
    expect(page.locator("#mss-iso")).to_be_visible()
    expect(page.locator("#mss-data")).to_be_hidden()
    page.fill("#mss-lower", "2")
    page.fill("#mss-upper", "11")
    page.fill("#isof-cal-u", "0.01")
    page.fill("#isof-lin-data", "\n".join(f"{x} " + " ".join(str(v) for v in row) for x, row in zip(A1_X, A1_Y)))
    page.fill("#isof-proc-data", "\n".join(" ; ".join(" ".join(str(v) for v in part) for part in op) for op in A4))
    page.click("#mss-save")
    expect(page.locator("#ms-more-detail")).to_contain_text("ISO 22514-7 study 1", timeout=20000)
    expect(page.locator("#ms-more-detail")).to_contain_text("Q_MS = 3.71")
    expect(page.locator("#ms-more-detail")).to_contain_text("Q_MP = 9.3 %")
    expect(page.locator("#ms-more-detail")).to_contain_text("pooled with the repeatability")
    expect(page.locator("#ms-checks")).to_contain_text("ISO 22514-7: Q_MP 9.3 %")
    page.fill("#isof-lin-data", "2 3 4")
    page.click("#mss-save")
    expect(page.locator("#errors")).to_contain_text("at least 3 readings", timeout=20000)
    for field in ("#mss-lower", "#mss-upper", "#isof-cal-u", "#isof-lin-data", "#isof-proc-data"):
        page.fill(field, "")  # the fields of a failed study stay: the next study starts empty

    # an attribute system: the uncertainty range of 12.3
    page.click("#ms-back")
    page.click("#ms-new")
    page.fill("#mse-name", "Go/no-go gauge")
    page.select_option("#mse-kind", "attribute")
    page.click("#mse-save")
    expect(page.locator("#ms-title")).to_have_text("Go/no-go gauge", timeout=20000)
    expect(page.locator("#mss-kind option[value=uncertainty_range]")).to_be_attached()
    page.select_option("#mss-kind", "uncertainty_range")
    expect(page.locator("#mss-lower")).to_be_visible()
    res = figure_6_results()
    header = "ref " + " ".join(f"{op}:{t + 1}" for op in res for t in range(3))
    rows = [header] + [f"{ref} " + " ".join("ok" if res[op][t][i] else "ng" for op in res for t in range(3)) for i, ref in enumerate(FIG6)]
    page.fill("#mss-data", "\n".join(rows))
    page.click("#mss-save")
    expect(page.locator("#errors")).to_contain_text("lower and the upper", timeout=20000)
    page.fill("#mss-lower", "0.45")
    page.fill("#mss-upper", "0.55")
    page.click("#mss-save")
    expect(page.locator("#ms-more-detail")).to_contain_text("Uncertainty range, study 1", timeout=20000)
    expect(page.locator("#ms-more-detail")).to_contain_text("Q_attr = 23.8 %")
    expect(page.locator("#ms-checks")).to_contain_text("Uncertainty range Q_attr 23.8 %")
    assert problems == []
    ctx.close()


def test_linearity_monitoring_in_the_browser(server, browser):
    from tests.test_iso22514_7 import A1_X, A1_Y

    expect = playwright_sync.expect
    ctx = browser.new_context(viewport={"width": 1300, "height": 1000}, locale="en")
    page = ctx.new_page()
    problems = []
    page.on("pageerror", lambda e: problems.append(str(e)))
    page.goto(server)
    sign_in(page)
    page.click("nav.tabs button[data-tab=msa]")
    page.click("#ms-new")
    page.fill("#mse-name", "Microscope")
    page.fill("#mse-resolution", "0.005")
    page.fill("#mse-tolerance", "9")
    page.click("#mse-save")
    expect(page.locator("#ms-title")).to_have_text("Microscope", timeout=20000)
    expect(page.locator("#lm-card")).to_be_hidden()
    page.select_option("#mss-kind", "iso_study")
    page.fill("#isof-lin-data", "\n".join(f"{x} " + " ".join(str(v) for v in row) for x, row in zip(A1_X, A1_Y)))
    page.click("#mss-save")
    expect(page.locator("#lm-card")).to_be_visible(timeout=20000)
    ok = "\n".join(f"{x} {0.2358 + 0.987 * x + 0.02:.4f} {0.2358 + 0.987 * x - 0.03:.4f}" for x in (2, 6, 10))
    page.fill("#lm-data", ok)
    page.click("#lm-run")
    expect(page.locator("#lm-out")).to_contain_text("still valid", timeout=20000)
    drift = "\n".join(f"{x} {0.2358 + 0.987 * x + 0.6:.4f}" for x in (2, 6, 10))
    page.fill("#lm-data", drift)
    page.click("#lm-run")
    expect(page.locator("#lm-out")).to_contain_text("has to be updated", timeout=20000)
    page.fill("#lm-data", "2 1")
    page.click("#lm-run")
    expect(page.locator("#errors")).to_contain_text("at least 2 standards", timeout=20000)
    assert problems == []
    ctx.close()


def test_a_dfd_file_is_read_in_the_browser(server, browser, tmp_path):
    from tests.test_dfq import FILE

    expect = playwright_sync.expect
    ctx = browser.new_context(viewport={"width": 1300, "height": 1000}, locale="en")
    page = ctx.new_page()
    problems = []
    page.on("pageerror", lambda e: problems.append(str(e)))
    page.goto(server)
    sign_in(page)
    page.click("nav.tabs button[data-tab=tools]")
    path = tmp_path / "shaft.dfd"
    path.write_bytes(FILE.replace("K2111/3 0.05", "K2111/3 0.05\r\nK1900/1 late remark").encode())
    page.set_input_files("#dfd-file", str(path))
    expect(page.locator("#dfd-out")).to_contain_text("Part 4711-A, Shaft 12 mm: 3 characteristic(s).", timeout=20000)
    expect(page.locator("#dfd-out table")).to_contain_text("Diameter")
    expect(page.locator("#dfd-out table")).to_contain_text("12.02")
    expect(page.locator("#dfd-out")).to_contain_text("a part field K1900 follows a characteristic field")
    page.locator("#dfd-out button").first.click()
    expect(page.locator("#dfd-out")).to_contain_text("The limits of Diameter are in the analysis form")
    assert page.locator("#a-lsl").input_value() == "11.98" and page.locator("#a-usl").input_value() == "12.02"
    bad = tmp_path / "bad.dfd"
    bad.write_bytes(b"hello world")
    page.set_input_files("#dfd-file", str(bad))
    expect(page.locator("#errors")).to_contain_text("K0100", timeout=20000)
    assert problems == []
    ctx.close()


def test_an_improvement_cycle_in_the_browser(server, browser):
    expect = playwright_sync.expect
    ctx = browser.new_context(viewport={"width": 1300, "height": 1000}, locale="en")
    page = ctx.new_page()
    problems = []
    page.on("pageerror", lambda e: problems.append(str(e)))
    page.goto(server)
    sign_in(page)
    page.click("nav.tabs button[data-tab=monitor]")
    page.click("#mon-new")
    page.fill("#me-name", "Bore diameter")
    page.fill("#me-characteristic", "diameter")
    page.fill("#me-lsl", "9.5")
    page.fill("#me-usl", "10.5")
    page.fill("#me-mu", "10")
    page.fill("#me-sigma", "0.1")
    page.locator("#me-ocap tr[data-key=default] input").nth(0).fill("Stop")
    page.locator("#me-ocap tr[data-key=default] input").nth(1).fill("Setter")
    page.click("#me-save")
    expect(page.locator("#md-title")).to_have_text("Bore diameter", timeout=20000)
    page.click("nav.tabs button[data-tab=improve]")
    page.fill("#imp-title", "Fewer rejects of the bore")
    page.fill("#imp-target", "1.5")
    page.fill("#imp-plan", "Replace the spindle bearings")
    page.click("#imp-create")
    expect(page.locator("#imp-d-title")).to_have_text("Fewer rejects of the bore", timeout=20000)
    expect(page.locator("#imp-d-status")).to_have_text("planned")
    expect(page.locator("#imp-d-body")).to_contain_text("No baseline yet")  # the monitor has no points yet
    page.click("#imp-act-btn")
    expect(page.locator("#errors")).to_be_visible(timeout=20000)  # nothing was said about what was done
    page.fill("#imp-act-text", "Bearings replaced")
    page.click("#imp-act-btn")
    expect(page.locator("#imp-d-status")).to_have_text("implemented, to be verified", timeout=20000)
    expect(page.locator("#imp-act-text")).to_be_hidden()
    page.click("#imp-act-btn")
    expect(page.locator("#errors")).to_contain_text("at least 25 valid points", timeout=20000)
    page.click("#imp-back")
    expect(page.locator("#imp-list")).to_contain_text("Fewer rejects of the bore", timeout=20000)
    assert problems == []
    ctx.close()


def test_excel_import_saves_a_template_and_uses_it_for_the_next_file(server, browser, tmp_path):
    """An Excel file with notes above the column names: choose the sheet and line, save the template, and the
    next file (columns in another order) is filled in by choosing the template."""
    from openpyxl import Workbook

    expect = playwright_sync.expect

    def book(path, order):
        rng = np.random.default_rng(7)
        wb = Workbook()
        ws = wb.active
        ws.title = "Notes"
        ws.append(["nothing here"])
        data = wb.create_sheet("CMM")
        data.append(["Line 4 report"])
        data.append([])
        cols = {"批號": lambda g: f"B{g // 5 + 1:02d}", "直徑": lambda g: round(10 + float(rng.normal(0, 0.1)), 4), "機台": lambda g: f"M{g % 2 + 1}"}
        data.append(order)
        for g in range(50):
            data.append([cols[c](g) for c in order])
        wb.save(path)
        return path

    first = book(tmp_path / "first.xlsx", ["批號", "直徑", "機台"])
    second = book(tmp_path / "second.xlsx", ["機台", "直徑", "批號"])

    ctx = browser.new_context(viewport={"width": 1200, "height": 900}, locale="zh-TW")
    page = ctx.new_page()
    problems = []
    page.on("console", lambda m: problems.append(m.text) if m.type == "error" else None)
    page.on("pageerror", lambda e: problems.append(str(e)))
    page.goto(server)
    sign_in(page)

    page.set_input_files("#file", str(first))
    expect(page.locator("#detected")).to_contain_text("Excel")
    expect(page.locator("#sheet-wrap")).to_be_visible()
    page.select_option("#import-sheet", "CMM")
    expect(page.locator("#detected")).to_contain_text("CMM")
    page.fill("#import-header-row", "3")
    page.press("#import-header-row", "Tab")
    expect(page.locator("#col-value")).to_have_value("直徑")  # the columns are read again from the new header line
    expect(page.locator("#col-subgroup")).to_have_value("批號")
    expect(page.locator("#template-note")).to_contain_text("沒有符合")
    page.locator("#tab-import details summary", has_text="存成匯入範本").click()
    page.fill("#tpl-name", "第四線三次元")
    page.locator("#col-tags input[value=機台]").check()
    page.click("#tpl-save")
    expect(page.locator("#template-note")).to_contain_text("已儲存範本")
    expect(page.locator("#import-template")).to_have_value(page.locator("#import-template option").nth(1).get_attribute("value"))
    page.click("#import-btn")
    expect(page.locator("#data-counts")).to_contain_text("共 50 筆")

    # the next file: choose the template, nothing else
    page.click("nav.tabs button[data-tab=import]")
    page.set_input_files("#file", str(second))
    expect(page.locator("#detected")).to_contain_text("Excel")
    page.select_option("#import-sheet", "CMM")
    expect(page.locator("#detected")).to_contain_text("CMM")
    page.fill("#import-header-row", "3")
    page.press("#import-header-row", "Tab")
    expect(page.locator("#col-value")).to_have_value("直徑")  # the file is read again from line 3
    expect(page.locator("#template-note")).to_contain_text("1 個已存的範本符合")
    page.select_option("#import-template", label="第四線三次元")
    expect(page.locator("#col-value")).to_have_value("直徑")
    expect(page.locator("#col-subgroup")).to_have_value("批號")
    expect(page.locator("#col-tags input[value=機台]")).to_be_checked()
    page.click("#import-btn")
    expect(page.locator("#data-counts")).to_contain_text("共 50 筆")
    assert problems == []
    ctx.close()
