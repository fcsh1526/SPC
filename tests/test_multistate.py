import numpy as np
import pytest
from scipy import stats

from spc.core import multistate as ms
from spc.data import Dataset
from spc.service.state_tests import state_tests_for_dataset
from tests.conftest import logged_in_client, make_app
from tests.test_api import err


def groups(seed=1):
    r = np.random.default_rng(seed)
    return [r.normal(0, 1, 20), r.normal(0.5, 2, 25), r.normal(0, 1, 15)]


def test_bartlett_and_fisher_equal_scipy():
    for seed in range(5):
        g = groups(seed)
        b, f = ms.bartlett(g), ms.fisher(g)
        ref_b, ref_f = stats.bartlett(*g), stats.f_oneway(*g)
        assert b["statistic"] == pytest.approx(ref_b.statistic, rel=1e-12) and b["p"] == pytest.approx(ref_b.pvalue, rel=1e-9)
        assert f["statistic"] == pytest.approx(ref_f.statistic, rel=1e-12) and f["p"] == pytest.approx(ref_f.pvalue, rel=1e-9)
        assert b["df"] == 2 and f["df1"] == 2 and f["df2"] == 57


def test_the_critical_value_of_grubbs_is_the_published_one():
    # NIST/SEMATECH e-Handbook table, two-sided, alpha = 0.05
    for n, value in ((10, 2.2900), (25, 2.8217)):
        assert ms.grubbs_critical(n) == pytest.approx(value, abs=1e-4)
    assert ms.grubbs_critical(10, 0.01) > ms.grubbs_critical(10, 0.05)


def test_grubbs_finds_the_most_extreme_value_only_when_it_is_too_far():
    x = list(np.random.default_rng(3).normal(10, 0.2, 30))
    assert ms.grubbs(x)["outlier"] is False
    x[11] = 12.0
    g = ms.grubbs(x)
    assert g["outlier"] and g["index"] == 11 and g["value"] == 12.0 and g["g"] > g["g_crit"]
    assert ms.grubbs([1, 1, 1, 1])["g"] is None  # no variation: nothing to test


def test_the_tests_see_what_was_built_into_the_states():
    r = np.random.default_rng(8)
    same = {"a": r.normal(10, 0.3, 40), "b": r.normal(10, 0.3, 40), "c": r.normal(10, 0.3, 40)}
    t = ms.state_tests(same)
    assert not t["variance_differs"] and not t["location_differs"] and t["outliers"] == []
    diff = {"a": r.normal(10, 0.3, 40), "b": r.normal(11, 1.0, 40), "c": r.normal(10, 0.3, 40), "tiny": [1.0, 2.0]}
    t = ms.state_tests(diff)
    assert t["variance_differs"] and t["location_differs"] and t["skipped"] == ["tiny"]


@pytest.mark.parametrize("bad", [{"a": [1, 2, 3]}, {"a": [1, 2, 3], "b": [1, 1, 1]}, {"a": [1, 2], "b": [3, 4]}])
def test_too_little_or_degenerate_data_is_refused(bad):
    with pytest.raises(ValueError):
        ms.state_tests(bad)


def test_alpha_is_checked():
    with pytest.raises(ValueError):
        ms.state_tests({"a": [1, 2, 4], "b": [2, 3, 7]}, alpha=0.7)


def test_the_states_come_from_subgroup_labels_or_a_tag_and_marked_values_do_not_count():
    r = np.random.default_rng(5)
    v = np.concatenate([r.normal(10, 0.2, 30), r.normal(10.5, 0.2, 30)])
    labels = ["m1"] * 30 + ["m2"] * 30
    ds = Dataset.from_values(v, subgroup=labels)
    out = state_tests_for_dataset(ds)
    assert set(out["states"]) == {"m1", "m2"} and out["by"] == "subgroup" and out["location_differs"]
    ds2 = Dataset.from_values(v, tags={"machine": labels}).mark_invalid([0, 1], "test", "me")
    out2 = state_tests_for_dataset(ds2, by="machine")
    assert out2["states"]["m1"]["n"] == 28 and out2["by"] == "machine"
    with pytest.raises(ValueError):
        state_tests_for_dataset(Dataset.from_values(v))  # no labels and no tag
    with pytest.raises(ValueError):
        state_tests_for_dataset(ds2, by="nothing")


def test_the_api_runs_the_tests_on_a_stored_data_set():
    app = make_app()
    c = logged_in_client(app, "eng")
    r = np.random.default_rng(6)
    v = np.concatenate([r.normal(10, 0.2, 30), r.normal(10, 0.6, 30)])
    key = app.state.store.add(Dataset.from_values(v, subgroup=["a"] * 30 + ["b"] * 30), 1, "two states")
    res = c.post(f"/api/datasets/{key}/state-tests", json={})
    assert res.status_code == 200 and res.json()["variance_differs"] is True and res.json()["bartlett"]["df"] == 1
    assert c.post(f"/api/datasets/{key}/state-tests", json={"alpha": 0.9}).status_code in (400, 422)
    plain = app.state.store.add(Dataset.from_values(v), 1, "no states")
    assert err(c.post(f"/api/datasets/{plain}/state-tests", json={}))["code"] == "invalid_input"


# ------------------------------------------------------------------ ISO 22514-8 procedure and annex B

def test_the_table_for_the_forced_variance_and_the_tests_of_annex_b():
    assert ms.forced_variance([143.1] * 5, 0.1) == pytest.approx(0.0016) and ms.forced_variance([140.2, 140.2, 140.2, 140.1], 0.1) == pytest.approx(0.0074)
    assert ms.forced_variance([140.2, 140.0, 140.2, 140.3, 140.6], 0.1) is None  # six scale marks: not forced
    assert ms.forced_variance([1.0] * 12, 0.1) == pytest.approx(0.001)  # more than 10 values, range 0: 0,10 x resolution squared
    with pytest.raises(ValueError):
        ms.bartlett_forced([[1.0, 1.0, 1.0], [2.0, 2.0, 2.0]])  # no variation and no resolution: it cannot be tested
    assert ms.bartlett_forced([[1.0, 1.0, 1.0], [2.0, 2.0, 2.0]], 0.1)["forced"] == [True, True]
    # without forcing and with the bias correction it is the bartlett of scipy
    g = groups(2)
    assert ms.bartlett_forced(g)["statistic"] == pytest.approx(stats.bartlett(*g).statistic, rel=1e-12)
    assert ms.bartlett_forced(g, bias_correction=False)["statistic"] < ms.bartlett_forced(g)["statistic"] * 1.2


def test_two_states_use_the_f_test_and_the_t_test_or_the_welch_test():
    r = np.random.default_rng(4)
    a, b = r.normal(10, 1, 15), r.normal(10.8, 1, 20)
    t = ms.t_test(a, b)
    assert t["statistic"] == pytest.approx(stats.ttest_ind(a, b).statistic, rel=1e-12) and t["p"] == pytest.approx(stats.ttest_ind(a, b).pvalue, rel=1e-9)
    w = ms.t_test(a, b, equal_variances=False)
    assert w["statistic"] == pytest.approx(stats.ttest_ind(a, b, equal_var=False).statistic, rel=1e-12) and w["p"] == pytest.approx(stats.ttest_ind(a, b, equal_var=False).pvalue, rel=1e-9) and w["welch"]
    f = ms.f_test_variances(a, b)
    assert f["statistic"] == pytest.approx(a.var(ddof=1) / b.var(ddof=1)) and 0 < f["p"] <= 1 and f["critical_low"] < 1 < f["critical_high"]


def test_outliers_are_removed_one_at_a_time_and_not_more_than_a_third():
    states = {"a": [10.0, 10.1, 9.9, 10.0, 10.1, 9.95], "b": [10.1, 10.0, 12.0, 10.05, 9.95, 10.0], "c": [10.0, 10.05, 9.95, 10.1, 10.0, 9.9]}
    s = ms.screen_outliers(states)
    assert [r["value"] for r in s["removed"]] == [12.0] and s["removed"][0]["state"] == "b" and s["removed"][0]["delta_a"] == pytest.approx(12.0 - 10.02)
    assert len(s["states"]["b"]) == 5 and s["first"]["states"]["b"]["outlier"] is True
    many = {"a": [1.0, 1.01, 5.0, 1.02, 9.0, 1.0], "b": [1.0, 1.02, 1.01, 7.0, 1.0, 1.01]}
    try:
        s2 = ms.screen_outliers(many)
        assert len(s2["removed"]) <= 4
    except ValueError as exc:
        assert "one third" in str(exc)


def test_the_types_of_global_dispersion_and_the_capability_of_table_2():
    r = np.random.default_rng(10)
    same = {"a": r.normal(10, 0.2, 8), "b": r.normal(10, 0.2, 8), "c": r.normal(10, 0.2, 8)}
    assert ms.analyse(same, 9.0, 11.0)["type"] == 0
    shift = {"a": r.normal(10, 0.2, 10), "b": r.normal(11, 0.2, 10), "c": r.normal(10.5, 0.2, 10)}
    t1, t2 = ms.analyse(shift, 8.0, 13.0), ms.analyse(shift, 8.0, 13.0, delta_m_variable=True, delta_m_star=1.5)
    assert t1["type"] == 1 and t2["type"] == 2
    assert t1["pm"] == pytest.approx((5.0 - t1["delta_m"]) / (6 * t1["sigma_pooled"])) and t2["pm"] == pytest.approx(5.0 / (6 * t2["sigma_pooled"] + 1.5))
    assert t1["pmk_l"] == pytest.approx((min(v["x50"] for v in t1["states"].values()) - 8.0) / (3 * t1["sigma_pooled"]))
    wide = {"a": r.normal(10, 0.1, 10), "b": r.normal(10, 0.6, 10), "c": r.normal(10, 0.1, 10)}
    t3 = ms.analyse(wide, 8.0, 12.0)
    assert t3["type"] == 3 and t3["pm"] == pytest.approx(4.0 / max(v["d_l"] + v["d_u"] for v in t3["states"].values()))
    mixed = {"a": r.normal(10, 0.1, 10), "b": r.normal(11, 0.6, 10), "c": r.normal(10, 0.1, 10)}
    assert ms.analyse(mixed, 7.0, 14.0)["type"] == 4 and ms.analyse(mixed, 7.0, 14.0, delta_m_variable=True)["type"] == 5
    # the analyst can decide where the tests are not enough
    assert ms.analyse(shift, 8.0, 13.0, locations_equal=True)["type"] == 0
    # the physical outlier widens the dispersion on the side where it can occur
    out = {"a": [10.0, 10.1, 9.9, 10.0, 8.0], "b": [10.2, 10.3, 10.1, 10.2, 10.25]}
    w = ms.analyse(out, 7.0, 13.0, outlier_physical=True, outlier_direction="negative")
    assert w["outliers"] and w["states"]["a"]["d_l"] > w["states"]["a"]["d_u"]
    with pytest.raises(ValueError):
        ms.analyse({"a": [1.0, 2.0]}, 0.0, 3.0)


def test_the_worked_example_a3_of_the_standard_is_reproduced():
    from spc.validation.iso22514 import ADAPTERS

    a = ms.analyse(ADAPTERS, 19.8, 20.2, outlier_physical=True, outlier_direction="negative")
    assert a["type"] == 1 and a["dof"] == 23 and a["sigma_pooled"] == pytest.approx(0.0123, abs=5e-5) and a["delta_m"] == pytest.approx(0.096, abs=5e-4)
    assert (round(a["pm"], 2), round(a["pmk_u"], 2), round(a["pmk_l"], 2)) == (1.25, 2.17, 1.08)
    assert a["widths"]["statistic"] == pytest.approx(3.4297, abs=1e-3) and a["locations"]["statistic"] == pytest.approx(46.85, abs=5e-3) and a["locations"]["critical"] == pytest.approx(2.62, abs=5e-3)


def test_the_api_runs_the_procedure_and_refuses_what_cannot_be_analysed():
    from spc.validation.iso22514 import ADAPTERS

    app = make_app()
    c = logged_in_client(app, "eng")
    labels = [k for k, v in ADAPTERS.items() for _ in v]
    key = app.state.store.add(Dataset.from_values([x for v in ADAPTERS.values() for x in v], subgroup=labels), 1, "adapters")
    body = {"lsl": 19.8, "usl": 20.2, "outlier_physical": True, "outlier_direction": "negative"}
    r = c.post(f"/api/datasets/{key}/multistate", json=body)
    assert r.status_code == 200, r.text
    j = r.json()
    assert j["type"] == 1 and round(j["pm"], 2) == 1.25 and round(j["pmk"], 2) == 1.08 and j["outliers"][0]["value"] == 19.95 and j["by"] == "subgroup"
    assert j["widths"]["name"] == "bartlett" and j["locations"]["name"] == "fisher" and j["locations"]["equal"] is False
    # the analyst decides where the tests are not enough
    d = c.post(f"/api/datasets/{key}/multistate", json={**body, "locations_equal": True}).json()
    assert d["type"] == 0
    assert c.post(f"/api/datasets/{key}/multistate", json={"lsl": 20.2, "usl": 19.8}).status_code == 400
    assert c.post(f"/api/datasets/{key}/multistate", json={"lsl": 1}).status_code == 422
    plain = app.state.store.add(Dataset.from_values([1.0, 2, 3, 4, 5, 6]), 1, "no states")
    assert err(c.post(f"/api/datasets/{plain}/multistate", json=body))["code"] == "invalid_input"
    two = app.state.store.add(Dataset.from_values([10.0, 10.1, 9.9, 10.2, 11.0, 11.1, 10.9, 11.2], subgroup=["a"] * 4 + ["b"] * 4), 1, "two")
    t = c.post(f"/api/datasets/{two}/multistate", json={"lsl": 8, "usl": 13}).json()
    assert t["widths"]["name"] == "fisher" and t["locations"]["name"] == "t"


def test_a_report_carries_the_machine_performance_of_the_states_into_its_archive():
    import copy

    from spc.report import reproduce, verify_archive
    from spc.validation.iso22514 import ADAPTERS

    app = make_app()
    c = logged_in_client(app, "eng")
    labels = [k for k, v in ADAPTERS.items() for _ in v]
    key = app.state.store.add(Dataset.from_values([x for v in ADAPTERS.values() for x in v], subgroup=labels), 1, "adapters")
    base = {"analysis": {"stage": "machine", "lsl": 19.8, "usl": 20.2}, "language": "en"}
    ms = {"outlier_physical": True, "outlier_direction": "negative", "locations_equal": None}
    r = c.post(f"/api/datasets/{key}/reports", json={**base, "multistate": ms})
    assert r.status_code == 200, r.text
    rid = r.json()["id"]
    arc = c.get(f"/api/reports/{rid}/archive.json").json()
    m = arc["multistate"]
    assert (m["lsl"], m["usl"], m["by"]) == (19.8, 20.2, "subgroup") and m["options"]["outlier_physical"] is True
    assert m["result"]["type"] == 1 and round(m["result"]["pm"], 2) == 1.25 and round(m["result"]["pmk"], 2) == 1.08
    assert verify_archive(arc) and reproduce(arc).reproduced
    html = c.get(f"/api/reports/{rid}").text
    assert "Annex D: machine performance of the states" in html and "Pm = 1.25" in html and "19.95" in html and "Decided by the analyst" not in html.split("Annex D")[0]
    assert "Decided by the analyst: the outlier is real" in html or "the outlier is real and can occur downwards only" in html
    zh = c.post(f"/api/datasets/{key}/reports", json={**base, "language": "zh-TW", "multistate": ms}).json()["id"]
    assert "附錄 D：各狀態的機器性能" in c.get(f"/api/reports/{zh}").text
    x = c.get(f"/api/reports/{rid}/report.xlsx")
    assert x.status_code == 200
    import io

    from openpyxl import load_workbook

    cells = [str(cell.value) for ws in load_workbook(io.BytesIO(x.content)).worksheets for row in ws.iter_rows() for cell in row if cell.value]
    assert any("Annex D" in v for v in cells) and any(v.startswith("Pm = 1.25") for v in cells)
    # a changed number breaks the digest, and a stored result that the data no longer give is found by the reproduction
    bad = copy.deepcopy(arc)
    bad["multistate"]["result"]["pm"] = 9.0
    assert not verify_archive(bad)
    from spc.report.archive import build_archive  # the same archive with a wrong result and a fresh digest: only the reproduction can tell

    body = {k: v for k, v in bad.items() if k != "integrity"}
    import hashlib

    from spc.report.archive import canonical

    bad["integrity"] = {**arc["integrity"], "digest": hashlib.sha256(canonical(body)).hexdigest()}
    rep = reproduce(bad)
    assert rep.integrity_ok and not rep.reproduced and any(d.startswith("multistate.pm") or "multistate.pm" in d for d in rep.differences)
    # no multistate: nothing added; the limits of the analysis are used when the block names none; without limits or states: refused
    plain = c.post(f"/api/datasets/{key}/reports", json=base).json()["id"]
    assert "multistate" not in c.get(f"/api/reports/{plain}/archive.json").json() and "Annex D" not in c.get(f"/api/reports/{plain}").text
    assert err(c.post(f"/api/datasets/{key}/reports", json={"analysis": {"stage": "machine", "lsl": 19.8}, "multistate": {}}))["code"] == "multistate_needs_limits"
    assert c.post(f"/api/datasets/{key}/reports", json={"analysis": {"stage": "machine", "lsl": 19.8}, "multistate": {"usl": 20.2}}).status_code == 200  # the limits of the block complete those of the analysis
    assert err(c.post(f"/api/datasets/{key}/reports", json={**base, "multistate": {"by": "nothing"}}))["code"] == "invalid_input"
    assert any(e["action"] == "report_created" and e["detail"].get("multistate", {}).get("type") == 1 for e in app.state.audit.list(100))
