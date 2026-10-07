"""Draft 8.5: multi-stage machining (8.5.1) and GD&T with MMR/LMR (8.5.3)."""

import numpy as np
import pytest
from scipy import stats

from spc.core import gdt
from spc.core import multistage as ms
from tests.conftest import logged_in_client, make_app
from tests.test_api import csv_text, upload


@pytest.fixture
def client():
    return logged_in_client(make_app(max_upload=200_000))

# ------------------------------------------------------------------ 8.5.1 scope of inspection


def test_scope_reproduces_the_draft_example_2():
    # 3 identical machines and 7 pallets, one part per pallet: the draft counts 3 * 7 * 50 = 1050 parts for a complete acceptance
    s = ms.inspection_scope(1, 7, 1, 3)
    assert s["combinations"] == 21 and s["full_acceptance_parts"] == 1050
    assert s["parts_per_combination"] == 5 and s["total_parts"] == 105 and s["saving"] == 945
    assert s["minimum_total_binding"] is False


def test_scope_minimum_total_of_50_raises_the_parts_per_combination():
    s = ms.inspection_scope(1, 2, 2, 1)  # 4 combinations: 5 each gives 20, the total of 50 needs 13 each
    assert s["combinations"] == 4 and s["parts_per_combination"] == 13 and s["total_parts"] == 52
    assert s["minimum_total_binding"] is True
    one = ms.inspection_scope()  # a single combination
    assert one["parts_per_combination"] == 50 and one["total_parts"] == 50


def test_scope_identical_combinations_and_fewer_carriers():
    base = ms.inspection_scope(1, 7, 1, 3)
    waived = ms.inspection_scope(1, 7, 1, 3, geometrically_identical=6)
    assert waived["inspected_combinations"] == 15 and waived["total_parts"] == 75
    fewer = ms.inspection_scope(2, 7, 1, 3, measured_carriers=3)
    assert fewer["combinations"] == 42 and fewer["combinations_in_study"] == 18
    assert fewer["parts_in_study"] == 90 and fewer["parts_other_carriers"] == (7 - 3) * 2 * 5
    assert fewer["total_parts"] == 90 + 40 and base["total_parts"] == 105
    for bad in (dict(carriers=0), dict(machines=-1), dict(geometrically_identical=21), dict(measured_carriers=8)):
        with pytest.raises(ValueError):
            ms.inspection_scope(**{"components_per_carrier": 1, "carriers": 7, "spindles": 1, "machines": 3, **bad})


def test_form_parts_are_spread_evenly_over_the_spindles():
    d = ms.spindle_distribution(50, 3)
    assert d["per_spindle"] == [17, 17, 16] and sum(d["per_spindle"]) == 50


# ------------------------------------------------------------------ 8.5.1 analysis of the combinations


def shifted(seed=3, shift=0.25):
    rng = np.random.default_rng(seed)
    x = rng.normal(10, 0.1, 120)
    pallet = np.tile(["P1", "P2", "P3"], 40)
    spindle = np.repeat(["S1", "S2"], 60)
    x[pallet == "P3"] += shift
    return x, pallet, spindle


def test_a_deviating_pallet_is_found_and_the_numbers_match_independent_formulas():
    x, pallet, spindle = shifted()
    r = ms.analyse_combinations(x, {"pallet": pallet, "spindle": spindle}, 9.5, 10.5)
    assert r["n"] == 120 and r["mean"] == pytest.approx(x.mean()) and r["sd"] == pytest.approx(x.std(ddof=1))
    assert r["indices"]["pk"] == pytest.approx(min(10.5 - x.mean(), x.mean() - 9.5) / (3 * x.std(ddof=1)))
    groups = [x[(pallet == p) & (spindle == s)] for p in ("P1", "P2", "P3") for s in ("S1", "S2")]
    f, p = stats.f_oneway(*groups)
    assert r["between"]["anova"]["f"] == pytest.approx(f) and r["between"]["anova"]["p_value"] == pytest.approx(p)
    sst = ((x - x.mean()) ** 2).sum()
    ssb = sum(g.size * (g.mean() - x.mean()) ** 2 for g in groups)
    assert r["between"]["eta_squared"] == pytest.approx(ssb / sst)
    swithin = sum(((g - g.mean()) ** 2).sum() for g in groups)
    assert r["between"]["pooled_sd"] == pytest.approx(np.sqrt(swithin / (120 - 6)))
    lev = stats.levene(*groups, center="median")
    assert r["between"]["variances"]["p_value"] == pytest.approx(lev.pvalue)
    by = {(c["labels"]["pallet"], c["labels"]["spindle"]): c for c in r["combinations"]}
    g = x[(pallet == "P3") & (spindle == "S1")]
    rest = np.delete(x, np.where((pallet == "P3") & (spindle == "S1"))[0])
    assert by[("P3", "S1")]["p_value"] == pytest.approx(stats.ttest_ind(g, rest, equal_var=False).pvalue)
    assert all(by[("P3", s)]["different"] for s in ("S1", "S2"))
    assert r["bonferroni_alpha"] == pytest.approx(0.05 / 6)
    pal = next(f for f in r["by_factor"] if f["factor"] == "pallet")
    assert pal["different"] and pal["balanced"]
    assert next(f for f in r["by_factor"] if f["factor"] == "spindle")["different"] is False


def test_identical_combinations_are_not_flagged_and_the_coverage_is_complete():
    x, pallet, spindle = shifted(shift=0.0, seed=11)
    r = ms.analyse_combinations(x, {"pallet": pallet, "spindle": spindle}, 9.5, 10.5)
    assert r["between"]["anova"]["different"] is False
    assert r["coverage"]["complete"] and r["coverage"]["missing_count"] == 0 and r["coverage"]["expected_combinations"] == 6


def test_missing_and_short_combinations_and_too_few_parts_are_reported():
    x = np.array([10.0, 10.1, 9.9, 10.05, 10.02, 9.98, 10.01, 10.03, 9.97, 10.0, 10.2, 9.8])
    pallet = ["A"] * 5 + ["B"] * 5 + ["C"] * 2
    machine = ["m1"] * 5 + ["m2"] * 5 + ["m1"] * 2
    r = ms.analyse_combinations(x, {"pallet": pallet, "machine": machine}, expected={"machine": ["m1", "m2"]})
    cov = r["coverage"]
    assert cov["expected_combinations"] == 6 and cov["combinations"] == 3 and cov["missing_count"] == 3
    assert cov["short_count"] == 1 and cov["short"] == [{"pallet": "C", "machine": "m1"}]
    assert cov["enough_total"] is False and cov["complete"] is False
    assert r["indices"] is None


def test_analysis_refuses_bad_input():
    with pytest.raises(ValueError):
        ms.analyse_combinations([1, 2], {"a": ["x", "y"]})
    with pytest.raises(ValueError):
        ms.analyse_combinations([1, 2, 3, 4], {})
    with pytest.raises(ValueError):
        ms.analyse_combinations([1, 2, 3, 4], {"a": ["x", "y"]})
    with pytest.raises(ValueError):
        ms.analyse_combinations([1, 2, 3, 4], {"a": list("xyxy")}, lsl=5, usl=1)


# ------------------------------------------------------------------ 8.5.3 GD&T

BORE = dict(lower=20.0, upper=20.2, position_tolerance=0.2)
BORE_ARGS = (20.0, 20.2, 0.2)


def test_virtual_sizes_of_the_draft_example_and_the_other_cases():
    assert gdt.virtual_size("bore", "mmc", 20.0, 20.2, 0.2) == (20.0, pytest.approx(19.8))  # MMS = LD, MMVS = LD - TP
    assert gdt.virtual_size("bore", "lmc", 20.0, 20.2, 0.2) == (20.2, pytest.approx(20.4))
    assert gdt.virtual_size("pin", "mmc", 19.8, 20.0, 0.1) == (20.0, pytest.approx(20.1))
    assert gdt.virtual_size("pin", "lmc", 19.8, 20.0, 0.1) == (19.8, pytest.approx(19.7))
    with pytest.raises(ValueError):
        gdt.virtual_size("bore", "mmc", 20.2, 20.0, 0.2)
    with pytest.raises(ValueError):
        gdt.virtual_size("bore", "mmc", 20.0, 20.2, 0)


def test_position_deviation_is_the_diameter_of_the_circle():
    assert gdt.position_deviation(0.03, 0.04) == pytest.approx(0.1)  # 2 * sqrt(0.03^2 + 0.04^2)


@pytest.mark.parametrize("kind,req,lower,upper,tp", [("bore", "mmc", 20.0, 20.2, 0.2), ("bore", "lmc", 20.0, 20.2, 0.2),
                                                     ("pin", "mmc", 19.8, 20.0, 0.1), ("pin", "lmc", 19.8, 20.0, 0.1)])
def test_clearance_is_zero_exactly_at_the_bonus_boundary(kind, req, lower, upper, tp):
    # the position tolerance grows by the departure of the size from the material limit; C = 0 where x_P equals that allowance
    for xd in np.linspace(lower, upper, 6):
        limit = lower if (kind, req) in (("bore", "mmc"), ("pin", "lmc")) else upper
        bonus = abs(xd - limit)
        allowed = tp + bonus
        assert gdt.clearance(xd, allowed, kind, req, lower, upper, tp) == pytest.approx(0.0, abs=1e-12)
        assert gdt.clearance(xd, allowed - 0.01, kind, req, lower, upper, tp) > 0
        assert gdt.clearance(xd, allowed + 0.01, kind, req, lower, upper, tp) < 0


def test_the_sign_printed_in_the_draft_would_reward_a_position_error():
    # draft example: a bore at the minimum size 20.0 with the full position deviation 0.2 just assembles (C = 0)
    assert gdt.clearance(20.0, 0.2, "bore", "mmc", *BORE_ARGS) == pytest.approx(0.0, abs=1e-12)
    assert gdt.clearance(20.0, 0.0, "bore", "mmc", *BORE_ARGS) == pytest.approx(0.2)  # a perfect position leaves 0.2
    # the printed "+ x_P" would give 20.0 + 0.2 - 19.8 = 0.4 for the bad hole: better than the perfect one (0.2)
    assert 20.0 + 0.2 - 19.8 > 20.0 + 0.0 - 19.8


def parts(n=300, seed=5, mean_size=20.1, sd_size=0.03, sd_pos=0.03):
    rng = np.random.default_rng(seed)
    return rng.normal(mean_size, sd_size, n), rng.normal(0, sd_pos, n), rng.normal(0, sd_pos, n)


def test_clearance_capability_of_a_normal_clearance_matches_the_definition():
    xd, dx, dy = parts()
    r = gdt.analyse_clearance(xd, dx=dx, dy=dy, distribution="normal", bootstrap_n=0, **BORE)
    c = xd - gdt.position_deviation(dx, dy) - 19.8
    assert r["n"] == 300 and r["clearance"] == pytest.approx(list(c))
    mu, sd = c.mean(), c.std()
    q = stats.norm.ppf([0.00135, 0.5, 0.99865], mu, sd)
    assert r["quantiles"]["c_50"] == pytest.approx(q[1], rel=1e-9) and r["quantiles"]["c_0135"] == pytest.approx(q[0], rel=1e-6)
    assert r["indices"]["pk"] == pytest.approx((q[1] - 0) / (q[1] - q[0]), rel=1e-6)
    z = gdt.analyse_clearance(xd, dx=dx, dy=dy, distribution="normal", method="Z", bootstrap_n=0, **BORE)
    assert z["indices"]["pk"] == pytest.approx(mu / sd / 3, rel=1e-6)
    assert r["observed"]["negative"] == int(np.sum(c < 0)) and r["sign_note"] == "draft_prints_plus_for_position"


def test_a_part_that_does_not_assemble_is_counted_and_a_bad_process_has_a_low_index():
    xd, dx, dy = parts(mean_size=20.02, sd_pos=0.06)
    r = gdt.analyse_clearance(xd, dx=dx, dy=dy, bootstrap_n=0, **BORE)
    assert r["observed"]["negative"] > 0 and r["indices"]["pk"] < 1.0 and r["indices"]["ppm"] > 0


def test_auto_distribution_is_chosen_for_the_folded_position_and_the_bootstrap_is_repeatable():
    xd, dx, dy = parts(n=200)
    a = gdt.analyse_clearance(xd, dx=dx, dy=dy, bootstrap_n=40, **BORE)
    b = gdt.analyse_clearance(xd, dx=dx, dy=dy, bootstrap_n=40, **BORE)
    assert a["indices"]["ci_pk"] == b["indices"]["ci_pk"] and a["indices"]["ci_pk"] is not None
    assert a["distribution"]["name"] in ("normal", "weibull", "gamma", "lognormal", "johnson_su", "box_cox", "mixture")
    xp = gdt.position_deviation(dx, dy)
    c = gdt.analyse_clearance(xd, xp=xp, bootstrap_n=0, **BORE)  # xp given directly gives the same clearances
    assert c["clearance"] == pytest.approx(a["clearance"])


def test_gdt_input_is_checked():
    xd, dx, dy = parts(n=20)
    with pytest.raises(ValueError):
        gdt.analyse_clearance(xd, **BORE)  # neither xp nor dx, dy
    with pytest.raises(ValueError):
        gdt.analyse_clearance(xd, xp=-np.ones(20), **BORE)
    with pytest.raises(ValueError):
        gdt.analyse_clearance(xd, xp=np.zeros(19), **BORE)
    with pytest.raises(ValueError):
        gdt.analyse_clearance(xd[:4], dx=dx[:4], dy=dy[:4], **BORE)
    with pytest.raises(ValueError):
        gdt.analyse_clearance(xd, dx=dx, dy=dy, distribution="empirical", method="Z", **BORE)
    small = gdt.analyse_clearance(xd, dx=dx, dy=dy, bootstrap_n=0, distribution="normal", **BORE)
    assert any(w["code"] == "gdt_small_sample" for w in small["warnings"])


# ------------------------------------------------------------------ API


def test_scope_api(client):
    r = client.post("/api/multistage/scope", json={"carriers": 7, "machines": 3}).json()
    assert r["combinations"] == 21 and r["full_acceptance_parts"] == 1050 and r["total_parts"] == 105
    assert sum(r["spindle_form"]["per_spindle"]) == 50
    bad = client.post("/api/multistage/scope", json={"carriers": 2, "geometrically_identical": 5})
    assert bad.status_code == 400 and bad.json()["error"]["code"] == "invalid_input"
    assert client.post("/api/multistage/scope", json={"carriers": 0}).status_code == 422


def test_multistage_on_a_stored_dataset(client):
    ds = upload(client, csv_text(k=20, n=5, seed=2)).json()
    r = client.post(f"/api/datasets/{ds['id']}/multistage", json={"factors": ["machine"], "lsl": 9.5, "usl": 10.5})
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["n"] == 100 and [c["labels"]["machine"] for c in body["combinations"]] == ["M1", "M2"]
    assert body["coverage"]["expected_combinations"] == 2 and body["indices"]["pk"] > 0
    sub = client.post(f"/api/datasets/{ds['id']}/multistage", json={"factors": ["machine", "subgroup"]}).json()
    assert sub["coverage"]["expected_combinations"] == 40
    missing = client.post(f"/api/datasets/{ds['id']}/multistage", json={"factors": ["nope"]})
    assert missing.status_code == 400 and "nope" in missing.json()["error"]["message"]
    assert client.post("/api/datasets/zzz/multistage", json={"factors": ["machine"]}).status_code == 404


def test_gdt_api(client):
    xd, dx, dy = parts(n=120)
    body = {"xd": xd.tolist(), "dx": dx.tolist(), "dy": dy.tolist(), "lower": 20.0, "upper": 20.2, "position_tolerance": 0.2, "bootstrap_n": 20}
    r = client.post("/api/gdt/clearance", json=body)
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["virtual_size"] == pytest.approx(19.8) and out["n"] == 120 and out["indices"]["pk"] > 0
    assert client.post("/api/gdt/clearance", json={**body, "dx": None, "dy": None}).json()["error"]["code"] == "invalid_input"
    assert client.post("/api/gdt/clearance", json={**body, "kind": "plug"}).status_code == 422


# ------------------------------------------------------------------ 8.5.2 multivariate Pm and Pmk
from spc.core import multivariate_perf as mp  # noqa: E402
from spc.core import nested as nest  # noqa: E402
from spc.core.charts import trend as trend_mod  # noqa: E402


def brute_distance(mean, cov, lower, upper, n=400_000, seed=0):
    """Smallest Mahalanobis distance from the mean to points of the tolerance ellipsoid's border, found by sampling the border."""
    d = len(mean)
    m, a = (np.array(lower) + upper) / 2, (np.array(upper) - lower) / 2
    y = np.random.default_rng(seed).normal(size=(n, d))
    y /= np.linalg.norm(y, axis=1)[:, None]
    dx = m + a * y - mean
    return float(np.sqrt(np.min(np.einsum("ij,jk,ik->i", dx, np.linalg.inv(cov), dx))))


def test_the_one_dimensional_case_gives_the_univariate_indices():
    # d = 1 is the check of the reading u_p = Phi^-1((1 + p) / 2): Pmk = min(U - mu, mu - L) / (3 sigma), Pm = (U - L) / (6 sigma)
    for mu in (10.0, 10.2, 9.7, 12.0):
        c, inside = mp.distance_to_tolerance([mu], [[0.0225]], [9.0], [11.0])
        assert c == pytest.approx(min(11 - mu, mu - 9) / 0.15 if 9 <= mu <= 11 else (mu - 11) / 0.15) and inside == (9 <= mu <= 11)
        p, u = mp._u_of(c, 1)
        assert u == pytest.approx(c, rel=1e-9)
    c, _ = mp.distance_to_tolerance([10.0], [[0.0225]], [9.0], [11.0])
    assert mp._u_of(c, 1)[0] == pytest.approx(2 * stats.norm.cdf(c) - 1)


def test_distance_to_the_tolerance_border_matches_a_brute_force_search():
    rng = np.random.default_rng(2)
    for d in (2, 3):
        for _ in range(4):
            M = rng.normal(size=(d, d))
            cov = M @ M.T + 0.3 * np.eye(d)
            lo, up = -rng.uniform(1, 3, d), rng.uniform(1, 3, d)
            for mean in ((lo + up) / 2, (lo + up) / 2 + rng.normal(0, 0.5, d), (lo + up) / 2 + rng.normal(0, 4, d)):
                c, inside = mp.distance_to_tolerance(mean, cov, lo, up)
                b = brute_distance(mean, cov, lo, up)
                assert c <= b + 1e-9 and c == pytest.approx(b, abs=3e-3)
                assert inside == bool(np.sum(((mean - (lo + up) / 2) / ((up - lo) / 2)) ** 2) <= 1)


def test_a_circular_tolerance_and_uncorrelated_data_have_a_closed_form():
    # identity covariance, centred, tolerance circle of radius a: c = a, p = 1 - exp(-a^2 / 2) (chi-square with 2 d.f.)
    c, _ = mp.distance_to_tolerance([0, 0], np.eye(2), [-3, -3], [3, 3])
    assert c == pytest.approx(3.0)
    p, u = mp._u_of(3.0, 2)
    assert p == pytest.approx(1 - np.exp(-4.5)) and u == pytest.approx(stats.norm.isf(np.exp(-4.5) / 2))


def centred_sample(n=400, seed=7, scale=(0.1, 0.2), rho=0.5):
    rng = np.random.default_rng(seed)
    z = rng.multivariate_normal([0, 0], [[1, rho], [rho, 1]], size=n)
    z -= z.mean(axis=0)
    return z * np.array(scale)


def test_pm_pmk_of_a_centred_sample_are_equal_and_a_shifted_one_has_lower_pmk():
    x = centred_sample() + [10.0, 5.0]
    r = mp.performance(x, [9.0, 4.0], [11.0, 6.0])
    assert r["case"] == 1 and r["pm"] == pytest.approx(r["pmk"], rel=1e-9) and r["pm"] > 1
    shifted = mp.performance(x + [0.3, 0.0], [9.0, 4.0], [11.0, 6.0])
    assert shifted["pm"] == pytest.approx(r["pm"], rel=1e-9) and shifted["pmk"] < r["pmk"]  # Pm ignores the location
    outside = mp.performance(x + [5.0, 0.0], [9.0, 4.0], [11.0, 6.0])
    assert outside["case"] == 2 and outside["pmk"] < 0 and outside["centre_inside"] is False


def test_performance_does_not_change_when_every_axis_is_rescaled():
    x = centred_sample() + [10.0, 5.0] + [0.2, -0.1]
    a = mp.performance(x, [9.0, 4.0], [11.0, 6.0])
    f = np.array([3.0, 0.25])
    b = mp.performance(x * f, np.array([9.0, 4.0]) * f, np.array([11.0, 6.0]) * f)
    assert b["pm"] == pytest.approx(a["pm"], rel=1e-9) and b["pmk"] == pytest.approx(a["pmk"], rel=1e-9)


def test_a_larger_spread_lowers_pm_and_the_sample_check_reports_the_normality_tests():
    tight = mp.performance(centred_sample(scale=(0.1, 0.2)) + 5, [4, 4], [6, 6])
    wide = mp.performance(centred_sample(scale=(0.3, 0.6)) + 5, [4, 4], [6, 6])
    assert wide["pm"] < tight["pm"]
    assert 0 <= tight["normality"]["skewness_p"] <= 1 and tight["univariate"][0]["name"] == "x1"
    assert tight["expected_outside"] is not None and wide["expected_outside"] > tight["expected_outside"]


def test_multivariate_input_is_checked():
    x = centred_sample(n=30) + 5
    with pytest.raises(ValueError):
        mp.performance(x, [4, 4], [6])  # one limit missing
    with pytest.raises(ValueError):
        mp.performance(x, [6, 4], [4, 6])
    with pytest.raises(ValueError):
        mp.performance(x[:5], [4, 4], [6, 6])
    with pytest.raises(ValueError):
        mp.performance(np.column_stack([x[:, 0], x[:, 0] * 2]), [4, 4], [6, 6])  # singular covariance
    with pytest.raises(ValueError):
        mp.performance(x[:, :1], [4], [6])


# ------------------------------------------------------------------ nested sources of variation


def nested_data(a=6, b=4, c=3, s=(2.0, 1.0, 0.5), seed=1):
    rng = np.random.default_rng(seed)
    lot, part, x = [], [], []
    for i in range(a):
        ea = rng.normal(0, s[0])
        for j in range(b):
            eb = rng.normal(0, s[1])
            for _ in range(c):
                lot.append(f"L{i}")
                part.append(f"P{j}")
                x.append(10 + ea + eb + rng.normal(0, s[2]))
    return np.array(x), lot, part


def test_balanced_nested_anova_equals_the_textbook_formulas():
    x, lot, part = nested_data()
    r = nest.analyse(x, {"lot": lot, "part": part})
    a, b, c = 6, 4, 3
    cells = x.reshape(a, b, c)
    ms_e = ((cells - cells.mean(axis=2, keepdims=True)) ** 2).sum() / (a * b * (c - 1))
    ms_b = c * ((cells.mean(axis=2) - cells.mean(axis=(1, 2))[:, None]) ** 2).sum() / (a * (b - 1))
    ms_a = b * c * ((cells.mean(axis=(1, 2)) - cells.mean()) ** 2).sum() / (a - 1)
    assert r["balanced"] and r["error"]["variance"] == pytest.approx(ms_e)
    by = {t["level"]: t for t in r["table"]}
    assert by["part"]["variance"] == pytest.approx((ms_b - ms_e) / c) and by["lot"]["variance"] == pytest.approx((ms_a - ms_b) / (b * c))
    assert by["lot"]["coefficients"]["lot"] == pytest.approx(b * c) and by["lot"]["coefficients"]["part"] == pytest.approx(c)
    assert by["lot"]["df"] == a - 1 and by["part"]["df"] == a * (b - 1) and r["error"]["df"] == a * b * (c - 1)
    assert by["lot"]["f"] == pytest.approx(ms_a / ms_b) and by["lot"]["p_value"] == pytest.approx(stats.f.sf(ms_a / ms_b, a - 1, a * (b - 1)))
    assert sum(t["share"] for t in r["table"]) + r["error"]["share"] == pytest.approx(1.0)
    assert r["largest"] == "lot"


def test_unbalanced_estimators_are_unbiased_for_the_known_components():
    # the method of moments is unbiased before truncation: the average of many unbalanced experiments hits the true variances
    true = np.array([4.0, 1.0, 0.25])
    rng = np.random.default_rng(10)
    est = []
    for _ in range(600):
        lot, part, x = [], [], []
        for i in range(5):
            ea = rng.normal(0, 2.0)
            for j in range(int(rng.integers(2, 5))):
                eb = rng.normal(0, 1.0)
                for _k in range(int(rng.integers(2, 5))):
                    lot.append(f"L{i}")
                    part.append(f"P{j}")
                    x.append(ea + eb + rng.normal(0, 0.5))
        r = nest.analyse(x, {"lot": lot, "part": part})
        est.append([r["table"][0]["variance_raw"], r["table"][1]["variance_raw"], r["error"]["variance"]])
        assert not r["balanced"] and r["f_exact"] is False
    assert np.mean(est, axis=0) == pytest.approx(true, rel=0.12)


def test_nested_groups_are_built_from_the_path_and_reused_labels_are_separate():
    x, lot, part = nested_data(a=3, b=2, c=4)
    same = nest.analyse(x, {"lot": lot, "part": part})
    assert same["table"][1]["groups"] == 6  # P0 and P1 inside each of the 3 lots
    one = nest.analyse(x, {"part": part, "lot": lot})  # the other order is a different model: P is outer, lot inside
    assert one["table"][0]["groups"] == 2 and one["table"][1]["groups"] == 6


def test_negative_estimates_are_truncated_and_marked_and_bad_designs_are_refused():
    rng = np.random.default_rng(3)
    x = rng.normal(0, 1, 24)
    r = nest.analyse(x, {"g": list("ab" * 12), "h": [str(i % 6) for i in range(24)]})
    assert any(t["truncated"] for t in r["table"]) or r["table"][0]["variance"] >= 0
    for t in r["table"]:
        assert t["variance"] >= 0
    with pytest.raises(ValueError):
        nest.analyse(x[:3], {"g": list("abc")})
    with pytest.raises(ValueError):
        nest.analyse(x, {"g": list("a" * 24)})  # one outer group
    with pytest.raises(ValueError):
        nest.analyse(x, {"g": list("ab" * 12), "h": [str(i) for i in range(24)]})  # every value its own group
    with pytest.raises(ValueError):
        nest.analyse(x, {})


# ------------------------------------------------------------------ the regression control chart for trends


def wear(k=60, cycle=20, slope=0.02, sd=0.05, seed=9):
    rng = np.random.default_rng(seed)
    t = np.arange(k) % cycle
    return 10 + slope * t + rng.normal(0, sd, k)


def test_trend_chart_matches_an_independent_regression():
    y = wear()
    t = np.arange(60) % 20
    ch = trend_mod.fit(y, cycle=20)
    ref = stats.linregress(t, y)
    assert ch.slope == pytest.approx(ref.slope) and ch.intercept == pytest.approx(ref.intercept) and ch.slope_p == pytest.approx(ref.pvalue)
    assert ch.r2 == pytest.approx(ref.rvalue ** 2)
    res = y - (ref.intercept + ref.slope * t)
    assert ch.sigma == pytest.approx(np.sqrt(np.sum(res ** 2) / 58)) and ch.df == 58
    u = stats.norm.isf(0.00135)
    assert ch.ucl == pytest.approx(ch.center + u * ch.sigma) and ch.lcl == pytest.approx(ch.center - u * ch.sigma)
    lo, hi = ref.slope - stats.t.isf(0.025, 58) * ref.stderr, ref.slope + stats.t.isf(0.025, 58) * ref.stderr
    assert ch.slope_ci == pytest.approx((lo, hi)) and not ch.violations


def test_a_process_with_a_trend_is_not_alarmed_by_its_trend_but_by_a_jump_from_the_line():
    y = wear()
    flat = trend_mod.fit(np.arange(60) * 0.0 + y)  # without a cycle the saw tooth is a poor fit
    assert flat.r2 < 0.3
    y2 = y.copy()
    y2[45] += 0.6  # a jump away from the line
    ch = trend_mod.fit(y2, cycle=20)
    assert [v.index for v in ch.violations if v.rule == "beyond_limits"] == [45]
    shift = y.copy()
    shift[30:] += 0.12  # a sustained shift above the line: a run on one side
    runs = trend_mod.fit(shift, cycle=20).violations
    assert any(v.rule == "run" for v in runs)


def test_trend_chart_input_is_checked():
    with pytest.raises(ValueError):
        trend_mod.fit([1, 2, 3])
    with pytest.raises(ValueError):
        trend_mod.fit(wear(), cycle=2)
    with pytest.raises(ValueError):
        trend_mod.fit(wear(k=10), cycle=20)
    with pytest.raises(ValueError):
        trend_mod.fit(np.arange(20.0))  # exactly on a line


# ------------------------------------------------------------------ API of the three


def test_multivariate_api(client):
    x = centred_sample(n=60) + [10.0, 5.0]
    body = {"data": x.tolist(), "lower": [9, 4], "upper": [11, 6], "names": ["x", "y"]}
    r = client.post("/api/multivariate/performance", json=body)
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["names"] == ["x", "y"] and out["pm"] == pytest.approx(out["pmk"], rel=1e-6) and out["case"] == 1
    assert client.post("/api/multivariate/performance", json={**body, "upper": [11]}).status_code in (400, 422)
    assert client.post("/api/multivariate/performance", json={**body, "lower": [11, 4]}).json()["error"]["code"] == "invalid_input"


def lots_csv():
    x, lot, part = nested_data(a=5, b=3, c=4, seed=4)
    lines = ["lot,part,value"] + [f"{lo},{pa},{v:.4f}" for lo, pa, v in zip(lot, part, x)]
    return "\n".join(lines).encode()


def test_nested_api(client):
    ds = upload(client, lots_csv(), value="value", subgroup="", tags=["lot", "part"]).json()
    r = client.post(f"/api/datasets/{ds['id']}/nested", json={"levels": ["lot", "part"]})
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["n"] == 60 and out["balanced"] and [t["level"] for t in out["table"]] == ["lot", "part"]
    assert client.post(f"/api/datasets/{ds['id']}/nested", json={"levels": ["nope"]}).status_code == 400


def test_trend_api(client):
    y = wear(k=60)
    lines = ["lot,value"] + [f"S{i // 4}, {v:.4f}" for i, v in enumerate(y)]
    # 15 subgroups of 4, and a cycle of 5 subgroups
    ds = upload(client, "\n".join(lines).encode(), value="value", subgroup="lot", tags=[]).json()
    r = client.post(f"/api/datasets/{ds['id']}/trend-chart", json={"cycle": 5, "lsl": 9.5, "usl": 11.0})
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["n_points"] == 15 and out["subgroup_size"] == 4 and out["cycle"] == 5 and len(out["center"]) == 15
    assert out["indices"]["pk"] > 0 and out["indices"]["distribution"]
    bad = client.post(f"/api/datasets/{ds['id']}/trend-chart", json={"cycle": 99})
    assert bad.status_code == 400


# ------------------------------------------------------------------ annex E: the results in a report and its archive
import json  # noqa: E402

from spc.report.archive import reproduce  # noqa: E402
from tests.test_api import REPORT_BODY  # noqa: E402


def special_request(seed=1):
    rng = np.random.default_rng(seed)
    return {
        "scope": {"carriers": 7, "machines": 3},
        "multistage": {"factors": ["machine"], "lsl": 9.5, "usl": 10.5},
        "nested": {"levels": ["machine", "subgroup"]},
        "trend": {"cycle": 5, "lsl": 9.5, "usl": 10.5},
        "gdt": {"xd": rng.normal(20.1, 0.03, 60).tolist(), "dx": rng.normal(0, 0.03, 60).tolist(), "dy": rng.normal(0, 0.03, 60).tolist(),
                "lower": 20.0, "upper": 20.2, "position_tolerance": 0.2, "bootstrap_n": 20},
        "multivariate": {"data": rng.multivariate_normal([10, 5], [[0.01, 0.004], [0.004, 0.02]], 60).tolist(), "lower": [9, 4], "upper": [11, 6]},
    }


def make_report(client, special, language="en"):
    ds = upload(client, csv_text(k=20, n=5)).json()
    r = client.post(f"/api/datasets/{ds['id']}/reports", json={**REPORT_BODY, "language": language, "special": special})
    assert r.status_code == 200, r.text
    return r.json()


def test_special_results_in_the_report_the_archive_and_the_excel_file(client):
    out = make_report(client, special_request())
    html = client.get(out["urls"]["html"]).text
    for text in ("Annex E", "Multi-stage machining: scope of inspection", "nested variance components", "regression control chart", "assembly clearance", "Pm and Pmk"):
        assert text in html
    assert "21 combinations" in html and "1050" in html
    archive = client.get(out["urls"]["archive"]).json()
    assert set(archive["special"]) == {"scope", "multistage", "nested", "trend", "gdt", "multivariate"}
    assert "clearance" not in archive["special"]["gdt"]["result"]  # the per-part arrays are made again from the request
    check = client.post("/api/archive/check", content=client.get(out["urls"]["archive"]).content).json()
    assert check == {"integrity_ok": True, "reproduced": True, "same_engine_version": True, "differences": []}
    xlsx = client.get(f"/api/reports/{out['id']}/report.xlsx")
    assert xlsx.status_code == 200
    from io import BytesIO

    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(xlsx.content))
    cells = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None)
    assert "Annex E" in cells and "assembly clearance" in cells.lower()


def test_report_in_chinese_has_the_annex(client):
    out = make_report(client, {"scope": {"carriers": 7, "machines": 3}}, "zh-TW")
    assert "附錄 E" in client.get(out["urls"]["html"]).text and "多段加工：檢驗範圍" in client.get(out["urls"]["html"]).text


def test_a_changed_special_result_is_found_by_the_digest_and_by_a_new_calculation(client):
    out = make_report(client, special_request())
    archive = client.get(out["urls"]["archive"]).json()
    forged = json.loads(json.dumps(archive))
    forged["special"]["gdt"]["result"]["indices"]["pk"] += 1.0
    assert client.post("/api/archive/check", json=forged).json()["integrity_ok"] is False
    # a forger who also repairs the digest is still caught: the stored request gives another result
    from spc.report.archive import canonical
    import hashlib

    body = {k: v for k, v in forged.items() if k != "integrity"}
    forged["integrity"] = {"algorithm": "sha256", "digest": hashlib.sha256(canonical(body)).hexdigest()}
    res = client.post("/api/archive/check", json=forged).json()
    assert res["integrity_ok"] is True and res["reproduced"] is False and any("special.gdt" in d for d in res["differences"])
    assert reproduce(archive).reproduced


def test_special_report_input_is_checked(client):
    ds = upload(client).json()
    bad = client.post(f"/api/datasets/{ds['id']}/reports", json={**REPORT_BODY, "special": {"multistage": {"factors": ["nope"]}}})
    assert bad.status_code == 400 and bad.json()["error"]["code"] == "invalid_input"
    assert client.post(f"/api/datasets/{ds['id']}/reports", json={**REPORT_BODY, "special": {"gdt": {"xd": [1, 2]}}}).status_code == 422


# ------------------------------------------------------------------ several cavities, stations or clamping devices (draft 8.2.6)
from spc.service.special import cavity_study  # noqa: E402
from spc.core.capability.indices import overall_indices  # noqa: E402
from spc.data import Dataset  # noqa: E402


def cavity_dataset(shift=0.0, seed=61, per=40):
    rng = np.random.default_rng(seed)
    labels = np.repeat(["C1", "C2", "C3", "C4"], per)
    x = rng.normal(10, 0.05, labels.size)
    x[labels == "C3"] += shift
    return Dataset.from_values(x, tags={"cavity": labels})


def test_each_cavity_is_a_machine_of_its_own_and_the_variation_is_split():
    ds = cavity_dataset(shift=0.15)
    r = cavity_study(ds, "cavity", 9.7, 10.3)
    x, labels = ds.values, ds.tags["cavity"]
    for c in r["cavities"]:
        v = x[labels == c["label"]]
        idx = overall_indices(v, 9.7, 10.3)
        assert c["n"] == 40 and c["pm"] == pytest.approx(idx.p) and c["pmk"] == pytest.approx(idx.pk) and c["mean"] == pytest.approx(v.mean())
    assert r["whole"]["pmk"] == pytest.approx(overall_indices(x, 9.7, 10.3).pk) and r["whole"]["n"] == 160
    shifted = next(c for c in r["cavities"] if c["label"] == "C3")
    assert shifted["pmk"] < min(c["pmk"] for c in r["cavities"] if c["label"] != "C3")  # the shifted cavity has the lowest Pmk
    t = r["variance"]["table"][0]
    assert t["level"] == "cavity" and t["significant"] and t["share"] > 0.5 and t["share"] + r["variance"]["error"]["share"] == pytest.approx(1.0)
    assert "C3" in [c["labels"]["cavity"] for c in r["combinations"]["combinations"] if c["different"]]
    even = cavity_study(cavity_dataset(), "cavity", 9.7, 10.3)
    assert not even["variance"]["table"][0]["significant"] and not any(c["different"] for c in even["combinations"]["combinations"])


def test_the_cavity_study_without_limits_and_with_a_small_cavity():
    ds = cavity_dataset()
    r = cavity_study(ds, "cavity", None, None)
    assert r["whole"] is None and all(c["pm"] is None and c["pmk"] is None for c in r["cavities"])
    tiny = Dataset.from_values([10.0, 10.1, 9.9, 10.05] + [10.0 + 0.01 * i for i in range(40)], tags={"cavity": ["A"] * 4 + ["B"] * 40})
    out = cavity_study(tiny, "cavity", 9.5, 10.5)
    assert out["cavities"][0]["pmk"] is None and out["cavities"][1]["pmk"] is not None  # four values are too few for an index
    with pytest.raises(ValueError):
        cavity_study(ds, "nope", 9.5, 10.5)


def test_cavities_api_report_annex_and_archive(client):
    rng = np.random.default_rng(62)
    lines = ["lot,value,cavity"] + [f"S{i // 4},{10 + rng.normal(0, 0.05) + (0.2 if i % 3 == 2 else 0):.4f},K{i % 3 + 1}" for i in range(120)]
    ds = upload(client, "\n".join(lines).encode(), value="value", subgroup="", tags=["cavity"]).json()
    r = client.post(f"/api/datasets/{ds['id']}/cavities", json={"factor": "cavity", "lsl": 9.7, "usl": 10.4})
    assert r.status_code == 200, r.text
    assert [c["label"] for c in r.json()["cavities"]] == ["K1", "K2", "K3"] and r.json()["variance"]["table"][0]["significant"]
    assert client.post(f"/api/datasets/{ds['id']}/cavities", json={"factor": "nope"}).status_code == 400
    out = client.post(f"/api/datasets/{ds['id']}/reports", json={**REPORT_BODY, "special": {"cavities": {"factor": "cavity", "lsl": 9.7, "usl": 10.4}}})
    assert out.status_code == 200, out.text
    html = client.get(out.json()["urls"]["html"]).text
    assert "several cavities, stations or clamping devices (draft 8.2.6)" in html and "K3" in html
    check = client.post("/api/archive/check", content=client.get(out.json()["urls"]["archive"]).content).json()
    assert check == {"integrity_ok": True, "reproduced": True, "same_engine_version": True, "differences": []}


# ------------------------------------------------------------------ the machine performance study named in a report (draft 9.3)

def test_a_report_can_name_a_closed_machine_study_and_refuses_an_open_one():
    from tests.test_study import RECORD, add_dataset, confirm_all, proven

    app = make_app()
    c = {u: logged_in_client(app, u) for u in ("admin", "eng", "view")}
    eng = c["eng"]
    sid = eng.post("/api/studies", json={"record": proven(c, {**RECORD, "dataset_id": add_dataset(app)})}).json()["study"]["id"]
    ds = upload(eng, csv_text(k=20, n=5)).json()
    body = {**REPORT_BODY, "machine_study_id": sid}
    refused = eng.post(f"/api/datasets/{ds['id']}/reports", json=body)
    assert refused.status_code == 409 and refused.json()["error"]["code"] == "machine_study_open"
    confirm_all(eng, sid)
    eng.post(f"/api/studies/{sid}/close", json={"reason": "all conditions met"})
    out = eng.post(f"/api/datasets/{ds['id']}/reports", json=body)
    assert out.status_code == 200, out.text
    html = eng.get(out.json()["urls"]["html"]).text
    assert "Machine performance study" in html and "closed" in html
    assert eng.post(f"/api/datasets/{ds['id']}/reports", json={**REPORT_BODY, "machine_study_id": 9999}).status_code == 404
    entry = [e for e in c["admin"].get("/api/audit?limit=20").json()["entries"] if e["action"] == "report_created"][0]
    assert entry["detail"]["machine_study"]["id"] == sid
