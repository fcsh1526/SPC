import io

import numpy as np
import pytest
from openpyxl import load_workbook

from spc.data import ColumnMap, Dataset, load_csv, to_csv
from spc.report import generate, report_xlsx
from spc.report.xlsx import XLSX_TYPE, dataset_xlsx
from spc.service import AnalysisRequest
from tests.conftest import logged_in_client, make_app
from tests.test_api import csv_text, err, upload

EVIL = '=HYPERLINK("http://example.com","click")'


def sample(n=125, seed=1, mark=True):
    rng = np.random.default_rng(seed)
    ds = Dataset.from_values(rng.normal(10, 0.1, n), subgroup=[f"L{i // 5 + 1}" for i in range(n)],
                             tags={"machine": [f"M{i % 2}" for i in range(n)]})
    return ds.mark_invalid([3], EVIL, "Eva (eng)") if mark else ds


def req(**kw):
    base = dict(stage="production", lsl=9.5, usl=10.5, model="A1", characteristic_class="major")
    base.update(kw)
    return AnalysisRequest(**base)


def workbook(ds=None, request=None, lang="en", meta=None):
    g = generate(ds or sample(), request or req(), meta, lang, now="2026-10-01T00:00:00+00:00", report_id="x1")
    return g, load_workbook(io.BytesIO(report_xlsx(g.archive)))


def cells(ws):
    return [c for row in ws.iter_rows() for c in row if c.value is not None]


def check_rows(ws):
    return {ws.cell(r, 1).value: r for r in range(11, 30) if ws.cell(r, 1).value}


# ------------------------------------------------------------------ structure and content

def test_sheets_are_named_in_the_language_of_the_report():
    _, en = workbook(lang="en")
    assert en.sheetnames == ["Summary", "Report elements", "Data", "Control charts", "Check", "Annex", "Log"]
    _, zh = workbook(lang="zh-TW")
    assert zh.sheetnames == ["摘要", "報告要素", "資料", "管制圖", "驗算", "附錄", "紀錄"]


def test_summary_holds_the_numbers_of_the_report():
    g, wb = workbook()
    ws = wb["Summary"]
    ix = g.archive["result"]["indices"]
    assert ws["A3"].value == "Report no." and ws["B3"].value == "x1"
    assert ws["B5"].value == g.archive["integrity"]["digest"]
    assert [ws["A8"].value, ws["A9"].value] == ["Cp.G", "Cpk.G"]
    assert ws["B8"].value == pytest.approx(ix["p"]) and ws["B9"].value == pytest.approx(ix["pk"])
    assert ws["C9"].value == pytest.approx(ix["ci_pk"][0]) and ws["D9"].value == pytest.approx(ix["ci_pk"][1])
    assert isinstance(ws["B9"].value, float)  # numbers, not text
    assert any("Cp.G" in str(c.value) and "meets" in str(c.value) for c in cells(ws))  # a conclusion sentence


def test_elements_sheet_lists_the_20_plus_2_elements():
    from spc.report.meta import ReportMeta

    meta = ReportMeta(process="turning", machine="CNC-07", part_name="shaft", unit="mm", recommendations="keep sampling",
                      uncertainty=0.02)
    g, wb = workbook(meta=meta)
    ws = wb["Report elements"]
    nos = {ws.cell(r, 1).value for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value is not None}
    assert nos == {1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 14, 15, 16, 17, 18, 19, 20, 21, 22}  # 11 to 13 are figures
    text = " ".join(str(c.value) for c in cells(ws))
    assert "turning" in text and "CNC-07" in text and "keep sampling" in text and "Guard band" in text
    values = {ws.cell(r, 3).value: ws.cell(r, 4).value for r in range(2, ws.max_row + 1)}
    assert values["Values used (neff)"] == g.archive["result"]["counts"]["n_used"]


def test_no_within_subgroup_index_anywhere():
    _, wb = workbook()
    text = " ".join(str(c.value) for ws in wb for c in cells(ws))
    assert "Cwk" not in text and " Cw " not in text


def test_everything_is_arial_and_numbers_are_numbers():
    _, wb = workbook()
    for ws in wb:
        for c in cells(ws):
            assert c.font.name == "Arial", (ws.title, c.coordinate)
    data = wb["Data"]
    assert isinstance(data["C2"].value, float) and isinstance(data["A2"].value, (int, float))


# ------------------------------------------------------------------ data sheet and safety

def test_data_sheet_has_every_value_with_status_and_marks():
    g, wb = workbook()
    ws = wb["Data"]
    assert ws.max_row == 126 and ws["C2"].value == pytest.approx(sample().values[0])
    used = sum(ws.cell(r, 4).value for r in range(2, 127))
    assert used == g.archive["result"]["counts"]["n_used"] == 120  # the subgroup of the marked value is left out
    assert ws["E5"].value == "marked invalid" and ws["D5"].value == 0
    assert ws["E6"].value == "not used"  # the other values of that subgroup
    head = [c.value for c in ws[1]]
    assert head[:5] == ["No.", "File row", "Value", "Used (1 = yes)", "Status"] and "machine" in head
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref.startswith("A1:")


def test_text_from_people_is_never_a_formula():
    _, wb = workbook()
    reason = next(c for c in cells(wb["Data"]) if c.value == EVIL)
    assert reason.data_type == "s"
    assert wb["Annex"]["C3"].value == EVIL and wb["Annex"]["C3"].data_type == "s"
    formulas = [(ws.title, c.coordinate) for ws in wb for c in cells(ws) if c.data_type == "f"]
    assert formulas and {title for title, _ in formulas} == {"Check"}  # formulas live on the Check sheet only


def test_other_user_text_fields_are_safe_too():
    from spc.report.meta import ReportMeta

    meta = ReportMeta(process="=1+1", machine="+SUM(A1)", persons="@cmd", recommendations="-2+3")
    _, wb = workbook(meta=meta)
    values = [c for c in cells(wb["Report elements"]) if str(c.value) in ("=1+1", "+SUM(A1)", "@cmd", "-2+3")]
    assert len(values) == 4 and all(c.data_type == "s" for c in values)


def test_csv_export_guards_formulas_and_import_takes_the_guard_off():
    ds = sample()
    text = to_csv(ds)
    assert f"'{EVIL}" in text.replace('""', '"') or "'=HYPERLINK" in text
    back = load_csv(text.encode("utf-8"), ColumnMap(value="value", subgroup="subgroup", valid="valid", invalid_reason="invalid_reason",
                                                    invalid_by="invalid_by", invalid_at="invalid_at", tags=("machine",)))
    assert next(iter(back.invalid_info().values()))[0] == EVIL
    plain = Dataset.from_values([1.0, -2.5, 3.0], subgroup=["a", "-1", "b"])
    assert "-1" in to_csv(plain) and "'-1" not in to_csv(plain)  # a plain number is no formula


# ------------------------------------------------------------------ charts

def test_control_charts_have_values_limits_and_native_charts():
    g, wb = workbook()
    ws = wb["Control charts"]
    part = g.archive["result"]["chart"]["location"]
    n = len(part["values"])
    assert ws["B3"].value == pytest.approx(part["values"][0]) and ws["C3"].value == pytest.approx(part["lcl"])
    assert ws["E3"].value == pytest.approx(part["ucl"]) and ws.max_row >= n + 2
    assert len(ws._charts) == 2
    assert [c.value for c in ws[2]][:5] == ["Label", "Value", "LCL", "CL", "UCL"]


def test_machine_study_has_no_control_chart_sheet():
    _, wb = workbook(request=req(stage="machine", model=None, characteristic_class=None))
    assert "Control charts" not in wb.sheetnames


def test_stepwise_limits_and_phases_are_written_per_point():
    rng = np.random.default_rng(4)
    ds = (Dataset.from_values(np.r_[rng.normal(10, 0.1, 40), rng.normal(10.6, 0.15, 40)])
          .add_restart([40], "new fixture", "A. Chen", new_limits=True))
    g, wb = workbook(ds, req(stage="preliminary", lsl=9.0, usl=12.0, model=None, characteristic_class=None, moving_n=3))
    ws = wb["Control charts"]
    assert ws["E3"].value != ws["E50"].value  # the limit changes with the phase
    elements = " ".join(str(c.value) for c in cells(wb["Report elements"]))
    assert "Phase 2" in elements and "Moving sample" in elements
    annex = " ".join(str(c.value) for c in cells(wb["Annex"]))
    assert "new fixture" in annex and "new limits from here" in annex
    data = wb["Data"]
    assert any(str(c.value) == "restart, new limits: new fixture" for c in data[42])  # row of value No. 41
    assert wb["Log"]["A2"].value == "chart restarted, new limits"


# ------------------------------------------------------------------ the Check sheet

formulas_lib = pytest.importorskip("formulas")


def evaluate_check(raw: bytes, tmp_path):
    path = tmp_path / "w.xlsx"
    path.write_bytes(raw)
    sol = formulas_lib.ExcelModel().loads(str(path)).finish().calculate()
    wb = load_workbook(path)
    ws = wb[wb.sheetnames[4]]
    out = {}
    for r in range(11, 30):
        if ws.cell(r, 1).value:
            get = lambda col: sol[f"'[w.xlsx]{ws.title.upper()}'!{col}{r}"].value[0, 0]
            out[ws.cell(r, 1).value] = (get("B"), ws.cell(r, 3).value, get("E"))
    return out


@pytest.mark.parametrize("request_kw", [{}, {"lsl": None}, {"usl": None}])
def test_excel_formulas_reproduce_the_numbers_of_the_program(tmp_path, request_kw):
    g = generate(sample(), req(**request_kw), None, "en", now="2026-10-01T00:00:00+00:00", report_id="x1")
    rows = evaluate_check(report_xlsx(g.archive), tmp_path)
    assert len(rows) >= 6
    for label, (excel, program, verdict) in rows.items():
        assert verdict == "agrees", (label, excel, program)
        assert excel == pytest.approx(program, rel=1e-6, abs=1e-9)
    assert any("interval" in k for k in rows) and any("ppm" in k for k in rows)


def test_check_sheet_marks_a_difference_when_the_data_change(tmp_path):
    g = generate(sample(), req(), None, "en", now="2026-10-01T00:00:00+00:00", report_id="x1")
    wb = load_workbook(io.BytesIO(report_xlsx(g.archive)))
    wb["Data"]["C10"].value = 11.0  # somebody edits a value in the workbook
    path = tmp_path / "edited.xlsx"
    wb.save(path)
    sol = formulas_lib.ExcelModel().loads(str(path)).finish().calculate()
    ws = wb["Check"]
    verdicts = {ws.cell(r, 1).value: sol[f"'[edited.xlsx]CHECK'!E{r}"].value[0, 0] for r in range(12, 20) if ws.cell(r, 1).value}
    assert verdicts["Mean"] == "differs"


def test_check_sheet_says_it_does_not_check_a_non_normal_report():
    ds = Dataset.from_values(9 + np.random.default_rng(3).lognormal(0, 0.4, 125))
    g, wb = workbook(ds, req(stage="preliminary", lsl=8.5, usl=13.0, model=None, distribution="lognormal", bootstrap_n=30))
    ws = wb["Check"]
    assert "fitted distribution" in str(ws["A3"].value) and "Lognormal" in str(ws["A3"].value)
    assert wb["Summary"]["A8"].value.endswith(".G")


def test_report_in_chinese_has_chinese_labels_and_the_same_numbers():
    g, wb = workbook(lang="zh-TW")
    assert wb["摘要"]["A3"].value == "報告編號"
    assert wb["摘要"]["B9"].value == pytest.approx(g.archive["result"]["indices"]["pk"])
    assert wb["驗算"]["A1"].value == "以 Excel 公式驗算"


# ------------------------------------------------------------------ dataset workbook

def test_dataset_workbook_has_data_marks_and_log():
    ds = sample().add_restart([10], "tool change", "A")
    wb = load_workbook(io.BytesIO(dataset_xlsx(ds, "en")))
    assert wb.sheetnames == ["Data", "Log"]
    ws = wb["Data"]
    assert ws.max_row == 126 and ws["E5"].value == "marked invalid" and ws["D5"].value == 0 and ws["D2"].value == 1
    assert sum(ws.cell(r, 4).value for r in range(2, 127)) == 124
    assert [wb["Log"].cell(r, 1).value for r in (2, 3)] == ["marked invalid", "chart restarted"]
    assert wb["Log"]["D2"].value == EVIL and wb["Log"]["D2"].data_type == "s"
    assert load_workbook(io.BytesIO(dataset_xlsx(ds, "zh-TW"))).sheetnames == ["資料", "紀錄"]


# ------------------------------------------------------------------ API

@pytest.fixture
def client():
    return logged_in_client(make_app(max_upload=200_000))


def test_report_workbook_download(client):
    ds = upload(client).json()
    rid = client.post(f"/api/datasets/{ds['id']}/reports",
                      json={"analysis": {"lsl": 9, "usl": 11, "model": "A1", "characteristic_class": "major"}, "language": "en"}).json()["id"]
    r = client.get(f"/api/reports/{rid}/report.xlsx")
    assert r.status_code == 200 and r.headers["content-type"] == XLSX_TYPE
    assert f"spc-report-{rid}.xlsx" in r.headers["content-disposition"] and r.headers["cache-control"] == "no-store"
    wb = load_workbook(io.BytesIO(r.content))
    assert wb["Summary"]["B3"].value == rid
    assert client.get("/api/reports/nope/report.xlsx").status_code == 404


def test_dataset_workbook_download_and_language(client):
    ds = upload(client).json()
    r = client.get(f"/api/datasets/{ds['id']}/export.xlsx")
    assert r.status_code == 200 and r.headers["content-type"] == XLSX_TYPE
    assert load_workbook(io.BytesIO(r.content)).sheetnames == ["Data", "Log"]
    zh = client.get(f"/api/datasets/{ds['id']}/export.xlsx", params={"lang": "zh-TW"})
    assert load_workbook(io.BytesIO(zh.content)).sheetnames == ["資料", "紀錄"]
    assert client.get(f"/api/datasets/{ds['id']}/export.xlsx", params={"lang": "fr"}).status_code == 422
    assert client.get("/api/datasets/nope/export.xlsx").status_code == 404


def test_viewers_may_download_and_anonymous_users_may_not():
    app = make_app(max_upload=200_000)
    eng, view = logged_in_client(app, "eng"), logged_in_client(app, "view")
    ds = eng.post("/api/datasets", params={"value": "diameter", "subgroup": "lot"}, content=csv_text()).json()
    rid = eng.post(f"/api/datasets/{ds['id']}/reports", json={"analysis": {"lsl": 9, "usl": 11}}).json()["id"]
    assert view.get(f"/api/reports/{rid}/report.xlsx").status_code == 200
    assert view.get(f"/api/datasets/{ds['id']}/export.xlsx").status_code == 200
    from fastapi.testclient import TestClient

    assert TestClient(app).get(f"/api/reports/{rid}/report.xlsx").status_code == 401
