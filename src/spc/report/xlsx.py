"""Excel output: the report as a workbook, and a dataset as a workbook.

Report workbook (sheets): Summary, Report elements (the 20+2 elements as rows), Data (every value with
its status), Control charts (plotted values and limits, with native Excel charts), Check (Excel formulas
that recalculate the normal-distribution results from the Data sheet and compare them with the
program), Annex (marked values, restarts, traceability).

Rules:
* Text that came from a person is always stored as a string. openpyxl would turn a text that starts
  with "=" into a formula. `_text` sets the cell type explicitly, so such a text stays text.
* Numbers are stored as numbers. Percent and ppm values are not rounded in the cell, only displayed rounded.
* Cw/Cwk never appear here, as in the HTML report.
* Formulas on the Check sheet use only functions that Excel 2007 and LibreOffice both know.
  Excel calculates them when the file is opened. The Program column next to them holds the numbers of the
  report, so a viewer that does not calculate formulas still shows the report's numbers.
"""

from __future__ import annotations

import io

import numpy as np
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from spc.data.dataset import Dataset
from spc.profile import ReportTemplate
from spc.report.builder import Report
from spc.report.texts import T

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="DCE6F1")
INPUT_FONT = Font(name=FONT, size=10, color="0000FF")  # blue: an input, as in financial models
THIN = Side(style="thin", color="BFBFBF")
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _font(bold=False, size=10, color=None) -> Font:
    return Font(name=FONT, size=size, bold=bold, color=color)


def _text(ws, row: int, col: int, value, bold=False, wrap=False):
    """A text cell. The type is set to string, so a text such as '=1+1' is shown and never calculated."""
    c = ws.cell(row=row, column=col, value=str(value))
    c.data_type = "s"
    c.font = _font(bold)
    c.alignment = Alignment(vertical="top", wrap_text=wrap)
    return c


def _num(ws, row: int, col: int, value, fmt: str | None = None, font: Font | None = None):
    c = ws.cell(row=row, column=col, value=None if value is None else float(value))
    c.font = font or _font()
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(vertical="top")
    return c


def _any(ws, row: int, col: int, value, fmt: str | None = None):
    if value is None or value == "":
        return ws.cell(row=row, column=col)
    if isinstance(value, bool):
        return _text(ws, row, col, value)
    if isinstance(value, (int, float, np.integer, np.floating)):
        return _num(ws, row, col, value, fmt)
    return _text(ws, row, col, value, wrap=True)


def _header(ws, row: int, labels, start_col: int = 1) -> None:
    for i, label in enumerate(labels):
        c = _text(ws, row, start_col + i, label, bold=True, wrap=True)
        c.fill = HEAD_FILL
        c.border = Border(bottom=THIN)


def _widths(ws, widths) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _stamp(s) -> str:
    return str(s).replace("T", " ")


def _status_rows(dataset: Dataset, used: set[int]):
    info = dataset.invalid_info()
    for i in range(dataset.n_total):
        yield i, ("invalid" if i in info else ("used" if i in used else "unused"))


def _data_sheet(wb: Workbook, dataset: Dataset, used: set[int], L):
    """Columns A No., B File row, C Value, D Used(1/0), E Status, F.. subgroup, time, tags, marks, restart."""
    ws = wb.create_sheet(L("x.sheet_data"))
    info, restarts, phases = dataset.invalid_info(), dataset.restart_info(), dataset.phase_positions()
    head = [L("x.col_pos"), L("f.col_row"), L("f.col_value"), L("x.col_used"), L("x.col_status")]
    if dataset.subgroup is not None:
        head.append(L("x.col_subgroup"))
    if dataset.timestamp is not None:
        head.append(L("x.col_time"))
    head += list(dataset.tags)
    head += [L("f.col_reason"), L("f.col_by"), L("f.col_at"), L("x.col_restart")]
    _header(ws, 1, head)
    for i, status in _status_rows(dataset, used):
        r = i + 2
        _num(ws, r, 1, i + 1, "0")
        _num(ws, r, 2, int(dataset.source_rows[i]), "0")
        _num(ws, r, 3, dataset.values[i])
        _num(ws, r, 4, 1 if status == "used" else 0, "0")
        _text(ws, r, 5, L(f"x.status_{status}"))
        col = 6
        if dataset.subgroup is not None:
            _text(ws, r, col, dataset.subgroup[i]); col += 1
        if dataset.timestamp is not None:
            _text(ws, r, col, _stamp(dataset.timestamp[i])); col += 1
        for tag in dataset.tags:
            _text(ws, r, col, dataset.tags[tag][i]); col += 1
        reason, by, at = info.get(i, ("", "", ""))
        for k, v in enumerate((reason, by, at)):
            if v:
                _text(ws, r, col + k, v)
        col += 3
        if i in restarts:
            _text(ws, r, col, f"{L('x.restart_phase' if i in phases else 'x.restart')}: {restarts[i][0]}")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(head))}{dataset.n_total + 1}"
    _widths(ws, [8, 10, 14, 12, 16] + [16] * (len(head) - 9) + [30, 22, 22, 36])
    return ws


def _log_sheet(wb: Workbook, dataset: Dataset, L):
    ws = wb.create_sheet(L("x.sheet_log"))
    _header(ws, 1, [L("x.col_action"), L("x.col_n"), L("x.col_positions"), L("f.col_reason"), L("f.col_by"), L("f.col_at")])
    for r, e in enumerate(dataset.log, start=2):
        _text(ws, r, 1, L(f"x.action_{e.action}"))
        _num(ws, r, 2, len(e.positions), "0")
        _text(ws, r, 3, ", ".join(str(p + 1) for p in e.positions))
        _text(ws, r, 4, e.reason, wrap=True)
        _text(ws, r, 5, e.by)
        _text(ws, r, 6, e.at)
    ws.freeze_panes = "A2"
    _widths(ws, [24, 8, 30, 40, 24, 22])
    return ws


# ---------------------------------------------------------------------------------------- report workbook

def _template(rep: Report) -> ReportTemplate:
    return ReportTemplate.from_dict(rep.profile["report"]) if rep.profile else ReportTemplate()


def _elements_rows(rep: Report, L) -> list[tuple[int | str, str, object, str | None]]:
    """(element no., item, value, number format) in the order of the report."""
    m, f, r = rep.meta, rep.facts, rep.result
    unit = f" {m.unit}" if m.unit else ""
    ix, tg, st = r["indices"], r["targets"], r["stability"]
    rows: list[tuple[int | str, str, object, str | None]] = []
    add = lambda no, item, value, fmt=None: rows.append((no, item, value, fmt))
    tpl = _template(rep)
    add(1, L("f.process"), m.process); add(1, L("f.machine"), m.machine); add(1, L("f.site"), m.site)
    add(2, L("f.process_ref"), m.process_ref); add(2, L("f.machine_ref"), m.machine_ref)
    add(3, L("el.3"), m.persons)
    if f["period"]:
        p = f["period"]
        add(4, L("f.start"), p["start"]); add(4, L("f.end"), p["end"]); add(4, L("f.duration"), p["duration"])
    else:
        add(4, L("el.4"), m.period_text)
    add(5, L("f.part_name"), m.part_name); add(5, L("f.part_number"), m.part_number)
    add(6, L("f.characteristic"), m.characteristic)
    if m.unit:
        add(6, L("f.unit"), m.unit)
    add(6, L("f.lsl") + unit, rep.request.lsl); add(6, L("f.usl") + unit, rep.request.usl)
    if f["tolerance"] is not None:
        add(6, L("f.tolerance") + unit, f["tolerance"])
    if m.target is not None:
        add(6, L("f.target") + unit, m.target)
    add(7, L("el.7"), m.technical_conditions); add(8, L("el.8"), m.deviations)
    add(9, L("f.n_total"), f["n_total"], "0"); add(9, L("f.n_marked"), f["n_marked"], "0")
    add(9, L("f.n_unused"), f["n_unused"], "0"); add(9, L("f.n_used"), f["n"], "0")
    add(10, L("f.subgroup_size"), f["subgroup_size"], "0"); add(10, L("f.subgroups"), f["k"], "0")
    if f["interval_min"] is not None:
        add(10, L("f.frequency_measured"), L("v.interval_min", value=f"{f['interval_min']:.1f}"))
    if m.sampling_frequency.strip():
        add(10, L("f.frequency_text"), m.sampling_frequency)
    if st["assessed"]:
        add(14, L("f.chart"), L("v.chart_" + f["chart_kind"]))
        add(14, L("f.risk"), f["alpha"], "0.000000")
        add(14, L("f.criteria"), "; ".join(f["criteria"]) or "–")
        add(14, L("f.stability_mode"), L("v.mode_" + r["params"]["stability_mode"]))
        add(14, L("f.stable"), L("v.class_" + st["class"]))
        add(14, "", L("v.points_checked", n=st["n_alarm_points"], k=st["n_points_checked"],
                      expected=f"{st['expected_false_alarms']:.2f}", threshold=st["threshold"]))
        if f["moving_n"]:
            add(14, L("f.moving"), L("v.moving_line", n=f["moving_n"], k=len(f["restarts"])))
        for i, p in enumerate(f["phases"]):
            add(14, L("f.phase", k=i + 1), L("v.phase_line", row=p["from_row"], n=p["n"], mu=f"{p['mu']:.5g}", sigma=f"{p['sigma']:.4g}"))
    else:
        add(14, L("el.14"), L("v.not_for_machine"))
    level = f["level"]
    add(15, L("f.mean"), f["mean"]); add(15, L("f.mean") + f" ({level}), " + L("x.ci_low"), f["mean_ci"][0])
    add(15, L("f.mean") + f" ({level}), " + L("x.ci_high"), f["mean_ci"][1])
    add(15, L("f.median"), f["q_mid"])
    add(15, L("f.method"), L("v.median_fitted" if f["dist"] else "v.mean_method"))
    add(16, L("f.sd"), f["sd"]); add(16, L("f.sd") + f" ({level}), " + L("x.ci_low"), f["sd_ci"][0])
    add(16, L("f.sd") + f" ({level}), " + L("x.ci_high"), f["sd_ci"][1])
    add(16, L("f.q_low"), f["q_low"]); add(16, L("f.q_high"), f["q_high"])
    add(16, L("f.spread"), f["q_high"] - f["q_low"]); add(16, L("f.method"), L("v.sd_method"))
    d = f["dist"]
    if d:
        params = "; ".join(f"{k} = {v if not isinstance(v, list) else v}" for k, v in d["params"].items())
        how = L("v.how_auto" if d["requested"] == "auto" else "v.how_manual")
        add(17, L("f.model"), L("v.fitted_model", name=L("v.dist_" + d["name"]), how=how, params=params))
        chosen = next((c for c in d["candidates"] if c["family"] == d["name"] and c["ok"]), None)
        if chosen:
            add(17, "", L("v.fit_quality", aic=f"{chosen['aic']:.1f}", ad=f"{chosen['ad']:.2f}"))
    else:
        add(17, L("f.model"), L("v.normal_assumed"))
    nm = f["normality"]
    if nm and nm["p_value"] is not None:
        add(17, L("f.normality"), L("v.p_value", test=nm["test"], p=f"{nm['p_value']:.4f}"))
    if tg and not tg.get("blocked"):
        add(18, L("f.target_class"), L("v.class_" + tg["class"]))
        if tg["p"] is not None and ix["p"] is not None:
            add(18, L("f.target_p", name=f["name_p"]), tg["p"], "0.00")
        add(18, L("f.target_pk", name=f["name_pk"]), tg["pk"], "0.00")
        add(18, L("f.target_base"), tg["n_base"], "0"); add(18, L("f.target_used"), tg["n"], "0")
        add(18, L("f.confidence_target"), tg["confidence"], "0.00%")
    elif tg:
        add(18, L("f.target_class"), L("v.class_" + tg["class"])); add(18, "", L("c.target_blocked", base=tg["n_base"]))
    else:
        add(18, "", L("c.no_target"))
    add(18, L("f.confidence_est"), level)
    if d:
        b = d["bootstrap"]
        interval = L("v.interval_boot", used=b["succeeded"], requested=b["requested"], seed=b["seed"]) if b and ix["ci_pk"] else L("v.interval_none")
        add(18, L("f.method"), L("v.method_" + f["method"], interval=interval))
    else:
        add(18, L("f.method"), L("v.normal_total"))
    if ix["p"] is not None:
        add(19, f["name_p"], ix["p"], "0.00")
    add(19, f["name_pk"], ix["pk"], "0.00")
    if f["ppm_below"] is not None:
        add(19, L("f.below_lsl") + " (ppm)", f["ppm_below"], "0.00")
    if f["ppm_above"] is not None:
        add(19, L("f.above_usl") + " (ppm)", f["ppm_above"], "0.00")
    add(19, L("f.total_out") + " (ppm)", f["ppm_total"], "0.00")
    for c in rep.conclusions:
        add(20, "", c)
    if m.recommendations.strip():
        add(20, L("c.recommendations"), m.recommendations)
    model = rep.request.model
    if tpl.show_element_21:
        add(21, L("f.model"), L("v.model_" + model) if model else L("v.model_unknown"))
        add(21, L("f.controlled"), L("v.yes" if rep.request.controlled_stable else "v.no"))
    g = f["guard"]
    if not tpl.show_element_22:
        pass
    elif g:
        add(22, L("f.uncertainty") + unit, g["U"]); add(22, L("f.coverage"), g["k"]); add(22, L("f.std_uncertainty") + unit, g["u"])
        add(22, L("f.guard_band", z=f"{g['z']:.3f}", risk=f"{g['risk'] * 100:g} %"), g["g"])
        if g["accept_low"] is not None:
            add(22, L("f.accept_low") + unit, g["accept_low"])
        if g["accept_high"] is not None:
            add(22, L("f.accept_high") + unit, g["accept_high"])
    else:
        add(22, L("el.22"), L("doc.not_given"))
    for fld in tpl.extra_fields:  # the customer's own fields
        add("C", fld.label(rep.lang), m.extra.get(fld.key, ""))
    return rows


def _summary_sheet(wb: Workbook, rep: Report, L) -> None:
    ws = wb.active
    ws.title = L("x.sheet_summary")
    f, r = rep.facts, rep.result
    ix, tg = r["indices"], r["targets"]
    tpl = _template(rep)
    _text(ws, 1, 1, tpl.title or L("doc.title"), bold=True).font = _font(True, 14)
    head = " · ".join(x for x in (tpl.organization, f"{L('doc.form')} {tpl.form_no}" if tpl.form_no else "",
                                  f"{L('v.revision')} {tpl.revision}" if tpl.revision else "") if x)
    _text(ws, 2, 1, (head + " — " if head else "") + L("doc.stage_" + f["stage"]))
    _text(ws, 3, 1, L("doc.report_id")); _text(ws, 3, 2, rep.report_id)
    _text(ws, 4, 1, L("doc.created")); _text(ws, 4, 2, _stamp(rep.generated_at))
    _text(ws, 5, 1, L("x.digest")); _text(ws, 5, 2, rep.archive_digest)
    _header(ws, 7, [L("f.index"), L("f.value"), L("x.ci_low"), L("x.ci_high"), L("f.target_col"), L("f.verdict")])
    ok = tg and not tg.get("blocked")
    lines = []
    if ix["p"] is not None:
        lines.append((f["name_p"], ix["p"], ix["ci_p"], tg["p"] if ok else None, tg["verdict_p"] if ok else None))
    lines.append((f["name_pk"], ix["pk"], ix["ci_pk"], tg["pk"] if ok else None, tg["verdict_pk"] if ok else None))
    for k, (name, value, ci, target, verdict) in enumerate(lines, start=8):
        _text(ws, k, 1, name, bold=True)
        _num(ws, k, 2, value, "0.00")
        _num(ws, k, 3, None if not ci else ci[0], "0.00"); _num(ws, k, 4, None if not ci else ci[1], "0.00")
        _num(ws, k, 5, target, "0.00")
        _text(ws, k, 6, L("v.verdict_" + verdict) if verdict else L("v.verdict_none"))
    row = 8 + len(lines) + 1
    _text(ws, row, 1, L("f.total_out") + " (ppm)", bold=True); _num(ws, row, 2, f["ppm_total"], "0.00")
    row += 2
    _text(ws, row, 1, L("el.20"), bold=True)
    for k, c in enumerate(rep.conclusions, start=row + 1):
        _text(ws, k, 1, c, wrap=False)
    _widths(ws, [34, 18, 16, 16, 12, 44])


def _elements_sheet(wb: Workbook, rep: Report, L) -> None:
    ws = wb.create_sheet(L("x.sheet_elements"))
    _header(ws, 1, [L("x.col_no"), L("x.col_element"), L("x.col_item"), L("x.col_value")])
    r = 2
    last = None
    for no, item, value, fmt in _elements_rows(rep, L):
        if no != last:
            if no == "C":
                _text(ws, r, 1, "–", bold=True)
                _text(ws, r, 2, L("doc.customer_fields", name=rep.profile["name"]), bold=True)
            else:
                _num(ws, r, 1, no, "0", _font(True)); _text(ws, r, 2, L(f"el.{no}"), bold=True)
            last = no
        _text(ws, r, 3, item, wrap=True)
        _any(ws, r, 4, value, fmt).alignment = Alignment(vertical="top", wrap_text=True, horizontal="left")
        r += 1
    _text(ws, r + 1, 3, L("x.figures_note"), wrap=True)
    ws.freeze_panes = "A2"
    _widths(ws, [8, 38, 44, 70])


def _chart_sheet(wb: Workbook, rep: Report, L) -> None:
    chart = rep.result["chart"]
    if not rep.result["stability"]["assessed"] or "location" not in chart:
        return
    ws = wb.create_sheet(L("x.sheet_charts"))
    series_names = {"xbar-s": "fig.series.s", "xbar-r": "fig.series.r", "median-r": "fig.series.r", "imr": "fig.series.mr"}
    loc_name = L({"imr": "fig.series.individual", "median-r": "fig.series.median"}.get(chart["kind"], "fig.series.mean"))
    var_name = L(series_names[chart["kind"]])
    col = 1
    top = 3
    for part_key, name, title_key in (("location", loc_name, "x.chart_loc"), ("variation", var_name, "x.chart_var")):
        part = chart[part_key]
        n = len(part["values"])
        _text(ws, 1, col, f"{L(title_key)}: {name}", bold=True)
        _header(ws, 2, [L("x.col_label"), L("x.chart_point"), L("fig.lcl"), L("fig.cl"), L("fig.ucl")], col)
        lim = lambda key, i: part[key][i] if isinstance(part[key], list) else part[key]
        for i in range(n):
            _text(ws, top + i, col, part["labels"][i])
            _num(ws, top + i, col + 1, part["values"][i])
            for k, key in enumerate(("lcl", "center", "ucl")):
                _num(ws, top + i, col + 2 + k, lim(key, i))
        lc = LineChart()
        lc.title = f"{L(title_key)}: {name}"
        lc.height, lc.width = 8, 20
        data = Reference(ws, min_col=col + 1, max_col=col + 4, min_row=2, max_row=top + n - 1)
        lc.add_data(data, titles_from_data=True)
        lc.set_categories(Reference(ws, min_col=col, min_row=top, max_row=top + n - 1))
        names = [L("x.chart_point"), L("fig.lcl"), L("fig.cl"), L("fig.ucl")]
        for s, nm in zip(lc.series, names):
            s.tx = SeriesLabel(v=nm)
            s.smooth = False
        lc.series[0].marker.symbol = "circle"
        lc.series[0].marker.size = 4
        for s in lc.series[1:]:
            s.marker.symbol = "none"
        ws.add_chart(lc, f"M{3 if part_key == 'location' else 21}")  # right of both tables
        col += 6
    _widths(ws, [10, 12, 12, 12, 12, 4, 10, 12, 12, 12, 12])
    ws.freeze_panes = "A3"


def _check_sheet(wb: Workbook, rep: Report, dataset: Dataset, L) -> None:
    """Excel formulas that recompute the normal-distribution results from the Data sheet."""
    ws = wb.create_sheet(L("x.sheet_check"))
    f, r, req = rep.facts, rep.result, rep.request
    ix = r["indices"]
    _text(ws, 1, 1, L("x.check_title"), bold=True).font = _font(True, 12)
    _text(ws, 2, 1, L("x.check_note"), wrap=True)
    ws.merge_cells("A2:E2")
    ws.row_dimensions[2].height = 62
    if f["dist"]:
        _text(ws, 3, 1, L("x.check_not_normal", name=L("v.dist_" + f["dist"]["name"]), method=f["method"]), bold=True, wrap=True)
        ws.merge_cells("A3:E3")
        ws.row_dimensions[3].height = 48
    N = dataset.n_total + 1
    vals, used = f"Data!$C$2:$C${N}", f"Data!$D$2:$D${N}"
    _header(ws, 5, [L("x.check_item"), L("x.check_excel"), L("x.check_program"), L("x.check_diff"), L("x.check_result")])

    # inputs: blue. The limits and the confidence level are the ones of the analysis request
    _text(ws, 6, 1, L("f.lsl")); c = _num(ws, 6, 2, req.lsl, font=INPUT_FONT)
    _text(ws, 7, 1, L("f.usl")); _num(ws, 7, 2, req.usl, font=INPUT_FONT)
    _text(ws, 8, 1, L("x.check_conf")); _num(ws, 8, 2, r["params"]["estimate_confidence"], "0.00%", INPUT_FONT)
    _text(ws, 9, 1, L("x.check_tol")); _num(ws, 9, 2, 1e-6, "0.0E+00", INPUT_FONT)
    LSL, USL, CONF, TOL = "$B$6", "$B$7", "$B$8", "$B$9"

    spec = [
        # (row, label, formula, program value, number format)
        (11, L("x.check_n"), f"=SUM({used})", f["n"], "0"),
        (12, L("f.mean"), f"=SUMPRODUCT({vals},{used})/B11", f["mean"], "0.000000"),
        (13, L("f.sd"), f"=SQRT(SUMPRODUCT({used},({vals}-B12)^2)/(B11-1))", f["sd"], "0.000000"),
    ]
    row = 14
    if ix["p"] is not None:
        spec.append((row, f["name_p"], f'=IF(AND(ISNUMBER({LSL}),ISNUMBER({USL})),({USL}-{LSL})/(6*B13),"")', ix["p"], "0.0000"))
        row += 1
    pu = f'IF(ISNUMBER({USL}),({USL}-B12)/(3*B13),"")'
    pl = f'IF(ISNUMBER({LSL}),(B12-{LSL})/(3*B13),"")'
    spec.append((row, f["name_pk"], f"=MIN({pu},{pl})", ix["pk"], "0.0000"))
    pk_row = row
    row += 1
    p_row = pk_row - 1 if ix["p"] is not None else None
    if ix["p"] is not None and ix["ci_p"]:
        spec.append((row, L("x.check_ci_low", name=f["name_p"]), f"=B{p_row}*SQRT(CHIINV(1-(1-{CONF})/2,B11-1)/(B11-1))", ix["ci_p"][0], "0.0000")); row += 1
        spec.append((row, L("x.check_ci_high", name=f["name_p"]), f"=B{p_row}*SQRT(CHIINV((1-{CONF})/2,B11-1)/(B11-1))", ix["ci_p"][1], "0.0000")); row += 1
    if ix["ci_pk"]:
        half = f"NORMSINV(1-(1-{CONF})/2)*SQRT(1/(9*B11)+B{pk_row}^2/(2*(B11-1)))"
        spec.append((row, L("x.check_ci_low", name=f["name_pk"]), f"=B{pk_row}-{half}", ix["ci_pk"][0], "0.0000")); row += 1
        spec.append((row, L("x.check_ci_high", name=f["name_pk"]), f"=B{pk_row}+{half}", ix["ci_pk"][1], "0.0000")); row += 1
    ppm = (f"=(IF(ISNUMBER({LSL}),NORMDIST({LSL},B12,B13,TRUE),0)+IF(ISNUMBER({USL}),1-NORMDIST({USL},B12,B13,TRUE),0))*1000000")
    spec.append((row, L("x.check_ppm"), ppm, ix["ppm"], "0.00"))
    for rr, label, formula, program, fmt in spec:
        _text(ws, rr, 1, label)
        c = ws.cell(row=rr, column=2, value=formula)  # a real formula: the only formulas in this workbook
        c.font = _font()
        c.number_format = fmt
        _num(ws, rr, 3, program, fmt)
        d = ws.cell(row=rr, column=4, value=f"=B{rr}-C{rr}")
        d.font = _font()
        d.number_format = "0.0E+00"
        e = ws.cell(row=rr, column=5, value=f'=IF(ABS(D{rr})<={TOL}*MAX(1,ABS(C{rr})),"{L("x.check_agrees")}","{L("x.check_differs")}")')
        e.font = _font(True)
    _widths(ws, [44, 20, 18, 16, 14])
    wb.calculation.fullCalcOnLoad = True


def _annex_sheet(wb: Workbook, rep: Report, L) -> None:
    ws = wb.create_sheet(L("x.sheet_annex"))
    f = rep.facts
    r = 1
    _text(ws, r, 1, L("doc.annex_a"), bold=True); r += 1
    _header(ws, r, [L("f.col_row"), L("f.col_value"), L("f.col_reason"), L("f.col_by"), L("f.col_at")]); r += 1
    if rep.marked:
        for x in rep.marked:
            _num(ws, r, 1, x["source_row"], "0"); _num(ws, r, 2, x["value"])
            _text(ws, r, 3, x["reason"], wrap=True); _text(ws, r, 4, x["by"]); _text(ws, r, 5, x["at"]); r += 1
    else:
        _text(ws, r, 1, L("f.none")); r += 1
    if f["restarts"]:
        r += 1
        _text(ws, r, 1, L("f.restarts_title"), bold=True); r += 1
        _header(ws, r, [L("f.col_row"), L("f.col_reason"), L("f.col_phase"), L("f.col_by"), L("f.col_at")]); r += 1
        for x in f["restarts"]:
            _num(ws, r, 1, x["source_row"], "0"); _text(ws, r, 2, x["reason"], wrap=True)
            _text(ws, r, 3, L("v.new_limits") if x["new_limits"] else ""); _text(ws, r, 4, x["by"]); _text(ws, r, 5, x["at"]); r += 1
    r += 1
    _text(ws, r, 1, L("doc.annex_b"), bold=True); r += 1
    tr = rep.trace
    for label, value in ((L("f.engine"), tr["engine_version"]), (L("f.fingerprint"), tr["fingerprint"]),
                         (L("f.edition"), tr["edition"]), (L("f.stability_mode"), L("v.mode_" + tr["stability_mode"])),
                         (L("x.digest"), rep.archive_digest)):
        _text(ws, r, 1, label); _text(ws, r, 2, value); r += 1
    if tr.get("customer"):
        _text(ws, r, 1, L("f.customer")); _text(ws, r, 2, tr["customer"]); r += 1
    if rep.profile:
        p = rep.profile
        _text(ws, r, 1, L("f.profile")); _text(ws, r, 2, f"{p['name']} ({L('v.revision')} {p['revision']})"); r += 1
        _text(ws, r, 1, L("f.profile_deviations")); _text(ws, r, 2, ", ".join(p["deviations"]) or L("v.none")); r += 1
    if tr.get("source"):
        s = tr["source"]
        _text(ws, r, 1, "SHA-256"); _text(ws, r, 2, s["sha256"]); r += 1
    if rep.control_plan:
        from spc.report.render import plan_approval_lines, plan_head_line, plan_heads, plan_rows

        r += 1
        _text(ws, r, 1, L("doc.annex_c"), bold=True); r += 1
        _text(ws, r, 1, plan_head_line(rep.control_plan, L), wrap=False); r += 1
        _header(ws, r, plan_heads(L)); r += 1
        for row in plan_rows(rep.control_plan, L):
            for c, v in enumerate(row, start=1):
                _text(ws, r, c, v, wrap=True)
            r += 1
        r += 1
        _text(ws, r, 1, L("cp.approvals"), bold=True); r += 1
        for k, v in plan_approval_lines(rep.control_plan, L):
            _text(ws, r, 1, k); _text(ws, r, 2, v, wrap=True); r += 1
    _widths(ws, [34, 40, 40, 24, 22, 24, 24, 20])


def render_xlsx(rep: Report, dataset: Dataset) -> bytes:
    """The report as an Excel workbook. `dataset` is the data the report was made from."""
    L = lambda key, **p: T(rep.lang, key, **p)
    used = {int(p) for rows in rep.result["chart"]["location"]["positions"] for p in rows}  # values behind the plotted points
    wb = Workbook()
    _summary_sheet(wb, rep, L)
    _elements_sheet(wb, rep, L)
    _data_sheet(wb, dataset, used, L)
    _chart_sheet(wb, rep, L)
    _check_sheet(wb, rep, dataset, L)
    _annex_sheet(wb, rep, L)
    _log_sheet(wb, dataset, L)
    wb.properties.title = f"{L('doc.title')} {rep.report_id}"
    wb.properties.creator = "SPC"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------------------- dataset workbook

def dataset_xlsx(dataset: Dataset, lang: str = "en") -> bytes:
    """The data with every mark and restart, and the log. Valid values have 1 in the column 'Used'."""
    L = lambda key, **p: T(lang, key, **p)
    wb = Workbook()
    wb.remove(wb.active)
    used = set(range(dataset.n_total)) - set(dataset.invalid_info())
    _data_sheet(wb, dataset, used, L)
    _log_sheet(wb, dataset, L)
    wb.properties.title = L("x.dataset_title")
    wb.properties.creator = "SPC"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
